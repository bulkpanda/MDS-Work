"""
boh3_dds4_utils.py
Reusable query, processing, and report-building functions for the DDS4/BOH3 cohorts.

Usage in notebook:
    from boh3_dds4_utils import *
"""

import os
import re
import json
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from pathlib import Path
from collections import defaultdict
from xml.sax.saxutils import escape
from sqlalchemy import text

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, PageBreak, Paragraph, Spacer,
)
from reportlab.lib.units import inch
import matplotlib.dates as mdates
import seaborn as sns
# These come from your existing Utils.py / variableUtils.py — imported here
# so callers only need `from boh3_dds4_utils import *`.
from Utils import createTable, addPlotImage, getBannerDrawer, getmodeArgs, readDf, runDdl, toInt, autoFitColumns, autopct
import variableUtils


# ═══════════════════════════════════════════════════════════════════════════
# 1. Helper functions for Excel and PDF report generation
# ═══════════════════════════════════════════════════════════════════════════
def _sanitizeSheetName(name):
    safe = re.sub(r'[\[\]\:\*\?\/\\]', '', str(name or '')).strip()
    return (safe[:31] or "Sheet")


def _autoFitColumns(ws, minWidth=10, maxWidth=60):
    for colCells in ws.columns:
        maxLen = 0
        colLetter = get_column_letter(colCells[0].column)
        for cell in colCells:
            if cell.value is not None:
                maxLen = max(maxLen, len(str(cell.value)))
        ws.column_dimensions[colLetter].width = max(minWidth, min(maxWidth, maxLen + 2))


def _writeTitle(ws, row, title):
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    return row + 2


def _writeTable(ws, df, startRow, startCol=1, title=None):
    r = startRow
    if title:
        ws.cell(row=r, column=startCol, value=title).font = Font(bold=True, size=12)
        r += 1
    for j, col in enumerate(df.columns, start=startCol):
        c = ws.cell(row=r, column=j, value=str(col))
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
    for _, rowData in df.iterrows():
        for j, col in enumerate(df.columns, start=startCol):
            c = ws.cell(row=r, column=j, value=rowData[col])
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return r + 2

# ═══════════════════════════════════════════════════════════════════════════
# 2. Data processing — DDL, upsert, clinic standardisation
# ═══════════════════════════════════════════════════════════════════════════

_CREATE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS dds4_boh3_forms (
    assessmentId BIGINT NOT NULL,
    form_code TEXT NOT NULL,
    cohort TEXT,
    subject TEXT,
    type TEXT,
    completed BOOLEAN,
    datetimeUtc TIMESTAMPTZ,
    student_number BIGINT,
    student_name TEXT,
    student_email TEXT,

    formId BIGINT,
    version INT,
    clinic TEXT,
    assessorId BIGINT,
    assessor_name TEXT,
    rotation TEXT,
    createdAt TIMESTAMPTZ,
    updatedAt TIMESTAMPTZ,

    patient_data JSONB,
    student_data JSONB,
    assessor_data JSONB,
    student_config JSONB,
    assessor_config JSONB,
    external_clinic TEXT,
    additional_concerns TEXT,
    submitted_by_student BOOLEAN,
    submitted_by_assessor BOOLEAN,

    insertedAt TIMESTAMPTZ DEFAULT now(),
    PRIMARY KEY (assessmentId, form_code)
);

CREATE INDEX IF NOT EXISTS idx_dds4_boh3_forms_datetimeUtc ON dds4_boh3_forms (datetimeUtc);
CREATE INDEX IF NOT EXISTS idx_dds4_boh3_forms_cohort ON dds4_boh3_forms (cohort);
CREATE INDEX IF NOT EXISTS idx_dds4_boh3_forms_student_name ON dds4_boh3_forms (student_name);

"""

def getClinicStandardizationSql(fromNames: list[str], toName: str, tableName: str = "dds4_boh3_forms"):
    fromNamesList = ", ".join(f"'{n}'" for n in fromNames)
    sql = f"""UPDATE {tableName}
          SET external_clinic = '{toName}'
          WHERE external_clinic IN (
            {fromNamesList}
          );"""
    return sql

def getBoh3Dds4FormsProcessSql(replace=False, tableName="dds4_boh3_forms"):
    """Return (createTableSql, upsertSql) for the dds4_boh3_forms pipeline."""
    if replace:
        onconflict = f"""
      ON CONFLICT (assessmentid, form_code) DO UPDATE SET
      cohort = EXCLUDED.cohort,
      subject = EXCLUDED.subject,
      type = EXCLUDED.type,
      completed = EXCLUDED.completed,
      datetimeutc = EXCLUDED.datetimeutc,
      student_number = EXCLUDED.student_number,
      student_name = EXCLUDED.student_name,
      student_email = EXCLUDED.student_email,

      formid = EXCLUDED.formid,
      version = EXCLUDED.version,
      assessorid = EXCLUDED.assessorid,
      assessor_name = EXCLUDED.assessor_name,
      rotation = EXCLUDED.rotation,
      createdat = EXCLUDED.createdat,
      updatedat = EXCLUDED.updatedat,

      patient_data = EXCLUDED.patient_data,
      student_data = EXCLUDED.student_data,
      assessor_data = EXCLUDED.assessor_data,
      student_config = EXCLUDED.student_config,
      assessor_config = EXCLUDED.assessor_config,

      external_clinic = EXCLUDED.external_clinic,
      additional_concerns = EXCLUDED.additional_concerns,
      submitted_by_student = EXCLUDED.submitted_by_student,
      submitted_by_assessor = EXCLUDED.submitted_by_assessor;
        """
    else:
        onconflict = "ON CONFLICT (assessmentId, form_code) DO NOTHING"

    createTableSql = _CREATE_TABLE_SQL.replace("dds4_boh3_forms", tableName)
    upsertSql = f"""
  INSERT INTO {tableName} (
    assessmentid, form_code,
    cohort, subject, type, completed, datetimeutc,
    student_number, student_name, student_email,
    formid, version, clinic, assessorid, assessor_name, rotation, createdat, updatedat,
    patient_data, student_data, assessor_data, student_config, assessor_config,
    external_clinic, additional_concerns, submitted_by_student, submitted_by_assessor
  )
  SELECT
    r.assessmentid,
    f.form_value->>'form_key' AS form_code,
    r.cohort, r.subject, r.type, r.completed, r.datetimeutc,
    r.student_number, r.student_name, r.student_email,
    NULLIF(f.form_value->>'id','')::bigint AS formid,
    NULLIF(f.form_value->>'version','')::int AS version,
    NULLIF(f.form_value->>'external_clinic','') AS clinic,
    NULLIF(f.form_value->>'assessor','')::bigint AS assessorId,
    f.form_value->>'assessor_name' AS assessor_name,
    f.form_value->>'rotation' AS rotation,
    NULLIF(f.form_value->>'created_at','')::timestamptz AS createdAt,
    NULLIF(f.form_value->>'updated_at','')::timestamptz AS updatedAt,
    f.form_value->'patient_data' AS patient_data,
    f.form_value->'student_data' AS student_data,
    f.form_value->'assessor_data' AS assessor_data,
    f.form_value->'student_config' AS student_config,
    f.form_value->'assessor_config' AS assessor_config,
    f.form_value->>'external_clinic' AS external_clinic,
    f.form_value->>'additional_concerns' AS additional_concerns,
    (f.form_value->>'submitted_by_student')::boolean AS submitted_by_student,
    (f.form_value->>'submitted_by_assessor')::boolean AS submitted_by_assessor
  FROM rawforms r
  CROSS JOIN LATERAL (
    SELECT a.value AS form_value
    FROM jsonb_array_elements(r.forms) a
    WHERE jsonb_typeof(r.forms) = 'array'
    UNION ALL
    SELECT e.value AS form_value
    FROM jsonb_each(r.forms) e
    WHERE jsonb_typeof(r.forms) = 'object'
  ) f
  WHERE r.cohort = ANY(:targetCohorts)
    AND (f.form_value ? 'form_key')
    {onconflict}
  """
    return createTableSql, upsertSql

# ═══════════════════════════════════════════════════════════════════════════
# 3. Cohort-level query functions
# ═══════════════════════════════════════════════════════════════════════════

def getWhereStatement(cohort, filters: dict = None):
    """Build a WHERE clause string + params dict from cohort and optional filters."""
    whereClauses = ["cohort = :cohort"]
    params = {"cohort": cohort}
    if filters:
        for key, value in filters.items():
            if isinstance(value, list):
                whereClauses.append(f"{key} IN :{key}")
            elif key.endswith("_min") or key.endswith("_max"):
                continue  # handled specially in callers like getTopItemCodes
            else:
                whereClauses.append(f"{key} = :{key}")
        params.update(filters)
    return " AND ".join(whereClauses), params


def _where(cohort, filters):
    """Shorthand: returns (whereClause, params)."""
    if filters:
        return getWhereStatement(cohort, filters)
    return "cohort = :cohort", {"cohort": cohort}


def getFullDf(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"SELECT * FROM {formsTable} WHERE {whereClause};"
    return readDf(engine, sql, params)


def getTotalForms(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"SELECT COUNT(*)::bigint AS totalForms FROM {formsTable} WHERE {whereClause};"
    return readDf(engine, sql, params)


def getAgeCounts(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT
      COUNT(*) FILTER (WHERE age BETWEEN 0 AND 6)::bigint  AS age0to6,
      COUNT(*) FILTER (WHERE age BETWEEN 7 AND 17)::bigint AS age7to17,
      COUNT(*) FILTER (WHERE age >= 18)::bigint            AS age18plus
    FROM (
      SELECT NULLIF(pd->>'patientAge','')::int AS age
      FROM {formsTable}
      CROSS JOIN LATERAL jsonb_array_elements(COALESCE(patient_data, '[]'::jsonb)) pd
      WHERE {whereClause}
      AND (pd->>'patientAttended')::boolean = true
    ) x;
    """
    return readDf(engine, sql, params)


def getAgeCountsBatch(engine, cohort, formsTable="dds4_boh3_forms"):
    """
    Batch version of getAgeCounts — one query for ALL students in the cohort.
    Returns a DataFrame with columns: student_name, age0to6, age7to17, age18plus.
    """
    sql = f"""
    SELECT
      f.student_name,
      COUNT(*) FILTER (WHERE NULLIF(pd->>'patientAge','')::int BETWEEN 0 AND 6)::bigint  AS age0to6,
      COUNT(*) FILTER (WHERE NULLIF(pd->>'patientAge','')::int BETWEEN 7 AND 17)::bigint AS age7to17,
      COUNT(*) FILTER (WHERE NULLIF(pd->>'patientAge','')::int >= 18)::bigint            AS age18plus
    FROM {formsTable} f
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(f.patient_data, '[]'::jsonb)) pd
    WHERE f.cohort = :cohort
      AND (pd->>'patientAttended')::boolean = true
      AND f.student_name IS NOT NULL
    GROUP BY f.student_name
    ORDER BY f.student_name;
    """
    return readDf(engine, sql, {"cohort": cohort})


def getAgeList(engine, cohort, formsTable="dds4_boh3_forms", filters=None, paramsAdd=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT NULLIF(pd->>'patientAge','')::int AS age
    FROM {formsTable}
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(patient_data, '[]'::jsonb)) pd
    WHERE {whereClause}
      AND NULLIF(pd->>'patientAge','') ~ '^\\d+$';
    """
    if paramsAdd:
        params.update(paramsAdd)
    return readDf(engine, sql, params)


def getAvgAge(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT ROUND(AVG(
      CASE
        WHEN (pd->>'patientAge') ~ '^\\s*\\d+\\s*$'
        AND (pd->>'patientAge')::int BETWEEN 0 AND 120
        THEN (pd->>'patientAge')::int
        ELSE NULL
      END
    )::numeric, 2) AS avgAge
    FROM {formsTable}
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(patient_data, '[]'::jsonb)) pd
    WHERE {whereClause};
    """
    return readDf(engine, sql, params)


def getPatientsPerStudentStats(engine, cohort, formsTable="dds4_boh3_forms", attended=True, filters=None):
    whereClause, params = _where(cohort, filters)
    params["attended"] = attended
    sql = f"""
    WITH perStudent AS (
      SELECT student_name, COUNT(*)::int AS patients
      FROM {formsTable}
      CROSS JOIN LATERAL jsonb_array_elements(COALESCE(patient_data, '[]'::jsonb)) pd
      WHERE {whereClause}
        AND student_name IS NOT NULL
        AND (pd->>'patientAttended')::boolean = :attended
      GROUP BY student_name
    )
    SELECT
      ROUND(AVG(patients)::numeric, 2) AS avgPatientsPerStudent,
      MIN(patients) AS minPatientsPerStudent,
      MAX(patients) AS maxPatientsPerStudent
    FROM perStudent;
    """
    return readDf(engine, sql, params)


def getPatientPerStudent(engine, cohort, formsTable="dds4_boh3_forms", attended=True, filters=None):
    whereClause, params = _where(cohort, filters)
    params["attended"] = attended
    sql = f"""
    SELECT
      f.student_name,
      COUNT(pd)::int AS patients
    FROM {formsTable} f
    LEFT JOIN LATERAL (
      SELECT *
      FROM jsonb_array_elements(
        CASE WHEN jsonb_typeof(f.patient_data) = 'array' THEN f.patient_data ELSE '[]'::jsonb END
      ) elem
      WHERE (elem->>'patientAttended')::boolean = :attended
    ) pd ON TRUE
    WHERE {whereClause}
      AND f.student_name IS NOT NULL
    GROUP BY f.student_name
    ORDER BY patients DESC, f.student_name;
    """
    return readDf(engine, sql, params)


def getClinicForStudent(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT student_name, COALESCE(NULLIF(external_clinic, ''), '(unknown)') AS clinic
    FROM {formsTable}
    WHERE {whereClause} AND student_name IS NOT NULL
    ORDER BY student_name, datetimeutc DESC;
    """
    return readDf(engine, sql, params)


def getClinicPatientCounts(engine, cohort, formsTable="dds4_boh3_forms", attended=True, filters=None):
    whereClause, params = _where(cohort, filters)
    params["attended"] = attended
    sql = f"""
    SELECT
      COALESCE(NULLIF(external_clinic,''), '(unknown)') AS clinic,
      COUNT(*)::bigint AS patients
    FROM {formsTable}
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(patient_data, '[]'::jsonb)) pd
    WHERE {whereClause}
      AND (pd->>'patientAttended')::boolean = :attended
    GROUP BY 1
    ORDER BY patients DESC, clinic;
    """
    return readDf(engine, sql, params)


def getSubmittedCounts(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT
      COUNT(*) FILTER (WHERE submitted_by_student IS TRUE)::bigint  AS submitted_By_Student,
      COUNT(*) FILTER (WHERE submitted_by_assessor IS TRUE)::bigint AS submitted_By_Assessor
    FROM {formsTable}
    WHERE {whereClause};
    """
    return readDf(engine, sql, params)


def getTopItemCodes(engine, cohort, formsTable="dds4_boh3_forms", limit=5, filters=None):
    whereClause, params = _where(cohort, filters)
    ageClause = ""
    if filters:
        if "age_min" in filters:
            ageClause += " AND NULLIF(pd->>'patientAge','')::int >= :age_min"
        if "age_max" in filters:
            ageClause += " AND NULLIF(pd->>'patientAge','')::int <= :age_max"
    sql = f"""
    SELECT ic->>'code' AS itemCode, COUNT(*)::bigint AS freq
    FROM {formsTable} f
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(f.patient_data, '[]'::jsonb)) pd
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(pd->'itemCodes', '[]'::jsonb)) ic
    WHERE {whereClause}
      AND ic ? 'code'
      {ageClause}
    GROUP BY 1
    ORDER BY freq DESC, itemCode
    LIMIT :limit;
    """
    params["limit"] = int(limit)
    return readDf(engine, sql, params)


def getItemCodesPerStudentBatch(engine, cohort, formsTable="dds4_boh3_forms"):
    """
    All item codes per student in one query. Can be used to create pivot tables later
    Returns: Student ID, Student Name, Item Code, Description, Total Qty
    """
    sql = f"""
    SELECT
      f.student_number AS "Student ID",
      f.student_name   AS "Student Name",
      ic->>'code'      AS "Item Code",
      MAX(ic->>'description') AS "Description",
      SUM(COALESCE(NULLIF(ic->>'quantity','')::int, 1))::int AS "Total Qty"
    FROM {formsTable} f
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(f.patient_data,'[]'::jsonb)) pd
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(pd->'itemCodes','[]'::jsonb)) ic
    WHERE f.cohort = :cohort
      AND ic ? 'code'
      AND (pd->>'patientAttended')::boolean = true
    GROUP BY f.student_number, f.student_name, ic->>'code'
    ORDER BY f.student_name, "Total Qty" DESC;
    """
    return readDf(engine, sql, {"cohort": cohort})


def getCafFinalEvalScoreStudent(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    WITH itemScores AS (
      SELECT f.assessmentid, f.form_code, f.cohort, kv.key AS mc,
        CASE kv.value
          WHEN 'Done well'      THEN 1.0
          WHEN 'Done'           THEN 0.8
          WHEN 'Mostly done'    THEN 0.6
          WHEN 'Sometimes done' THEN 0.4
          WHEN 'Not done'       THEN 0.0
          ELSE NULL
        END::numeric AS score
      FROM {formsTable} f
      CROSS JOIN LATERAL jsonb_each_text(
        COALESCE(f.student_data->'checklists'->'checklist-caf-final-eval', '{{}}'::jsonb)
      ) kv(key, value)
      WHERE {whereClause}
    )
    SELECT cohort, ROUND(AVG(score), 4)::numeric(6,4) AS avgStudentChecklistScore
    FROM itemScores WHERE score IS NOT NULL
    GROUP BY cohort;
    """
    return readDf(engine, sql, params)


def getAdditionalConcerns(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT cohort AS "Cohort", student_name AS "Student Name", date_trunc('day', datetimeutc) AS "Date",
           assessor_name AS "Assessor Name", additional_concerns AS "Additional Concerns"
    FROM {formsTable}
    WHERE {whereClause}
      AND additional_concerns IS NOT NULL
      AND TRIM(additional_concerns) <> '';
    """
    df = readDf(engine, sql, params)
    df["Date"] = pd.to_datetime(df["Date"]).dt.date
    return df


def getClinicalIncidentSummary(engine, cohort, formsTable="dds4_boh3_forms", filters=None, extractValue='clinical-incident', colName='Clinical Incidents'):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT
      f.cohort AS "Cohort", date_trunc('day', f.datetimeutc) AS "Date",
      f.assessor_name AS "Assessor Name", f.student_name AS "Student Name",
      STRING_AGG(ci->>'value', '; ' ORDER BY ci->>'name') AS "{colName}" 
    FROM {formsTable} f
    LEFT JOIN LATERAL jsonb_array_elements(
      f.assessor_data->'multi-select'->'{extractValue}'
    ) AS ci ON TRUE
    WHERE {whereClause} AND ci IS NOT NULL
    GROUP BY f.cohort, date_trunc('day', f.datetimeutc), f.assessor_name, f.student_name
    ORDER BY date_trunc('day', f.datetimeutc), f.assessor_name;
    """
    df = readDf(engine, sql, params)
    df["Date"] = pd.to_datetime(df["Date"]).dt.date
    return df


def getPatientCountPerRow(engine, cohort, formsTable="dds4_boh3_forms", attended=True, filters=None):
    whereClause, params = _where(cohort, filters)
    params["attended"] = attended
    sql = f"""
    SELECT assessmentid, form_code, student_name, external_clinic, rotation,
           COUNT(pd) AS patients
    FROM {formsTable}
    LEFT JOIN LATERAL (
      SELECT * FROM jsonb_array_elements(
        CASE WHEN jsonb_typeof(patient_data) = 'array' THEN patient_data ELSE '[]'::jsonb END
      ) elem
      WHERE (elem->>'patientAttended')::boolean = :attended
    ) pd ON TRUE
    WHERE {whereClause}
    GROUP BY assessmentid, form_code, student_name;
    """
    return readDf(engine, sql, params)


def getEntrustmentSummary(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT
      CASE NULLIF(assessor_data->'scales'->'scale-entrustment'->>'scale','')
        WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4
      END::smallint AS entrustment,
      COUNT(*)::bigint AS cnt
    FROM {formsTable}
    WHERE {whereClause}
      AND assessor_data->'scales'->'scale-entrustment'->>'scale' IS NOT NULL
    GROUP BY entrustment
    ORDER BY entrustment;
    """
    return readDf(engine, sql, params)


def getAvgEntrustment(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT ROUND(AVG(
      CASE NULLIF(assessor_data->'scales'->'scale-entrustment'->>'scale','')
        WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4 ELSE NULL
      END
    )::numeric, 2) AS avgEntrustment
    FROM {formsTable}
    WHERE {whereClause};
    """
    return readDf(engine, sql, params)


def getEntrustmentPerStudentBatch(engine, cohort, formsTable="dds4_boh3_forms"):
    sql = f"""
    WITH base AS (
      SELECT
        student_number,
        student_name,
        CASE NULLIF(assessor_data->'scales'->'scale-entrustment'->>'scale','')
          WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4
        END::smallint AS entrustment,
        assessor_data->'multi-select' AS multi
      FROM {formsTable}
      WHERE cohort = :cohort
      --  AND submitted_by_assessor
    )
    SELECT
      student_number                          AS "Student ID",
      student_name                            AS "Student Name",

      COUNT(*) FILTER (WHERE entrustment = 1) AS "Entrustment Lvl 1",
      COUNT(*) FILTER (WHERE entrustment = 2) AS "Entrustment Lvl 2",
      COUNT(*) FILTER (WHERE entrustment = 3) AS "Entrustment Lvl 3",
      COUNT(*) FILTER (WHERE entrustment = 4) AS "Entrustment Lvl 4",
      ROUND(AVG(entrustment)::numeric, 2)     AS "Entrustment Avg",

      
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-timeliness','[]'::jsonb))),0)::int AS "Weakness Timeliness",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-communication','[]'::jsonb))),0)::int AS "Weakness Communication",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-technical-skills','[]'::jsonb))),0)::int AS "Weakness Technical",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-person-centered-care','[]'::jsonb))),0)::int AS "Weakness PCC",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-professional-behaviour','[]'::jsonb))),0)::int AS "Weakness Professional",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-risk-management','[]'::jsonb))),0)::int AS "Weakness Risk Mgmt",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-knowledge-clinical-reasoning','[]'::jsonb))),0)::int AS "Weakness Clinical Reasoning",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-other','[]'::jsonb))),0)::int AS "Weakness Other",
      COALESCE(SUM(jsonb_array_length(COALESCE(multi->'strengths','[]'::jsonb))),0)::int AS "Commendations",
      COUNT(*) AS "Total Forms"
    FROM base
    GROUP BY student_number, student_name
    ORDER BY student_name;
    """
    return readDf(engine, sql, {"cohort": cohort})


def getNotSubmitted(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT assessmentid, form_code, student_name, assessor_name,
           datetimeutc::date AS datetimeutc, submitted_by_student, submitted_by_assessor,
           external_clinic, rotation
    FROM {formsTable}
    WHERE {whereClause}
      AND (NOT submitted_by_student OR NOT submitted_by_assessor);
    """
    return readDf(engine, sql, params)


def createPatientStatsPerClinic(engine, cohort, formsTable="dds4_boh3_forms", attended=True, filters=None):
    countsperrow = getPatientCountPerRow(engine, cohort, formsTable, attended, filters)
    studentClinicDf = countsperrow.groupby(
        ["external_clinic", "student_name"], as_index=False
    ).agg(totalPatients=("patients", "sum"))
    clinicStats = (
        studentClinicDf.groupby("external_clinic")["totalPatients"]
        .agg(avg="mean", min="min", max="max")
        .reset_index()
    )
    clinicStats["patientsSummary"] = (
        clinicStats["avg"].astype(int).astype(str)
        + " ("
        + clinicStats["min"].astype(str)
        + "-"
        + clinicStats["max"].astype(str)
        + ")"
    )
    return clinicStats[["external_clinic", "patientsSummary"]], studentClinicDf


def createPatientStatsPerClinicPerRotation(engine, cohort, formsTable="dds4_boh3_forms", attended=True):
    filepath = f"BOH3_DDS4/patient_stats_by_clinic_{cohort}.xlsx"
    studentclinicfilepath = f"BOH3_DDS4/patient_stats_by_clinic_{cohort}_detailed.xlsx"
    for rotation in ["Rotation 1", "Rotation 2", "Rotation 3"]:
        kwargs = getmodeArgs(filepath)
        kwargs2 = getmodeArgs(studentclinicfilepath)
        filters = {"rotation": rotation}
        resultDf, studentClinicDf = createPatientStatsPerClinic(
            engine, cohort, filters=filters, formsTable=formsTable, attended=attended
        )
        with pd.ExcelWriter(filepath, **kwargs) as writer:
            resultDf.to_excel(writer, sheet_name=rotation, index=False)
        with pd.ExcelWriter(studentclinicfilepath, **kwargs2) as writer:
            studentClinicDf.to_excel(writer, sheet_name=rotation, index=False)


def getCounts(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    whereClause, params = _where(cohort, filters)
    sql = f"""
    SELECT
      COUNT(*)::bigint AS totalForms,
      COUNT(*) FILTER (WHERE submitted_by_student IS TRUE)::bigint  AS "Submitted By Student",
      COUNT(*) FILTER (WHERE submitted_by_assessor IS TRUE)::bigint AS "Submitted By Assessor"
    FROM {formsTable}
    WHERE {whereClause};
    """
    return readDf(engine, sql, params)



# ═══════════════════════════════════════════════════════════════════════════
# 4. Cohort summary report (PDF)
# ═══════════════════════════════════════════════════════════════════════════

def getFrontPageSummaryTable(engine, cohort, formsTable="dds4_boh3_forms", filters=None):
    """Build the single-row summary DataFrame for the cohort front page."""
    totalFormsDf = getTotalForms(engine, cohort, formsTable, filters)
    ageDf = getAgeCounts(engine, cohort, formsTable, filters)
    avgAgeDf = getAvgAge(engine, cohort, formsTable, filters)
    ptsDf = getPatientsPerStudentStats(engine, cohort, formsTable, filters=filters)
    clinicDf = getClinicPatientCounts(engine, cohort, formsTable, filters=filters)
    submitDf = getSubmittedCounts(engine, cohort, formsTable, filters=filters)
    topItemsDf = getTopItemCodes(engine, cohort, formsTable, limit=5, filters=filters)
    cafScoreDf = getCafFinalEvalScoreStudent(engine, cohort, formsTable, filters=filters)
    entrustmentDf = getEntrustmentSummary(engine, cohort, formsTable, filters=filters)

    totalForms = int(totalFormsDf.iloc[0]["totalforms"])
    age0to6 = int(ageDf.iloc[0]["age0to6"])
    age7to17 = int(ageDf.iloc[0]["age7to17"])
    age18plus = int(ageDf.iloc[0]["age18plus"])
    avgAge = float(avgAgeDf.iloc[0]["avgage"]) if not avgAgeDf.empty else None

    avgPts = ptsDf.iloc[0]["avgpatientsperstudent"]
    minPts = ptsDf.iloc[0]["minpatientsperstudent"]
    maxPts = ptsDf.iloc[0]["maxpatientsperstudent"]

    submittedStudent = int(submitDf.iloc[0]["submitted_by_student"])
    submittedAssessor = int(submitDf.iloc[0]["submitted_by_assessor"])
    cafScore = cafScoreDf.iloc[0]["avgstudentchecklistscore"] if not cafScoreDf.empty else None

    ageStr = f"0–6: {age0to6}<br/> 7–17: {age7to17}<br/> 18+: {age18plus}"
    patientStatsStr = f"{int(avgPts)} ({int(minPts)} - {int(maxPts)})"
    clinicStr = "<br/>".join(f"{row.clinic}: {row.patients}" for _, row in clinicDf.iterrows())
    topItemsStr = ", ".join(f"{row.itemcode}" for _, row in topItemsDf.iterrows())
    cafScoreStr = f"{cafScore:.2f}" if cafScore is not None else "NA"
    entrustmentStr = "<br/>".join(f"S{int(row.entrustment)}: {row.cnt}" for _, row in entrustmentDf.iterrows())

    summary = {
        "Total forms": totalForms,
        "Submitted by students": submittedStudent,
        "Submitted by assessors": submittedAssessor,
        "Patient age distribution": ageStr,
        "Average patient age": f"{avgAge:.2f}" if avgAge is not None else "NA",
        "Patients per student (avg/min/max)": patientStatsStr,
        "Clinic patient counts": clinicStr,
        "Top item codes": topItemsStr,
        "Avg CAF final eval score (student)": cafScoreStr,
        "Entrustment levels": entrustmentStr,
        "Avg Entrustment": getAvgEntrustment(engine, cohort).iloc[0, 0],
    }
    return pd.DataFrame([summary])


def plotItemCodeFrequencies(itemCodesDf, ax, uniColor=None):
    if uniColor is None:
        uniColor = variableUtils.uniColor
    ax.bar(itemCodesDf["itemcode"], itemCodesDf["freq"], color=uniColor)
    ax.set_title("Top Item Codes", fontsize=14, color=uniColor)
    ax.set_xlabel("")
    ax.set_ylabel("Frequency", fontsize=10, color=uniColor)
    ax.set_xticklabels(itemCodesDf["itemcode"], rotation=90, ha="right", fontsize=8, color=uniColor)
    ax.set_ylim(0, itemCodesDf["freq"].max() * 1.2)
    for i, row in itemCodesDf.iterrows():
        ax.text(i, row["freq"], str(row["freq"]), ha="center", va="bottom", fontsize=6, color=uniColor)


def buildFrontPage(engine, cohort, filters, elements, subheadingColor,
                   subheadingStyle, tableTextStyleSmall, uniColor, figSize):
    metrics = getFrontPageSummaryTable(engine, cohort, filters=filters).transpose()
    metrics = metrics.reset_index()
    metrics.columns = ["Metric", "Value"]

    summaryTable = createTable(
        metrics, colRatio=[2, 1], customTextCols=[0, 1], bottomPadding=6,
        topPadding=6, title="", titleStyle=subheadingStyle, headerColor=subheadingColor,
        tableTextStyle=tableTextStyleSmall,
    )
    elements.append(summaryTable)

    allItemCodesDf = getTopItemCodes(engine, cohort, limit=25, filters=filters)
    fig, ax = plt.subplots(figsize=(figSize[0], figSize[1] / 4), dpi=200)
    plotItemCodeFrequencies(allItemCodesDf, ax, uniColor=uniColor)
    plt.close(fig)
    img = addPlotImage(fig, 0.9)
    elements.append(Spacer(1, 24))
    elements.append(img)


def buildCohortSummaryPdf(*, engine, cohort, outPath, bannerTitle,
                          subheadingStyle, subheadingColor, concernsDf, incidentsDf,
                          patientPerStudentDf, pageSize, rightMargin, leftMargin,
                          topMargin, bottomMargin, tableTextStyleSmall, uniColor, figSize,
                          superExcelPath=None):
    elements = []
    elements.append(Spacer(1, 72))
    doc = SimpleDocTemplate(outPath, pagesize=pageSize, rightMargin=rightMargin,
                            leftMargin=leftMargin, topMargin=topMargin, bottomMargin=bottomMargin)

    buildFrontPage(engine, cohort, filters=None, elements=elements,
                   subheadingColor=subheadingColor, subheadingStyle=subheadingStyle,
                   tableTextStyleSmall=tableTextStyleSmall, uniColor=uniColor, figSize=figSize)
    elements.append(PageBreak())

    # Item codes by age group
    ageGroups = [("0-6", {"age_min": 0, "age_max": 6}),
                 ("7-17", {"age_min": 7, "age_max": 17}),
                 ("18+", {"age_min": 18, "age_max": 120})]
    fig, axs = plt.subplots(3, 1, figsize=(figSize[0], figSize[1]), dpi=200)
    for ax, (ageGroupName, ageFilter) in zip(axs, ageGroups):
        itemCodeDf = getTopItemCodes(engine, cohort, limit=20, filters=ageFilter)
        plotItemCodeFrequencies(itemCodeDf, ax, uniColor=uniColor)
        ax.set_title(f"Age {ageGroupName}", fontsize=12, color=uniColor)
    plt.suptitle("Top Item Codes by Age Group", fontsize=14, color=uniColor)
    plt.subplots_adjust(hspace=0.5)
    plt.close(fig)
    img = addPlotImage(fig, 0.9)
    elements.append(img)
    elements.append(PageBreak())

    # Per-rotation summaries
    for rotation in ["Rotation 1", "Rotation 2", "Rotation 3", "Rotation 4"]:
        elements.append(Paragraph(f"{rotation} Summary", subheadingStyle))
        buildFrontPage(engine, cohort, filters={"rotation": rotation}, elements=elements,
                       subheadingColor=subheadingColor, subheadingStyle=subheadingStyle,
                       tableTextStyleSmall=tableTextStyleSmall, uniColor=uniColor, figSize=figSize)
        elements.append(PageBreak())

    # Concerns and incidents
    hasConcerns = concernsDf is not None and not concernsDf.empty
    hasIncidents = incidentsDf is not None and not incidentsDf.empty
    if hasConcerns or hasIncidents:
        if hasConcerns:
            concernsTable = createTable(
                concernsDf, colRatio=[1, 1, 1, 1, 3],
                customTextCols=list(range(concernsDf.shape[1])),
                bottomPadding=6, topPadding=6, title="Additional Concerns",
                titleStyle=subheadingStyle, headerColor=subheadingColor,
                tableTextStyle=tableTextStyleSmall,
            )
            elements.append(concernsTable)
            elements.append(Spacer(1, 24))
        if hasIncidents:
            incidentsTable = createTable(
                incidentsDf, colRatio=[1, 1, 1, 1, 3],
                customTextCols=list(range(incidentsDf.shape[1])),
                bottomPadding=6, topPadding=6, title="Clinical Incidents",
                titleStyle=subheadingStyle, headerColor=subheadingColor,
                tableTextStyle=tableTextStyleSmall,
            )
            elements.append(incidentsTable)

    # Patient count per student — BATCH age counts (no more N+1 loop)
    ageCountsDf = getAgeCountsBatch(engine, cohort)
    ageCountsDf['Total Patients'] = ageCountsDf['age0to6'] + ageCountsDf['age7to17'] + ageCountsDf['age18plus']
    ageCountsDf.sort_values('Total Patients', ascending=False, inplace=True)
    entrustmentDf = getEntrustmentPerStudentBatch(engine, cohort) # batch entrustment data
    entrustmentDf.sort_values('Entrustment Avg', ascending=False, inplace=True)

    # Load section mapping and merge with item codes to get section-level summaries later if needed
    itemsDf = getItemCodesPerStudentBatch(engine, cohort)
    mappingDf = pd.read_excel(variableUtils.itemSectionMappingFile)
    mappingDf["Item Code"] = mappingDf["Item Code"].astype(str).str.strip()
    itemsDf["Item Code"] = itemsDf["Item Code"].astype(str).str.strip()

    # Merge and pivot by section
    #Extract first code for mapping (handles "022/024" → "022")
    merged = itemsDf.copy()
    merged["Mapping Code"] = merged["Item Code"].str.split("/").str[0].str.strip()
    merged = merged.merge(mappingDf[["Item Code", "Section"]], left_on="Mapping Code", right_on="Item Code", how="left", suffixes=("", "_map"))
    merged.drop(columns=["Item Code_map", "Mapping Code"], inplace=True)
    merged["Section"] = merged["Section"].fillna("Unmapped")


    sectionPivot = merged.pivot_table(index=["Student Name"], columns="Section", values="Total Qty", aggfunc="sum", fill_value=np.nan).reset_index()
    sectionPivot.columns.name = None
    sectionPivot["Total"] = sectionPivot.iloc[:, 1:].sum(axis=1)
    sectionPivot.sort_values("Student Name", inplace=True)

    # save to excel for checking
    if superExcelPath:
        modeArgs = getmodeArgs(superExcelPath)
        with pd.ExcelWriter(superExcelPath, **modeArgs) as writer:
            ageCountsDf.to_excel(writer, sheet_name=f"Age Counts {cohort}", index=False)
            entrustmentDf.to_excel(writer, sheet_name=f"Entrustment {cohort}", index=False)
            incidentsDf.to_excel(writer, sheet_name=f"Incidents {cohort}", index=False)
            concernsDf.to_excel(writer, sheet_name=f"Concerns {cohort}", index=False)
            merged.to_excel(writer, sheet_name=f"Merged Item-Section {cohort}", index=False)
            sectionPivot.to_excel(writer, sheet_name=f"Section Pivot {cohort}", index=False)
        
        # autofit columns in the saved Excel file
        with pd.ExcelWriter(superExcelPath, engine='openpyxl', mode='a') as writer:
          workbook = writer.book
          for sheet_name in writer.sheets:
              ws = writer.sheets[sheet_name]
              autoFitColumns(ws)

    ageCountsDf["Age Count"] = (
        "0-6: " + ageCountsDf["age0to6"].astype(str)
        + "<br/> 7-17: " + ageCountsDf["age7to17"].astype(str)
        + "<br/> 18+: " + ageCountsDf["age18plus"].astype(str)
    )
    patientPerStudentDf = patientPerStudentDf.merge(
        ageCountsDf[["student_name", "Age Count"]],
        how="left", on="student_name",
    )
    # merge avg entrustment level per student
    patientPerStudentDf = patientPerStudentDf.merge(
        entrustmentDf[["Student Name", "Entrustment Avg"]],
        how="left", left_on="student_name", right_on="Student Name",
    )
    # drop the redundant "Student Name" column from entrustmentDf after merge
    patientPerStudentDf.drop(columns=["Student Name"], inplace=True)
    patientPerStudentDf.rename(columns={"student_name": "Student Name", "patients": "Patient Count", "patients_attended": "Attended",
                                        'patients_fta': "FTA", "Entrustment Avg": "Average ES"}, inplace=True)
    patientPerStudentTable = createTable(
        patientPerStudentDf, colRatio=[2, 1, 1, 2, 1],
        customTextCols=list(range(patientPerStudentDf.shape[1])),
        bottomPadding=6, topPadding=6, title="Patient Count per Student",
        titleStyle=subheadingStyle, headerColor=subheadingColor,
        tableTextStyle=tableTextStyleSmall, tableWidth= 0.9
    )
    elements.append(PageBreak())
    elements.append(patientPerStudentTable)

    doc.build(elements, onFirstPage=getBannerDrawer(bannerTitle, ""))


# ═══════════════════════════════════════════════════════════════════════════
# 5. Student-level query functions
# ═══════════════════════════════════════════════════════════════════════════

def getStudentTopItemCodes(engine, cohort, studentNumber, limit=10, formsTable="dds4_boh3_forms"):
    sql = f"""
    SELECT
      ic->>'code' AS itemCode,
      MAX(ic->>'description') AS description,
      SUM(COALESCE(NULLIF((ic->>'quantity')::int, NULL), 1))::int AS totalQty
    FROM {formsTable} f
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(f.patient_data,'[]'::jsonb)) pd
    CROSS JOIN LATERAL jsonb_array_elements(COALESCE(pd->'itemCodes','[]'::jsonb)) ic
    WHERE f.cohort = :cohort
      AND f.student_number = :studentNumber
      AND ic ? 'code'
    GROUP BY 1
    ORDER BY totalQty DESC, itemCode
    LIMIT :limit;
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber, "limit": int(limit)})


def getStudentsInCohort(engine, cohort, formsTable="dds4_boh3_forms"):
    sql = f"""
    SELECT DISTINCT student_number, student_name
    FROM {formsTable}
    WHERE cohort = :cohort
      AND student_name IS NOT NULL AND student_name <> '' AND student_name <> 'Test Student'
    ORDER BY student_name;
    """
    return readDf(engine, sql, {"cohort": cohort})


def getStudentPatientSummary(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    sql = f"""
    WITH pats AS (
      SELECT f.student_number, f.student_name, pd AS patient,
             (pd->>'patientAttended')::boolean AS attended,
             NULLIF(pd->>'patientAge','')::int AS age
      FROM {formsTable} f
      CROSS JOIN LATERAL jsonb_array_elements(COALESCE(f.patient_data,'[]'::jsonb)) pd
      WHERE f.cohort = :cohort AND f.student_number = :studentNumber
    )
    SELECT
      COUNT(*) FILTER (WHERE attended)::int                          AS totalAttended,
      COUNT(*) FILTER (WHERE NOT attended)::int                      AS totalNotAttended,
      COUNT(*) FILTER (WHERE attended AND age BETWEEN 0 AND 6)::int  AS age0to6,
      COUNT(*) FILTER (WHERE attended AND age BETWEEN 7 AND 17)::int AS age7to17,
      COUNT(*) FILTER (WHERE attended AND age >= 18)::int            AS age18plus
    FROM pats;
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber})


def getStudentSelfSummary(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    sql = f"""
    WITH base AS (
      SELECT assessmentid, form_code, datetimeutc, subject, clinic, assessor_name,
        student_data->'texts'->>'reflection' AS student_reflection,
        NULLIF(student_data->'scales'->'scale-practice-readiness'->>'scale','') AS practice_readiness,
        student_data->'checklists'->'checklist-caf-final-eval' AS caf
      FROM {formsTable}
      WHERE cohort = :cohort AND student_number = :studentNumber AND submitted_by_student
    ),
    cafRows AS (
      SELECT b.assessmentid, b.form_code, b.datetimeutc, b.subject, b.clinic, b.assessor_name,
        b.student_reflection, b.practice_readiness, kv.key AS mc, kv.value AS ratingText,
        CASE kv.value
          WHEN 'Done well' THEN 1.0 WHEN 'Done' THEN 0.8 WHEN 'Mostly done' THEN 0.6
          WHEN 'Sometimes done' THEN 0.4 WHEN 'Not done' THEN 0.0 ELSE NULL
        END::numeric AS ratingScore
      FROM base b
      LEFT JOIN LATERAL jsonb_each_text(COALESCE(b.caf, '{{}}'::jsonb)) kv(key, value) ON TRUE
    )
    SELECT assessmentid, form_code, datetimeutc, subject, clinic, assessor_name,
           practice_readiness, student_reflection,
           ROUND(AVG(ratingScore), 3)::numeric(6,3) AS caf_avg_score
    FROM cafRows
    GROUP BY assessmentid, form_code, datetimeutc, subject, clinic, assessor_name,
             practice_readiness, student_reflection
    ORDER BY datetimeutc, assessmentid;
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber})


def getStudentAssessorSummary(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    sql = f"""
    WITH base AS (
      SELECT assessmentid, form_code, datetimeutc, subject, clinic, assessor_name,
        NULLIF(assessor_data->'scales'->'scale-entrustment'->>'scale','') AS entrustment_scale,
        NULLIF(additional_concerns,'') AS additional_concerns,
        assessor_data->'multi-select' AS multi
      FROM {formsTable}
      WHERE cohort = :cohort AND student_number = :studentNumber AND submitted_by_assessor
    ),
    extracted AS (
      SELECT b.*,
        (SELECT string_agg(x->>'value', E'\\n' ORDER BY x->>'value')
         FROM jsonb_array_elements(COALESCE(b.multi->'strengths','[]'::jsonb)) x) AS strengths_text,
        (SELECT string_agg(x->>'value', E'\\n' ORDER BY x->>'value')
         FROM jsonb_array_elements(COALESCE(b.multi->'weakness-other','[]'::jsonb)) x) AS weakness_other_text,
        (SELECT string_agg(x->>'value', E'\\n' ORDER BY x->>'value')
         FROM jsonb_array_elements(COALESCE(b.multi->'clinical-incident','[]'::jsonb)) x) AS clinical_incident_text,
        COALESCE(jsonb_array_length(COALESCE(b.multi->'clinical-incident','[]'::jsonb)),0) AS clinical_incident_count
      FROM base b
    )
    SELECT assessmentid, form_code, datetimeutc, subject, clinic, assessor_name,
      CASE entrustment_scale
        WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4 ELSE NULL
      END::smallint AS entrustment_numeric,
      additional_concerns, clinical_incident_count,
      strengths_text, weakness_other_text, clinical_incident_text
    FROM extracted
    ORDER BY datetimeutc, assessmentid;
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber})


def getStudentSummaryTable(engine, cohort, studentNumber, studentName, formsTable="dds4_boh3_forms"):
    patientDf = getStudentPatientSummary(engine, cohort, studentNumber, formsTable)
    selfDf = getStudentSelfSummary(engine, cohort, studentNumber, formsTable)
    assessorDf = getStudentAssessorSummary(engine, cohort, studentNumber, formsTable)
    
    totalCounts = getCounts(engine, cohort, formsTable, filters={"student_number": studentNumber})
    totalForms = int(totalCounts.iloc[0]["totalforms"])
    submittedByStudent = int(totalCounts.iloc[0]["Submitted By Student"])
    submittedByAssessor = int(totalCounts.iloc[0]["Submitted By Assessor"])
    # totalForms = int(readDf(engine, f"""
    #     SELECT COUNT(*)::int AS n
    #     FROM {formsTable}
    #     WHERE cohort=:cohort AND student_number=:studentNumber;
    # """, {"cohort": cohort, "studentNumber": studentNumber}).iloc[0]["n"])

    avgEntrustment = assessorDf["entrustment_numeric"].dropna().astype(float).mean()
    entrustmentDistribution = assessorDf["entrustment_numeric"].value_counts().sort_index()
    avgCaf = selfDf["caf_avg_score"].dropna().astype(float).mean()

    totalConcerns = assessorDf["additional_concerns"].dropna().shape[0]
    totalIncidents = int(assessorDf["clinical_incident_count"].fillna(0).sum())
    p = patientDf.iloc[0] if not patientDf.empty else pd.Series(dtype="object")

    metrics = [
        ("Total Forms", str(toInt(totalForms))),
        ("Submitted by Student", str(toInt(submittedByStudent))),
        ("Submitted by Assessor", str(toInt(submittedByAssessor))),
        ("Patients Attended", str(toInt(p.get("totalattended", 0) or 0))),
        ("Patients FTA", str(toInt(p.get("totalnotattended", 0) or 0))),
        ("Patients Age 0–6", str(toInt(p.get("age0to6", 0) or 0))),
        ("Patients Age 7–17", str(toInt(p.get("age7to17", 0) or 0))),
        ("Patients Age 18+", str(toInt(p.get("age18plus", 0) or 0))),
        ("Avg Self CAF Score", None if pd.isna(avgCaf) else round(avgCaf, 2)),
        ("Avg Assessor Entrustment (1–4)", None if pd.isna(avgEntrustment) else round(avgEntrustment, 2)),
        ("Entrustment Distribution (1–4)",
         "<br/>".join(f"Lvl {idx}: {cnt}" for idx, cnt in entrustmentDistribution.items())),
        ("Additional Concerns", str(toInt(totalConcerns))),
        ("Clinical Incidents", str(toInt(totalIncidents))),
    ]
    return pd.DataFrame(metrics, columns=["Metric", "Value"]), selfDf, assessorDf


def getStudentTimeSeries(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    """Per-form entrustment (assessor) and practice readiness (student) over time."""
    sql = f"""
    SELECT
      datetimeutc::date AS date,
      CASE NULLIF(assessor_data->'scales'->'scale-entrustment'->>'scale','')
        WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4
      END::smallint AS entrustment,
      CASE NULLIF(student_data->'scales'->'scale-practice-readiness'->>'scale','')
      WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4
      END::smallint AS practice_readiness
    FROM {formsTable}
    WHERE cohort = :cohort AND student_number = :studentNumber AND (submitted_by_student AND submitted_by_assessor)
    ORDER BY datetimeutc;
    """
        
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber})


def getStudentRollup(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    sql = f"""
    WITH base AS (
      SELECT student_data->'texts'->>'reflection' AS reflection,
             NULLIF(student_data->'scales'->'scale-practice-readiness'->>'scale','') AS readiness,
             student_data->'checklists'->'checklist-caf-final-eval' AS caf
      FROM {formsTable}
      WHERE cohort = :cohort AND student_number = :studentNumber
    ),
    caf_vals AS (
      SELECT CASE kv.value
          WHEN 'Done well' THEN 1.0 WHEN 'Done' THEN 0.8 WHEN 'Mostly done' THEN 0.6
          WHEN 'Sometimes done' THEN 0.4 WHEN 'Not done' THEN 0.0 ELSE NULL
        END::numeric AS score
      FROM base b
      LEFT JOIN LATERAL jsonb_each_text(COALESCE(b.caf,'{{}}'::jsonb)) kv(key, value) ON TRUE
    )
    SELECT
      (SELECT COUNT(*)::int FROM base) AS forms_count,
      (SELECT COUNT(*)::int FROM base WHERE NULLIF(reflection,'') IS NOT NULL) AS reflections_count,
      (SELECT ROUND(AVG(score),3)::numeric(6,3) FROM caf_vals WHERE score IS NOT NULL) AS caf_avg,
      (SELECT PERCENTILE_CONT(0.5) WITHIN GROUP (ORDER BY score) FROM caf_vals WHERE score IS NOT NULL) AS caf_median;
    """
    roll = readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber})

    sql2 = f"""
      SELECT
        student_config->'scales'->'scale-practice-readiness'->'fields'
          -> (student_data->'scales'->'scale-practice-readiness'->>'scale') AS readiness,
        COUNT(*) AS n
      FROM {formsTable}
      WHERE cohort = :cohort AND student_number = :studentNumber
      GROUP BY readiness
      ORDER BY readiness;
    """
    readinessDf = readDf(engine, sql2, {"cohort": cohort, "studentNumber": studentNumber})
    return roll, readinessDf


def getAssessorRollup(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    """
    Get a single-row summary of all assessor-submitted forms for a given student, including average entrustment level,
    counts of various multi-select categories, and number of forms with additional concerns. This is used
    to populate the student-level summary section of the PDF report, and is designed to be efficient by doing all calculations 
    in a single query with CTEs
    """
    sql = f"""
    WITH base AS (
      SELECT
        NULLIF(assessor_data->'scales'->'scale-entrustment'->>'scale','') AS entrustment_scale,
        NULLIF(additional_concerns,'') AS additional_concerns,
        assessor_data->'multi-select' AS multi
      FROM {formsTable}
      WHERE cohort=:cohort AND student_number=:studentNumber and submitted_by_assessor
    ),
    ent AS (
      SELECT CASE entrustment_scale
        WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4 ELSE NULL
      END::numeric AS entrustment_num FROM base
    ),
    counts AS (
      SELECT
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'strengths','[]'::jsonb))),0) AS strengths_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-other','[]'::jsonb))),0) AS weakness_other_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-timeliness','[]'::jsonb))),0) AS weakness_timeliness_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-communication','[]'::jsonb))),0) AS weakness_communication_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-technical-skills','[]'::jsonb))),0) AS weakness_technical_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-person-centered-care','[]'::jsonb))),0) AS weakness_pcc_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-professional-behaviour','[]'::jsonb))),0) AS weakness_professional_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-risk-management','[]'::jsonb))),0) AS weakness_risk_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'weakness-knowledge-clinical-reasoning','[]'::jsonb))),0) AS weakness_reasoning_n,
        COALESCE(SUM(jsonb_array_length(COALESCE(multi->'clinical-incident','[]'::jsonb))),0) AS incidents_n
      FROM base
    )
    SELECT
      (SELECT COUNT(*)::int FROM base) AS forms_count,
      (SELECT COUNT(*)::int FROM base WHERE additional_concerns IS NOT NULL) AS concerns_forms_n,
      (SELECT ROUND(AVG(entrustment_num),3)::numeric(6,3) FROM ent WHERE entrustment_num IS NOT NULL) AS entrustment_avg,
      (SELECT PERCENTILE_CONT(0.5) WITHIN GROUP (ORDER BY entrustment_num) FROM ent WHERE entrustment_num IS NOT NULL) AS entrustment_median,
      strengths_n, weakness_other_n, weakness_timeliness_n, weakness_communication_n,
      weakness_technical_n, weakness_pcc_n, weakness_professional_n, weakness_risk_n,
      weakness_reasoning_n, incidents_n
    FROM counts;
    """
    sql_entrustment = f"""
      SELECT
        assessor_config->'scales'->'scale-entrustment'->'fields'
          -> (assessor_data->'scales'->'scale-entrustment'->>'scale') AS "Entrustment",
        COUNT(*)::int AS "Count"
      FROM {formsTable}
      WHERE cohort = :cohort AND student_number = :studentNumber
        AND submitted_by_assessor
        AND assessor_data->'scales'->'scale-entrustment'->>'scale' IS NOT NULL
      GROUP BY "Entrustment"
      ORDER BY "Entrustment";
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber}), readDf(engine, sql_entrustment, {"cohort": cohort, "studentNumber": studentNumber})


def getTopMultiSelectValues(engine, cohort, studentNumber, key, limit=6, formsTable="dds4_boh3_forms"):
    sql = f"""
    WITH vals AS (
      SELECT x->>'value' AS v
      FROM {formsTable} f
      CROSS JOIN LATERAL jsonb_array_elements(COALESCE(f.assessor_data->'multi-select'->:key,'[]'::jsonb)) x
      WHERE f.cohort=:cohort AND f.student_number=:studentNumber AND x ? 'value'
    )
    SELECT v AS value, COUNT(*)::int AS n FROM vals GROUP BY v ORDER BY n DESC, v LIMIT :limit;
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber, "key": key, "limit": int(limit)})


def getSelfReflections(engine, cohort, studentNumber, formsTable="dds4_boh3_forms"):
    sql = f"""
    SELECT
      datetimeutc::date AS datetimeutc,
      student_data->'texts'->>'reflection' AS self_reflection,
      (SELECT string_agg(ic->>'code', ', ' ORDER BY ic->>'code')
       FROM jsonb_array_elements(COALESCE(f.patient_data, '[]'::jsonb)) pd
       CROSS JOIN jsonb_array_elements(COALESCE(pd->'itemCodes', '[]'::jsonb)) ic
       WHERE ic ? 'code') AS item_codes
    FROM {formsTable} f
    WHERE cohort=:cohort AND student_number=:studentNumber
      AND submitted_by_student
      AND NULLIF(student_data->'texts'->>'reflection','') IS NOT NULL
    ORDER BY datetimeutc;
    """
    return readDf(engine, sql, {"cohort": cohort, "studentNumber": studentNumber})


# ═══════════════════════════════════════════════════════════════════════════
# 6. Student PDF report builder
# ═══════════════════════════════════════════════════════════════════════════

def weaknessPiePlot(weaknessCounts, totalWeaknesses, uniColor):
    fig, ax = plt.subplots(figsize=(variableUtils.figSize[0]*0.55, variableUtils.figSize[1]*0.32), dpi=200)
    palette = sns.color_palette("Set2", n_colors=weaknessCounts.shape[0])
    wedges, texts, autotexts = ax.pie(
        weaknessCounts["Count"],
        labels=None,
        colors=palette,
        autopct=lambda pct: autopct(pct, totalWeaknesses),
        pctdistance=0.75,
        startangle=140,
        wedgeprops=dict(linewidth=0.8, edgecolor="white"),
    )
    for at in autotexts:
        at.set_fontsize(5)
        at.set_color("white")
        at.set_fontweight("bold")

    ax.legend(
        wedges, [f"{l} ({v})" for l, v in zip(weaknessCounts["Weakness Type"], weaknessCounts["Count"])],
        loc="center left",
        bbox_to_anchor=(1.0, 0.5),
        fontsize=6.5,
        frameon=False,
        labelcolor=uniColor,
    )
    fig.suptitle("Weaknesses", fontsize=11, color=uniColor, fontweight="bold", y=0.9)
    fig.tight_layout()
    img = addPlotImage(fig, 0.8)
    return img

def buildStudentVsAssessorSection(engine, cohort, studentNumber, elements, styles,
                                  formsTable="dds4_boh3_forms",
                                  subheadingStyle=None, uniColor=None, tableTextStyleSmall=None):
    if subheadingStyle is None:
        subheadingStyle = styles.get("subheadingStyle")
    if uniColor is None:
        uniColor = variableUtils.uniColor

    selfRoll, readinessDf = getStudentRollup(engine, cohort, studentNumber, formsTable)
    assessorRoll, entrustmentDf = getAssessorRollup(engine, cohort, studentNumber, formsTable)
    topStrengths = getTopMultiSelectValues(engine, cohort, studentNumber, "strengths", 10, formsTable)
    topWeaknessOther = getTopMultiSelectValues(engine, cohort, studentNumber, "weakness-other", 30, formsTable)
    # topIncidents = getTopMultiSelectValues(engine, cohort, studentNumber, "clinical-incident", 6, formsTable)
    clinicalIncidents = getClinicalIncidentSummary(engine, cohort, filters={"student_number": studentNumber}, formsTable=formsTable)
    clinicalIncidents.drop(columns=["Cohort", 'Student Name'], inplace=True)

    weaknessOther = getClinicalIncidentSummary(engine, cohort, filters={"student_number": studentNumber}, 
                                               formsTable=formsTable, extractValue='weakness-other', colName = 'Weakness/Strength')
    weaknessOther.drop(columns=["Cohort", 'Student Name'], inplace=True)
    elements.append(PageBreak())

    # Practice readiness distribution
    if not readinessDf.empty:
        rd = readinessDf.copy()
        rd.columns = ["Practice Readiness", "Count"]
        rd = rd[rd["Practice Readiness"].notna() & (rd["Practice Readiness"] != "")
               & (rd["Practice Readiness"].str.lower() != "none")]
        elements.append(createTable(
            rd, colRatio=[3, 1], customTextCols=[0, 1], bottomPadding=6, topPadding=6,
            title="Student Judgement: Practice Readiness Distribution",
            titleStyle=subheadingStyle, headerColor=uniColor,
        ))
        elements.append(Spacer(1, 18))

    # Assessor summary
    elements.append(createTable(entrustmentDf, colRatio=[3, 1], customTextCols=[0, 1], bottomPadding=6, topPadding=6,
                        title="Assessor Judgement: Entrustment level distribution", titleStyle=subheadingStyle, headerColor=uniColor))
    
    a = assessorRoll.iloc[0].to_dict() if not assessorRoll.empty else {}
    assessorMetrics = pd.DataFrame([
        # ("Forms assessed", toInt(a.get("forms_count"))),
        # ("Average entrustment (1–4)", toInt(a.get("entrustment_avg"))),
        ("# Additional concerns", toInt(a.get("concerns_forms_n"))),
        ("# Clinical incidents", toInt(a.get("incidents_n"))),
        ("# Commendations given", toInt(a.get("strengths_n"))),
    ], columns=["Metric", "Value"])
    elements.append(createTable(
        assessorMetrics, colRatio=[3, 1], customTextCols=[0, 1], bottomPadding=6, topPadding=6,
        title="", titleStyle=subheadingStyle, headerColor=uniColor,
    ))

    elements.append(Spacer(1, 18))

    # Weakness counts
    weaknessCounts = pd.DataFrame([
        ("Time management", toInt(a.get("weakness_timeliness_n"))),
        ("Communication", toInt(a.get("weakness_communication_n"))),
        ("Technical skills", toInt(a.get("weakness_technical_n"))),
        ("Person-centred care", toInt(a.get("weakness_pcc_n"))),
        ("Professional behaviour", toInt(a.get("weakness_professional_n"))),
        ("Risk management", toInt(a.get("weakness_risk_n"))),
        ("Knowledge & reasoning", toInt(a.get("weakness_reasoning_n"))),
        # ("Other weaknesses noted", toInt(a.get("weakness_other_n"))),
    ], columns=["Weakness Type", "Count"])
    weaknessCounts = weaknessCounts[weaknessCounts["Count"] > 0]
    totalWeaknesses = weaknessCounts["Count"].sum()

    img = weaknessPiePlot(weaknessCounts, totalWeaknesses, uniColor)
    elements.append(img)
    # elements.append(createTable(
    #     weaknessCounts, colRatio=[4, 1], customTextCols=[0, 1], bottomPadding=6, topPadding=6,
    #     title="Weaknesses", titleStyle=subheadingStyle, headerColor=uniColor,
    # ))
    elements.append(Spacer(1, 18))

    def renderTopDf(df, title, valuename, showCounts=True):
        if df.empty:
            return
        tmp = df.copy()
        tmp.iloc[:, 0] = tmp.iloc[:, 0].str.replace("\n", "<br/>", regex=False)
        tmp.columns = [valuename, "Count"]
        if not showCounts:
            tmp.drop(columns=["Count"], inplace=True)
            customTextCols = [0]
            colRatio = [1]
        else:
            customTextCols = [0, 1]
            colRatio = [4, 1]
        elements.append(createTable(
            tmp, colRatio=colRatio, customTextCols=customTextCols, bottomPadding=6, topPadding=6,
            title=title, titleStyle=subheadingStyle, headerColor=uniColor, tableTextStyle=tableTextStyleSmall,
        ))
        elements.append(Spacer(1, 12))

    renderTopDf(topStrengths, "Top commendations", "Commendation")
    # renderTopDf(topWeaknessOther, "Other weaknesses and strengths", "Weakness/Strength", showCounts=False)
    otherWeaknessTable = createTable(
        weaknessOther, colRatio=[1, 1, 4], customTextCols=[0, 1, 2], bottomPadding=6, topPadding=6,
        title="Other weaknesses/strengths", titleStyle=subheadingStyle, headerColor=uniColor,
        tableTextStyle=tableTextStyleSmall,
    )
    elements.append(otherWeaknessTable)
    elements.append(Spacer(1, 18))

    ciTable = createTable(clinicalIncidents, colRatio=[1, 1, 4], customTextCols=[0, 1, 2], bottomPadding=6, topPadding=6,
                        title="Clinical Incidents", titleStyle=subheadingStyle, headerColor=uniColor, tableTextStyle=tableTextStyleSmall)
    elements.append(ciTable)

    # Self reflections
    reflectionsDf = getSelfReflections(engine, cohort, studentNumber, formsTable)
    if not reflectionsDf.empty:
        r = reflectionsDf.copy()
        r.columns = ["Date", "Self Reflection", "Item Codes"]
        r["Self Reflection"] = r["Self Reflection"].str.replace("\n", "<br/>")
        elements.append(createTable(
            r, colRatio=[1, 6, 1.2], customTextCols=[0, 1, 2], bottomPadding=6, topPadding=6,
            title="Student Judgement: Self Reflections",
            titleStyle=subheadingStyle, headerColor=uniColor,
            tableTextStyle=tableTextStyleSmall,
        ))
        elements.append(Spacer(1, 18))


def plotEntrustmentReadinessTimeSeries(df, title, uniColor=None, useDateAxis=False):
    """
    Time series of entrustment and practice readiness on the same axes with legend.
    useDateAxis=True  → matplotlib date axis with auto-spaced ticks (e.g. '21 Jan')
    useDateAxis=False → string dates on x-axis (default, matches existing rubricPlot style)
    """
    if uniColor is None:
        uniColor = variableUtils.uniColor

    df = df.copy()
    df['date'] = pd.to_datetime(df['date'])
    df.sort_values('date', inplace=True)

    # Aggregate by date — average if multiple forms on same day
    entDf = df.dropna(subset=['entrustment']).groupby('date')['entrustment'].mean().reset_index()
    prDf = df.dropna(subset=['practice_readiness']).groupby('date')['practice_readiness'].mean().reset_index()

    if not useDateAxis:
        entDf['date'] = entDf['date'].dt.strftime('%Y-%m-%d')
        prDf['date'] = prDf['date'].dt.strftime('%Y-%m-%d')

    fig, ax = plt.subplots(figsize=(variableUtils.figSize[0], variableUtils.figSize[1] / 3), dpi = 200)

    offset = 0.05

    if not entDf.empty:
        ax.plot(entDf['date'], entDf['entrustment'] - offset, marker='o', color='steelblue',
                label='Entrustment (Assessor)', linewidth=1.5, markersize=3)
    if not prDf.empty:
        ax.plot(prDf['date'], prDf['practice_readiness'] + offset, marker='s', color='darkorange',
                label='Practice Readiness (Student)', linewidth=1.5, markersize=3)

    ax.set_ylim(0.5, 4.5)
    ax.set_yticks([1, 2, 3, 4])
    ax.set_yticklabels(['1', '2', '3', '4'], fontsize=6, color=uniColor)

    if useDateAxis:
        ax.xaxis.set_major_locator(mdates.WeekdayLocator(interval=1))
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%d %b'))
        ax.tick_params(axis='x', rotation=45, labelsize=6)
    else:
        ax.tick_params(axis='x', rotation=90, labelsize=6)

    ax.set_xlabel('')
    ax.set_ylabel('Level', fontsize=8, color=uniColor)
    ax.set_title(title, fontsize=10, color=uniColor)
    ax.grid(True, axis='y', linestyle='-', alpha=0.3)
    ax.grid(True, axis='x', linestyle=':', alpha=0.2)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.legend(loc='lower left', fontsize=7, framealpha=0.9)
    fig.tight_layout()
    plt.close(fig)

    return fig


def buildStudentPdf(engine, cohort, studentNumber, studentName, outputPath,
                    formsTable="dds4_boh3_forms", pageSize=None,
                    leftMargin=36, rightMargin=36, topMargin=48, bottomMargin=36,
                    styles=None, subheadingStyle=None, subsubheadingStyleL=None,
                    uniColor=None, tableTextStyleSmall=None):
    metricsDf, selfDf, assessorDf = getStudentSummaryTable(
        engine, cohort, studentNumber, studentName, formsTable
    )
    doc = SimpleDocTemplate(str(outputPath), pagesize=pageSize, rightMargin=rightMargin,
                            leftMargin=leftMargin, topMargin=topMargin, bottomMargin=bottomMargin)
    elements = []
    elements.append(Spacer(1, 72))
    elements.append(Paragraph(
        "This is a summary report of your clinical activity so far in 2026. "
        "For detailed information please review your completed forms in the DASH program."
        "<br/> We are working on an interactive live dashboard for future reports.",
        subsubheadingStyleL,
    ))
    elements.append(createTable(
        metricsDf, colRatio=[2, 1], customTextCols=[0, 1], bottomPadding=6, topPadding=6,
        title="Summary", titleStyle=subheadingStyle, headerColor=uniColor,
    ))


    timeSeriesDf = getStudentTimeSeries(engine, cohort, studentNumber, formsTable)
    if not timeSeriesDf.empty:
        fig = plotEntrustmentReadinessTimeSeries(timeSeriesDf, title="Entrustment & Practice Readiness Over Time",
                                                  uniColor=uniColor, useDateAxis=False)
        elements.append(Spacer(1, 24))
        elements.append(addPlotImage(fig, 0.9))

    topItemsDf = getStudentTopItemCodes(engine, cohort, studentNumber, limit = 30, formsTable=formsTable)
    if not topItemsDf.empty:
        elements.append(Spacer(1, 18))
        topItemsDf2 = topItemsDf.copy()
        topItemsDf2.columns = ["Item Code", "Description", "Total Qty"]
        elements.append(createTable(
            topItemsDf2, colRatio=[1, 4, 1], customTextCols=[0, 1, 2], bottomPadding=6, topPadding=6,
            title="Top Procedures", titleStyle=subheadingStyle, headerColor=uniColor,
        ))



    buildStudentVsAssessorSection(
        engine, cohort, studentNumber, elements, styles, formsTable,
        subheadingStyle=subheadingStyle, uniColor=uniColor,
        tableTextStyleSmall=tableTextStyleSmall,
    )

    bannerTitle = f"Student Summary - {cohort}"
    bannerSubtitle = f"{studentName} ({studentNumber})" if studentNumber else studentName
    doc.build(elements, onFirstPage=getBannerDrawer(bannerTitle, bannerSubtitle))


def buildCohortStudentReports(engine, cohort, outputDir, formsTable="dds4_boh3_forms",
                              pageSize=None, leftMargin=36, rightMargin=36,
                              topMargin=48, bottomMargin=36, styles=None,
                              subheadingStyle=None, subsubheadingStyleL=None,
                              uniColor=None, tableTextStyleSmall=None):
    outputDir = Path(outputDir)
    outputDir.mkdir(parents=True, exist_ok=True)

    studentsDf = getStudentsInCohort(engine, cohort, formsTable)
    for _, row in studentsDf.iterrows():
        studentNumber = row["student_number"]
        studentName = row["student_name"]
        # if studentNumber !=1079946:  # temp filter for testing, remove in production
          # continue
        safeName = "".join(c for c in str(studentNumber) if c.isalnum() or c in (" ", "_", "-")).strip()
        outPath = outputDir / f"{safeName}.pdf"
        buildStudentPdf(
            engine=engine, cohort=cohort, studentNumber=studentNumber, studentName=studentName,
            outputPath=outPath, formsTable=formsTable, pageSize=pageSize,
            leftMargin=leftMargin, rightMargin=rightMargin, topMargin=topMargin,
            bottomMargin=bottomMargin, styles=styles, subheadingStyle=subheadingStyle,
            subsubheadingStyleL=subsubheadingStyleL, uniColor=uniColor,
            tableTextStyleSmall=tableTextStyleSmall,
        )
        # break  # for testing, remove in production to generate for all students


def buildEntrustmentTimeSeriesPdf(*, engine, cohort, outPath, bannerTitle,
                                  subheadingStyle, uniColor, formsTable="dds4_boh3_forms",
                                  pageSize=None, rightMargin=36, leftMargin=36,
                                  topMargin=48, bottomMargin=36):
    if pageSize is None:
        pageSize = variableUtils.pageSize

    elements = []
    elements.append(Spacer(1, 72))

    doc = SimpleDocTemplate(str(outPath), pagesize=pageSize, rightMargin=rightMargin,
                            leftMargin=leftMargin, topMargin=topMargin, bottomMargin=bottomMargin)

    studentsDf = getStudentsInCohort(engine, cohort, formsTable)
    studentsDf.sort_values("student_name", inplace=True)
    for i, row in studentsDf.iterrows():
        studentName = row["student_name"]
        studentNumber = row["student_number"]

        # elements.append(Paragraph(f"{studentName} ({studentNumber})", subheadingStyle))
        elements.append(Spacer(1, 12))

        timeSeriesDf = getStudentTimeSeries(engine, cohort, studentNumber, formsTable)
        if not timeSeriesDf.empty:
            fig = plotEntrustmentReadinessTimeSeries(timeSeriesDf, title=f"{studentName} ({studentNumber})", uniColor=uniColor,
                                                      useDateAxis=False)
            img = addPlotImage(fig, 0.9)
            elements.append(img)
        else:
            elements.append(Paragraph("No entrustment/readiness data available.", subheadingStyle))

        elements.append(Spacer(1, 36))

    doc.build(elements, onFirstPage=getBannerDrawer(bannerTitle, ""))
# ═══════════════════════════════════════════════════════════════════════════
# 7. Excel textual report (one sheet per student)
# ═══════════════════════════════════════════════════════════════════════════

_MAIN_SQL = """WITH base AS (
  SELECT student_name, assessor_name, datetimeutc::date AS date, student_data, assessor_data
  FROM dds4_boh3_forms
  WHERE cohort = :cohort AND student_name = :studentName
)
SELECT b.student_name, b.date, b.assessor_name,
  NULLIF(b.student_data->'texts'->>'reflection','') AS student_reflection,
  os.other_notes
FROM base b
LEFT JOIN LATERAL (
  SELECT NULLIF(string_agg(DISTINCT NULLIF(trim(v), ''), ', '), '') AS other_notes
  FROM (
    SELECT e->>'value' AS v
    FROM jsonb_array_elements(COALESCE(b.assessor_data->'multi-select'->'weakness-other','[]'::jsonb)) e
    UNION ALL
    SELECT e->>'value' AS v
    FROM jsonb_array_elements(COALESCE(b.assessor_data->'multi-select'->'strengths','[]'::jsonb)) e
    WHERE COALESCE(e->>'name','') = 'Other'
  ) t
) os ON TRUE
ORDER BY b.date, b.assessor_name;
"""

_WEAKNESS_SQL = """WITH base AS (
  SELECT assessor_data FROM dds4_boh3_forms
  WHERE cohort = :cohort AND student_name = :studentName
)
SELECT trim(e->>'value') AS weakness, COUNT(*)::int AS n
FROM base b
JOIN LATERAL jsonb_each(COALESCE(b.assessor_data->'multi-select','{}'::jsonb)) kv(key, arr) ON TRUE
JOIN LATERAL jsonb_array_elements(COALESCE(kv.arr,'[]'::jsonb)) e ON TRUE
WHERE kv.key LIKE 'weakness-%%' AND kv.key <> 'weakness-other'
  AND NULLIF(trim(e->>'value'), '') IS NOT NULL
GROUP BY weakness ORDER BY n DESC, weakness;
"""

_STRENGTH_SQL = """WITH base AS (
  SELECT assessor_data FROM dds4_boh3_forms
  WHERE cohort = :cohort AND student_name = :studentName
)
SELECT trim(e->>'value') AS strength, COUNT(*)::int AS n
FROM base b
JOIN LATERAL jsonb_array_elements(COALESCE(b.assessor_data->'multi-select'->'strengths','[]'::jsonb)) e ON TRUE
WHERE COALESCE(e->>'name','') <> 'Other' AND NULLIF(trim(e->>'value'), '') IS NOT NULL
GROUP BY strength ORDER BY n DESC, strength;
"""

_CONCERNS_SQL = """WITH base AS (
  SELECT student_name, assessor_name, datetimeutc::date AS date, additional_concerns, assessor_data
  FROM dds4_boh3_forms WHERE cohort = :cohort AND student_name = :studentName
)
SELECT b.student_name, b.date, b.assessor_name,
  NULLIF(b.additional_concerns,'') AS additional_concerns,
  ci.critical_incidents
FROM base b
LEFT JOIN LATERAL (
  SELECT NULLIF(string_agg(DISTINCT NULLIF(trim(e->>'value'), ''), ', '), '') AS critical_incidents
  FROM jsonb_array_elements(COALESCE(b.assessor_data->'multi-select'->'clinical-incident','[]'::jsonb)) e
) ci ON TRUE
WHERE NULLIF(b.additional_concerns,'') IS NOT NULL OR ci.critical_incidents IS NOT NULL
ORDER BY b.date, b.assessor_name;
"""




def buildStudentSheet(engine, wb, cohort, studentName):
    sheetName = _sanitizeSheetName(studentName)
    ws = wb.create_sheet(title=sheetName)

    mainDf = readDf(engine, _MAIN_SQL, {"cohort": cohort, "studentName": studentName})
    weaknessDf = readDf(engine, _WEAKNESS_SQL, {"cohort": cohort, "studentName": studentName})
    strengthDf = readDf(engine, _STRENGTH_SQL, {"cohort": cohort, "studentName": studentName})
    concernsDf = readDf(engine, _CONCERNS_SQL, {"cohort": cohort, "studentName": studentName})

    for df in [mainDf, weaknessDf, strengthDf, concernsDf]:
        if not df.empty:
            df.columns = [c.lower() for c in df.columns]

    if not mainDf.empty:
        keepCols = ["date", "assessor_name", "student_reflection", "other_notes"]
        mainOut = mainDf[keepCols].copy()
    else:
        mainOut = pd.DataFrame(columns=["date", "assessor_name", "student_reflection", "other_notes"])

    row = 1
    row = _writeTitle(ws, row, f"Student Summary: {studentName} ({cohort})")
    row = _writeTable(ws, mainOut, row, title="Entries (Reflection + Other Notes)")
    if weaknessDf.empty:
        weaknessDf = pd.DataFrame(columns=["weakness", "n"])
    row = _writeTable(ws, weaknessDf, row, title="Weaknesses (Counts)")
    if strengthDf.empty:
        strengthDf = pd.DataFrame(columns=["strength", "n"])
    row = _writeTable(ws, strengthDf, row, title="Strengths (Counts)")
    if not concernsDf.empty:
        concernsOut = concernsDf[["date", "assessor_name", "additional_concerns", "critical_incidents"]].copy()
        row = _writeTable(ws, concernsOut, row, title="Additional Concerns & Critical Incidents")
    _autoFitColumns(ws)


def exportStudentTextWorkbook(engine, cohort, outPath):
    studentsSql = """
    SELECT DISTINCT student_name FROM dds4_boh3_forms
    WHERE cohort = :cohort AND student_name IS NOT NULL AND student_name <> 'Test Student'
    ORDER BY student_name;
    """
    studentsDf = readDf(engine, studentsSql, {"cohort": cohort})
    studentNames = studentsDf["student_name"].dropna().tolist()

    wb = Workbook()
    wb.remove(wb.active)
    for studentName in studentNames:
        buildStudentSheet(engine, wb, cohort, studentName)
    wb.save(outPath)


# ═══════════════════════════════════════════════════════════════════════════
# 8. Individual detailed entry report
# ═══════════════════════════════════════════════════════════════════════════

INDIVIDUAL_ENTRY_SQL = """
WITH b AS (
  SELECT * FROM dds4_boh3_forms WHERE assessmentid = :assessmentId
),
student AS (
  SELECT
    b.assessmentid, b.form_code, b.cohort, b.subject, b.type,
    b.createdat::date AS created_date, b.updatedat::date AS updated_date,
    b.clinic, b.rotation,
    b.student_number, b.student_name, b.student_email,
    b.assessorid, b.assessor_name,
    NULLIF(b.student_data->'texts'->>'reflection','') AS student_reflection,
    b.student_data->'scales'->'scale-practice-readiness'->>'scale' AS practice_readiness_code,
    b.student_config->'scales'->'scale-practice-readiness'->'fields'
      -> (b.student_data->'scales'->'scale-practice-readiness'->>'scale') AS practice_readiness_text,
    b.assessor_data->'scales'->'scale-entrustment'->>'scale' AS entrustment_code,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC1' AS mc1,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC2' AS mc2,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC3' AS mc3,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC4' AS mc4,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC5' AS mc5,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC6' AS mc6,
    b.student_data->'checklists'->'checklist-caf-final-eval'->>'MC7' AS mc7,
    CASE b.assessor_data->'scales'->'scale-entrustment'->>'scale'
      WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4 ELSE NULL
    END AS entrustment_num,
    b.assessor_config->'scales'->'scale-entrustment'->'fields'
      -> (b.assessor_data->'scales'->'scale-entrustment'->>'scale') AS entrustment_text,
    NULLIF(b.assessor_data->'texts'->>'additional_comments','') AS assessor_comments,
    NULLIF(b.additional_concerns,'') AS additional_concerns,
    b.patient_data, b.assessor_data
  FROM b
),
agg AS (
  SELECT s.*,
    (SELECT NULLIF(string_agg(DISTINCT trim(e->>'value'), ', '), '')
     FROM jsonb_array_elements(COALESCE(s.assessor_data->'multi-select'->'strengths','[]'::jsonb)) e
     WHERE COALESCE(e->>'name','') <> 'Other' AND NULLIF(trim(e->>'value'), '') IS NOT NULL) AS strengths,
    (SELECT NULLIF(string_agg(DISTINCT trim(e->>'value'), ', '), '')
     FROM jsonb_array_elements(COALESCE(s.assessor_data->'multi-select'->'strengths','[]'::jsonb)) e
     WHERE COALESCE(e->>'name','') = 'Other' AND NULLIF(trim(e->>'value'), '') IS NOT NULL) AS strengths_other,
    (SELECT NULLIF(string_agg(DISTINCT trim(e->>'value'), ', '), '')
     FROM jsonb_each(COALESCE(s.assessor_data->'multi-select','{}'::jsonb)) kv(key, arr)
     JOIN LATERAL jsonb_array_elements(COALESCE(kv.arr,'[]'::jsonb)) e ON TRUE
     WHERE kv.key LIKE 'weakness-%%' AND kv.key <> 'weakness-other'
       AND NULLIF(trim(e->>'value'), '') IS NOT NULL) AS weaknesses,
    (SELECT NULLIF(string_agg(DISTINCT trim(e->>'value'), ', '), '')
     FROM jsonb_array_elements(COALESCE(s.assessor_data->'multi-select'->'weakness-other','[]'::jsonb)) e
     WHERE NULLIF(trim(e->>'value'), '') IS NOT NULL) AS weaknesses_other,
    (SELECT NULLIF(string_agg(DISTINCT trim(e->>'value'), ', '), '')
     FROM jsonb_array_elements(COALESCE(s.assessor_data->'multi-select'->'clinical-incident','[]'::jsonb)) e
     WHERE NULLIF(trim(e->>'value'), '') IS NOT NULL) AS clinical_incidents
  FROM student s
)
SELECT
  assessmentid, form_code, cohort, subject, type, created_date, updated_date, clinic, rotation,
  student_number, student_name, student_email,
  assessorid, assessor_name, practice_readiness_text,
  entrustment_num, entrustment_text,
  student_reflection, assessor_comments,
  strengths, strengths_other, weaknesses, weaknesses_other,
  clinical_incidents, additional_concerns, patient_data
FROM agg;
"""

INDIVIDUAL_MC_SQL = """
SELECT b.assessmentid, b.student_name, kv.key AS "MC Code",
  b.student_config->'checklists'->'checklist-caf-final-eval'->'fields'->> kv.key AS "Full MC Text",
  kv.value AS "MC Rating",
  CASE kv.value
    WHEN 'Done well' THEN 1.0 WHEN 'Done' THEN 0.8 WHEN 'Mostly done' THEN 0.6
    WHEN 'Sometimes done' THEN 0.4 WHEN 'Not done' THEN 0.0 ELSE NULL
  END AS "MC Score"
FROM dds4_boh3_forms b,
LATERAL jsonb_each_text(b.student_data->'checklists'->'checklist-caf-final-eval') kv
WHERE b.assessmentid = :assessmentId;
"""


def buildIndividualEntryPage(elements, row, mcDf,
                             uniColor=None, subheadingStyle=None,
                             tableTextStyleSmall=None):
    """Build reportlab elements for a single assessment entry."""
    infoData = [
        ("Creation Date", row["created_date"]),
        ("Updated Date", row["updated_date"]),
        ("Assessor", row["assessor_name"]),
        ("Clinic", row["clinic"]),
        ("Rotation", row["rotation"]),
    ]
    infoTable = createTable(
        pd.DataFrame(infoData, columns=["Field", "Value"]),
        colRatio=[1, 2], customTextCols=[0, 1], bottomPadding=4, topPadding=4,
        headerColor=uniColor, tableTextStyle=tableTextStyleSmall,
        title="Entry Information", titleStyle=subheadingStyle,
    )
    elements.append(infoTable)

    if not mcDf.empty:
        mcTable = createTable(
            mcDf[["MC Code", "Full MC Text", "MC Rating"]],
            colRatio=[1, 4, 1], customTextCols=[0, 1, 2], bottomPadding=4, topPadding=4,
            headerColor=uniColor, tableTextStyle=tableTextStyleSmall,
            title="CAF Checklist", titleStyle=subheadingStyle,
        )
        elements.append(Spacer(1, 12))
        elements.append(mcTable)

    reflectionsData = [
        ("Student Reflection", row["student_reflection"]),
        ("Assessor Comments", row["assessor_comments"]),
        ("Strengths Noted", row["strengths"]),
        ("Other Strengths", row["strengths_other"]),
        ("Weaknesses Noted", row["weaknesses"]),
        ("Other Weaknesses", row["weaknesses_other"]),
        ("Clinical Incidents", row["clinical_incidents"]),
        ("Additional Concerns", row["additional_concerns"]),
    ]
    reflectionsTable = createTable(
        pd.DataFrame(reflectionsData, columns=["Field", "Content"]),
        colRatio=[1, 3], customTextCols=[0, 1], bottomPadding=4, topPadding=4,
        headerColor=uniColor, tableTextStyle=tableTextStyleSmall,
        title="Reflections & Comments", titleStyle=subheadingStyle,
    )
    elements.append(Spacer(1, 12))
    elements.append(reflectionsTable)
