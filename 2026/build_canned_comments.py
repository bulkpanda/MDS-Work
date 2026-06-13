"""
Build a canned-comment library from the DASH assessment-form export.

Pipeline:
  1. Load the workbook and flatten the nested `forms` JSON (dedupe on each form's id).
  2. Extract student self-reflections and assessor feedback (two separate registers).
  3. Segment each reflection into sentence/line units; normalise (mask tooth numbers
     like "36D" -> <tooth> and measurements like "1.5mm" -> <num>) so the same advice
     about different teeth collapses to one template.
  4. Embed unique normalised units (MiniLM) and cluster into semantic "themes".
     A theme = one canned comment.
  5. For each theme, work out scope from where it is actually used:
        - GLOBAL      : broad, cross-context (generic praise / professionalism / etc.)
        - ITEM_CODE   : usage concentrated in <=3 procedure codes
        - CLINIC_TYPE : usage concentrated in <=2 clinic types (above base rate)
     Every theme is tagged on BOTH dimensions (clinic types + item codes it applies to),
     so the UI can filter by either while always showing the global core.
  6. Export an Excel workbook (README + Assessor + Student + item-code reference).

Requires: pandas, numpy, scikit-learn, sentence-transformers, openpyxl
"""

import json
import re
from collections import Counter

import numpy as np
import pandas as pd
from sklearn.cluster import MiniBatchKMeans

# ----------------------------------------------------------------------------- config
INPUT_XLSX   = "dash_forms_2025_anonymized.xlsx"
OUTPUT_XLSX  = "canned_comment_library.xlsx"
SEED         = 0
K_THEMES     = {"assessor": 500, "student": 600}   # cluster granularity per register
MIN_USAGE    = 4        # a theme must be used >= this many times to be worth canning
COVER        = 0.80     # tag the codes/clinic-types covering this share of a theme's use
CT_LIFT_MIN  = 1.30     # clinic-type must be over-represented vs base rate to be "specific"
EMB_MODEL    = "all-MiniLM-L6-v2"

# phrases that are placeholders, not feedback -> excluded from the library
PLACEHOLDERS = {"see above", "as above", "see below", "as below", "n a", "na", "nil",
                "none", "ditto", "good", "ok", "okay", "done", "yes", "no"}

# ------------------------------------------------------------------- text segmentation
TOOTH  = re.compile(r"\b\d{2}[A-Za-z]{0,3}\b")
NUM    = re.compile(r"\b\d+(\.\d+)?\s?(mm|cm)?\b")
WS     = re.compile(r"\s+")
BULLET = re.compile(r"^\s*[-*•·>\d]+[\).\s]*")


def split_units(text):
    parts = re.split(r"[\n\r]+|(?<=[.;!?])\s+", text)
    out = []
    for p in parts:
        p = BULLET.sub("", p).strip()
        if len(p) >= 3:
            out.append(p)
    return out


def normalise(unit):
    u = unit.lower().strip()
    u = TOOTH.sub("<tooth>", u)
    u = NUM.sub("<num>", u)
    u = re.sub(r"[^\w<>\s]", " ", u)
    return WS.sub(" ", u).strip()


def readable_template(norm_text):
    """Turn a normalised template into something human-readable for review."""
    t = norm_text.replace("<tooth>", "{tooth}").replace("<num>", "{measure}")
    return t.strip().capitalize()


# ----------------------------------------------------------------------- load + flatten
def load_units(path):
    df = pd.read_excel(path, dtype={"student_number": str})
    seen = {}
    for _, r in df.iterrows():
        for f in json.loads(r["forms"]):
            seen[f["id"]] = (f, r)          # dedupe on form id
    code_names = {}
    code_form_ct = Counter()        # distinct forms containing each code
    code_clinic = {}                # code -> Counter of clinic types (form level)
    recs = []
    for f, r in seen.values():
        ad, sd = f.get("assessor_data", {}), f.get("student_data", {})
        codes = tuple(sorted(f.get("checklists", {}).keys()))
        for code, obj in f.get("checklists", {}).items():
            code_names.setdefault(code, obj.get("name", ""))
            code_form_ct[code] += 1
            code_clinic.setdefault(code, Counter())[f.get("clinic_type")] += 1
        base = dict(cohort=r["cohort"], clinic_type=f.get("clinic_type"), codes=codes)
        for who, text in [("student", sd.get("reflection")),
                          ("assessor", ad.get("reflection"))]:
            text = (text or "").strip()
            if not text:
                continue
            for unit in split_units(text):
                norm = normalise(unit)
                if len(norm) >= 3:
                    recs.append(dict(base, who=who, raw=unit, norm=norm))
    meta = dict(code_names=code_names, code_form_ct=code_form_ct, code_clinic=code_clinic)
    return pd.DataFrame(recs), meta


# ----------------------------------------------------------------------------- cluster
def cluster_register(units, who, k):
    import os
    sub = units[units.who == who]
    counts = sub["norm"].value_counts()
    texts = counts.index.tolist()
    cache = f"emb_cache_{who}.npy"
    if os.path.exists(cache):
        emb = np.load(cache)
    else:
        from sentence_transformers import SentenceTransformer
        model = SentenceTransformer(EMB_MODEL)
        emb = model.encode(texts, batch_size=256, normalize_embeddings=True,
                           show_progress_bar=False).astype("float32")
        np.save(cache, emb)
    km = MiniBatchKMeans(n_clusters=min(k, len(texts)), random_state=SEED,
                         n_init=3, batch_size=2048).fit(emb)
    norm2theme = dict(zip(texts, km.labels_))
    sub = sub.copy()
    sub["theme"] = sub["norm"].map(norm2theme)
    return sub


# -------------------------------------------------------------------- scope assignment
def cover_set(value_counts, frac):
    s = value_counts.sort_values(ascending=False)
    c = s.cumsum() / s.sum()
    return list(s.index[: (c < frac).sum() + 1])


ALPHA = re.compile(r"[A-Za-z]")


def is_quality(raw):
    """Reject normalisation fragments / junk as canned text."""
    r = raw.strip()
    if not r or not r[0].isalpha():
        return False
    words = r.split()
    if not (2 <= len(words) <= 16):
        return False
    alpha_words = sum(bool(ALPHA.search(w)) for w in words)
    return alpha_words >= 2


def pick_representative(theme_rows):
    """Concise, human-readable canned text for a theme, plus a {tooth}/{measure}
    template derived from the SAME chosen phrase (so they never disagree)."""
    raw_counts = theme_rows["raw"].value_counts()
    ranked = raw_counts.index.tolist()
    quality = [t for t in ranked if is_quality(t)]
    if not quality:
        return None, "", []
    # prefer concise (2-12 words) among the most frequent quality phrases
    concise = [t for t in quality[:8] if 2 <= len(t.split()) <= 12]
    rep_raw = (concise[0] if concise else quality[0]).strip()
    masked = NUM.sub("<num>", TOOTH.sub("<tooth>", rep_raw))
    template = readable_template(masked) if ("<tooth>" in masked or "<num>" in masked) else ""
    variants = [t for t in quality if t != rep_raw][:3]
    return rep_raw, template, variants


GENERIC_CODE_SPREAD = 12   # appears in >= this many distinct codes -> treat as generic/global
GENERIC_CT_SPREAD   = 5    # or appears in >= this many clinic types


def build_library(sub, code_names, base_ct, base_code):
    # ---- 1. one record per theme ----
    raw_themes = []
    for theme, g in sub.groupby("theme"):
        usage = len(g)
        if usage < MIN_USAGE:
            continue
        rep, template, variants = pick_representative(g)
        if rep is None or normalise(rep) in PLACEHOLDERS:
            continue
        ct_counts = g["clinic_type"].value_counts()
        code_counts = g.explode("codes")["codes"].dropna().value_counts()
        raw_themes.append(dict(key=normalise(rep), rep=rep, template=template,
                               variants=variants, usage=usage,
                               ct_counts=ct_counts, code_counts=code_counts))

    # ---- 2. merge themes that resolved to the same phrase ----
    merged = {}
    for t in raw_themes:
        m = merged.get(t["key"])
        if m is None:
            merged[t["key"]] = t
        else:
            m["usage"] += t["usage"]
            m["ct_counts"] = m["ct_counts"].add(t["ct_counts"], fill_value=0)
            m["code_counts"] = m["code_counts"].add(t["code_counts"], fill_value=0)
            for v in t["variants"]:
                if v not in m["variants"] and len(m["variants"]) < 3:
                    m["variants"].append(v)

    # ---- 3. scope + tags from the (merged) footprint ----
    rows = []
    for t in merged.values():
        ctc, codec = t["ct_counts"], t["code_counts"]
        code_cov = cover_set(codec, COVER) if len(codec) else []
        ct_cov = cover_set(ctc, COVER) if len(ctc) else []
        n_codes = int((codec > 0).sum())
        n_cts = int((ctc > 0).sum())
        generic = (n_codes >= GENERIC_CODE_SPREAD) or (n_cts >= GENERIC_CT_SPREAD)

        if generic:
            scope = "GLOBAL"
        elif 0 < len(code_cov) <= 3:
            scope = "ITEM_CODE"
        elif len(ct_cov) <= 2:
            top_ct = ctc.index[0]
            lift = (ctc.iloc[0] / ctc.sum()) / max(base_ct.get(top_ct, 1e-9), 1e-9)
            scope = "CLINIC_TYPE" if lift >= CT_LIFT_MIN else "GLOBAL"
        else:
            scope = "GLOBAL"

        rows.append(dict(
            scope=scope,
            canned_text=t["rep"],
            template_if_procedure_specific=t["template"],
            applies_to_clinic_types=("ALL" if scope == "GLOBAL" else ", ".join(ct_cov)),
            applies_to_item_codes=("ALL" if scope != "ITEM_CODE" else ", ".join(code_cov)),
            item_code_names=("" if scope != "ITEM_CODE"
                             else " | ".join(f"{c}: {code_names.get(c,'')}" for c in code_cov)),
            frequency=int(t["usage"]),
            example_variants=" || ".join(t["variants"]),
        ))

    lib = pd.DataFrame(rows)
    order = {"GLOBAL": 0, "ITEM_CODE": 1, "CLINIC_TYPE": 2}
    lib = lib.sort_values(["scope", "frequency"],
                          key=lambda s: s.map(order) if s.name == "scope" else s,
                          ascending=[True, False]).reset_index(drop=True)
    lib.insert(0, "comment_id", range(1, len(lib) + 1))
    return lib


# ------------------------------------------------------------------------------- export
def write_excel(libs, code_ref, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HEAD = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    BODY = Font(name="Arial", size=10)
    FILL = PatternFill("solid", fgColor="2F5496")
    WRAP = Alignment(wrap_text=True, vertical="top")
    TOPL = Alignment(vertical="top")
    THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

    wb = Workbook()

    # ---- README ----
    ws = wb.active
    ws.title = "README"
    readme = [
        ("Canned Comment Library", ""),
        ("", ""),
        ("Built from", "dash_forms_2025_anonymized.xlsx (de-identified DASH assessment forms)"),
        ("Two registers", "Student self-reflections and Assessor feedback are kept separate."),
        ("", ""),
        ("How comments are grouped", ""),
        ("  GLOBAL", "Generic, cross-context (praise, professionalism, time management, "
                     "reflection quality). Show in EVERY context regardless of filter."),
        ("  ITEM_CODE", "Usage concentrated in <=3 procedure codes. Show when those codes are on the form."),
        ("  CLINIC_TYPE", "Usage concentrated in <=2 clinic types (above base rate). "
                          "Coarser grouping; useful when a form has many codes."),
        ("", ""),
        ("Tagging", "Every comment carries BOTH an applies_to_clinic_types and an "
                    "applies_to_item_codes tag, so you can filter the menu by either dimension. "
                    "'ALL' means the comment is not restricted on that dimension."),
        ("template_if_procedure_specific", "Where the same advice recurs about different teeth, "
                    "a {tooth}/{measure} template is suggested so one comment covers all teeth."),
        ("frequency", "How many comment fragments of this kind appeared in 2025 (theme "
                      "size) - the prioritisation/reuse signal. Higher = write the comment first."),
        ("", ""),
        ("Method note", "Phrases were segmented to sentence level, normalised (tooth numbers and "
                        "measurements masked), embedded (all-MiniLM-L6-v2) and clustered into themes; "
                        "one representative phrase per theme. Item code is the strongest content "
                        "driver; clinic type is a milder proxy; cohort and the rating scales carry "
                        "essentially no signal and are NOT used to group."),
        ("Intended use", "A starter set for human review - edit/merge/delete freely. A free-text "
                         "box should remain for the bespoke long tail."),
    ]
    for i, (a, b) in enumerate(readme, 1):
        ca, cb = ws.cell(i, 1, a), ws.cell(i, 2, b)
        ca.font = Font(name="Arial", bold=(i == 1 or a.strip().endswith(":") or a.isupper()
                                           or a in ("Built from", "Two registers", "How comments are grouped",
                                                    "Tagging", "Method note", "Intended use")),
                       size=(14 if i == 1 else 10))
        cb.font = BODY
        cb.alignment = WRAP
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 95

    # ---- library sheets ----
    def add_sheet(name, df, wraps):
        ws = wb.create_sheet(name)
        cols = list(df.columns)
        for j, col in enumerate(cols, 1):
            c = ws.cell(1, j, col)
            c.font, c.fill, c.alignment = HEAD, FILL, Alignment(vertical="center")
        for i, (_, row) in enumerate(df.iterrows(), 2):
            for j, col in enumerate(cols, 1):
                c = ws.cell(i, j, row[col])
                c.font, c.border = BODY, THIN
                c.alignment = WRAP if col in wraps else TOPL
        widths = {"comment_id": 10, "scope": 13, "canned_text": 48,
                  "template_if_procedure_specific": 38, "applies_to_clinic_types": 22,
                  "applies_to_item_codes": 22, "item_code_names": 50,
                  "frequency": 11, "example_variants": 55,
                  "item_code": 12, "name": 55, "clinic_type": 18, "n_forms": 10}
        for j, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 18)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df)+1}"

    wrap_cols = {"canned_text", "template_if_procedure_specific", "item_code_names",
                 "example_variants", "applies_to_clinic_types", "applies_to_item_codes", "name"}
    add_sheet("Assessor_Library", libs["assessor"], wrap_cols)
    add_sheet("Student_Library", libs["student"], wrap_cols)
    add_sheet("ItemCode_Reference", code_ref, wrap_cols)

    wb.save(path)


# ---------------------------------------------------------------------------------- main
def main():
    units, meta = load_units(INPUT_XLSX)
    code_names = meta["code_names"]

    base_ct = units["clinic_type"].value_counts(normalize=True).to_dict()
    base_code = units.explode("codes")["codes"].value_counts(normalize=True).to_dict()

    libs = {}
    for who in ("assessor", "student"):
        sub = cluster_register(units, who, K_THEMES[who])
        libs[who] = build_library(sub, code_names, base_ct, base_code)
        n = libs[who]
        print(f"{who:8} library: {len(n)} canned comments  "
              f"(GLOBAL {sum(n.scope=='GLOBAL')}, "
              f"ITEM_CODE {sum(n.scope=='ITEM_CODE')}, "
              f"CLINIC_TYPE {sum(n.scope=='CLINIC_TYPE')})")

    # item-code reference sheet (true distinct-form counts)
    ref = pd.DataFrame([
        dict(item_code=c, name=code_names.get(c, ""),
             clinic_type=meta["code_clinic"][c].most_common(1)[0][0],
             n_forms=meta["code_form_ct"][c])
        for c in meta["code_form_ct"]
    ]).sort_values("n_forms", ascending=False).reset_index(drop=True)

    write_excel(libs, ref, OUTPUT_XLSX)
    print(f"\nWrote {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
