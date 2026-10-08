"""
mcq_analysis_utils.py — Canvas "Student Analysis Report" MCQ exports -> coordinator workbook.

Handles fixed papers (everyone sees the same items), random-pool papers (each student gets
k of n items, unseen items are blank) and single- or multi-session exams (e.g. AM / PM).

Workbook (layout agreed with Kunal on the DENT90148 Fixed Pros analysis, Oct 2026):
  Summary Dashboard      cohort table, session comparison, difficulty tiers, score distribution, key findings
  <Session> Question Analysis   per item: key, counts, % correct, SD, D, item-rest r_pb, flag, distractors, note
  Student Scores         score, %, items wrong, Rasch θ (EAP) + SE — lowest first
  Performance Chart      % correct by question, one bar chart per session
  IRT Item Calibration   Rasch b, SE, infit/outfit, difficulty band (all sessions on one scale)

Main entry point: buildMcqAnalysisWorkbook(sessions, outPath, examTitle, ...).
Only the names in __all__ are exported, so `from mcq_analysis_utils import *` in main.ipynb
cannot shadow reportlab's Table / Paragraph / Image.
"""
import glob, math, re
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as colL

__all__ = ['loadCanvasQuizReport', 'itemStatistics', 'sessionStatistics', 'matchItemsAcrossSessions', 'fitRasch',
           'raschItemFit', 'kr20', 'difficultyFlag', 'buildMcqAnalysisWorkbook']

META_COLS = ['name', 'id', 'sis_id', 'section', 'section_id', 'section_sis_id', 'submitted', 'attempt']
TAIL_COLS = ['n correct', 'n incorrect', 'score']
NAVY, BLUE, HDR, UOM_BLUE, UOM_ORANGE = '1F3864', '2E75B6', '2E4053', '094183', 'E07B00'
TIER_STYLE = {'High': ('D5F5E3', '1E8449', 'High ≥90%'), 'Good': ('D6E4F0', '2E75B6', 'Good 75–89%'),
              'Moderate': ('FEF9E7', 'B7950B', 'Moderate 60–74%'), 'Low': ('FADBD8', 'C0392B', 'Low <60%')}
THIN = Border(*(Side(style='thin', color='D0D3D4'),) * 4)


# ───────────────────────────── loading / scoring ─────────────────────────────
def cleanStem(header):
    """'3099907: \\nQuestion.\\nIn a ...' -> ('3099907', 'In a ...')."""
    qid, txt = header.split(':', 1)
    txt = ' '.join(txt.split())
    txt = re.sub(r'^(Question\s*[.,:]?\s*)', '', txt, flags=re.I).lstrip(' ,.:;')
    return qid.strip(), txt


def loadCanvasQuizReport(csvPath, label='Exam', scoreOverrides=None, rescoreToKey=True):
    """Read one Canvas Student Analysis Report CSV.
    csvPath        : path or glob pattern (first match used)
    label          : session label shown in the workbook ('AM', 'PM', 'Exam' ...)
    scoreOverrides : {(studentNumber, 'Q19'): 0, ...} manual item marks applied last
    rescoreToKey   : re-mark every response against the item key (most common credited answer);
                     catches manual Canvas overrides, which are logged in ['corrections']
    Returns dict: label, df, resp (answers), sc (0/1, NaN = not shown), meta [(qid, stem)], keys,
                  corrections [(name, studentNo, Q, answer, from, to, reason)], isPool, maxScore."""
    path = sorted(glob.glob(csvPath))[0] if any(ch in csvPath for ch in '*?[') else csvPath
    df = pd.read_csv(path)
    if list(df.columns[:8]) != META_COLS or list(df.columns[-3:]) != TAIL_COLS:
        raise ValueError(f'Not a Canvas Student Analysis Report layout: {path}')
    qCols, sCols = list(df.columns[8:-3:2]), list(df.columns[9:-3:2])
    labels = [f'Q{i + 1}' for i in range(len(qCols))]
    resp = df[qCols].copy(); resp.columns = labels
    sc = df[sCols].apply(pd.to_numeric, errors='coerce'); sc.columns = labels
    sc = sc.where(resp.notna())                                   # unseen items stay NaN
    keys = {q: (resp.loc[sc[q] == 1, q].mode().iloc[0] if (sc[q] == 1).any() else '') for q in labels}
    corrections = []
    if rescoreToKey:
        scKey = pd.DataFrame({q: np.where(resp[q].isna(), np.nan, (resp[q] == keys[q]).astype(float)) for q in labels})
        diff = (sc != scKey) & resp.notna()
        for q in labels:
            for i in df.index[diff[q]]:
                corrections.append((df.at[i, 'name'], int(df.at[i, 'sis_id']), q, resp.at[i, q], sc.at[i, q], scKey.at[i, q],
                                    'Canvas mark differed from item key (manual override)'))
        sc = scKey
    for (sid, q), v in (scoreOverrides or {}).items():
        rows = df.index[df.sis_id.astype(str) == str(sid)]
        for i in rows:
            corrections.append((df.at[i, 'name'], int(sid), q, resp.at[i, q], sc.at[i, q], float(v), 'Manual override (scoreOverrides)'))
            sc.at[i, q] = float(v)
    nShown = resp.notna().sum(1)
    return dict(label=label, path=path, df=df, resp=resp, sc=sc, meta=[cleanStem(c) for c in qCols], keys=keys,
                corrections=corrections, isPool=bool(resp.isna().any().any()), nItems=len(labels),
                maxScore=int(nShown.mode().iloc[0]), nShown=nShown)


# ───────────────────────────── classical statistics ─────────────────────────────
def difficultyFlag(p):
    """% correct -> tier key (High / Good / Moderate / Low)."""
    return 'High' if p >= .9 else 'Good' if p >= .75 else 'Moderate' if p >= .6 else 'Low'


def kr20(sc):
    """KR-20 for a fixed (complete) 0/1 matrix; None if items are missing by design."""
    if sc.isna().any().any() or sc.shape[1] < 2: return None
    p = sc.mean(); var = sc.sum(1).var(ddof=1)
    return float(sc.shape[1] / (sc.shape[1] - 1) * (1 - (p * (1 - p)).sum() / var)) if var > 0 else None


def itemStatistics(sess):
    """Per-item table, using only students shown the item.
    D    = % correct in top 27% minus bottom 27%, students ranked by total % (classic upper-lower index)
    r_pb = correlation of item (0/1) with proportion correct on the student's OTHER items (item-rest)"""
    sc, resp = sess['sc'], sess['resp']
    tot, nSh = sc.sum(1), sc.notna().sum(1)
    rows = []
    for k, q in enumerate(sc.columns):
        m = sc[q].notna(); x = sc.loc[m, q]; n = int(m.sum())
        rest = ((tot - sc[q].fillna(0)) / (nSh - 1).clip(lower=1))[m]
        g = max(1, round(.27 * n)); order = (tot / nSh)[m].sort_values(ascending=False).index   # classic D: rank by total %
        dDisc = x[order[:g]].mean() - x[order[-g:]].mean() if n >= 4 else np.nan
        rpb = float(np.corrcoef(x, rest)[0, 1]) if x.std() > 0 and rest.std() > 0 else np.nan
        wrong = resp.loc[m & (sc[q] == 0), q].value_counts()
        rows.append(dict(q=q, qid=sess['meta'][k][0], stem=sess['meta'][k][1], key=sess['keys'][q], nShown=n,
                         nCorrect=int(x.sum()), nIncorrect=int(n - x.sum()), pCorrect=float(x.mean()) if n else np.nan,
                         sd=float(math.sqrt(x.mean() * (1 - x.mean()))) if n else np.nan, discD=dDisc, rPbis=rpb,
                         wrongAnswers='; '.join(f'{a} ({c})' for a, c in wrong.items()),
                         topWrong=(wrong.index[0], int(wrong.iloc[0])) if len(wrong) else None))
    return pd.DataFrame(rows)


def sessionStatistics(sess):
    """Score summary: totals out of maxScore, % (score / items shown), KR-20 for fixed papers."""
    sc = sess['sc']; tot = sc.sum(1); pct = tot / sc.notna().sum(1)
    return dict(n=len(sc), mean=tot.mean(), median=tot.median(), sd=tot.std(), min=tot.min(), max=tot.max(),
                pctMean=pct.mean(), pct=pct, tot=tot, kr20=kr20(sc))


def fmtP(p):
    """p-value text: '< 0.001' or '≈ 0.04'."""
    return '< 0.001' if p < .001 else f'≈ {p:.3f}' if p < .01 else f'≈ {p:.2f}'


def welchTest(a, b):
    """Welch t on two samples; p from the normal approximation (fine for n > ~30 per group)."""
    se = math.sqrt(a.var() / len(a) + b.var() / len(b)); t = (a.mean() - b.mean()) / se
    sp = math.sqrt(((len(a) - 1) * a.var() + (len(b) - 1) * b.var()) / (len(a) + len(b) - 2))
    return t, math.erfc(abs(t) / math.sqrt(2)), (a.mean() - b.mean()) / sp


def matchItemsAcrossSessions(sessList):
    """Same question used in several sessions (Canvas gives copies new IDs): match on normalised
    stem + identical key. Returns {(label, Q): linkId} where linkId is shared by matched copies,
    plus {linkId: 'AM Q7 = PM Q11'} display names."""
    norm = lambda t: re.sub(r'[^a-z0-9]', '', str(t).lower())
    groups = {}
    for s in sessList:
        for (qid, stem), q in zip(s['meta'], s['sc'].columns):
            groups.setdefault((norm(stem), norm(s['keys'][q])), []).append((s['label'], q))
    linkOf, display = {}, {}
    for members in groups.values():
        shared = len({lab for lab, _ in members}) > 1
        for lab, q in members:
            lid = ('LINK|' + '/'.join(f'{l} {x}' for l, x in members)) if shared else f'{lab}|{q}'
            linkOf[(lab, q)] = lid
            display[lid] = ' = '.join(f'{l} {x}' for l, x in members) if shared else q
    return linkOf, display


def signTestP(diffs):
    """Two-sided sign test on paired differences (zeros dropped)."""
    d = [x for x in diffs if x != 0]; n = len(d); k = max(sum(x > 0 for x in d), sum(x < 0 for x in d))
    return min(1.0, 2 * sum(math.comb(n, i) for i in range(k, n + 1)) / 2 ** n) if n else 1.0


# ───────────────────────────── Rasch (1PL) MML-EM ─────────────────────────────
def fitRasch(X, nQuad=61, maxIter=500, tol=1e-6):
    """Rasch via marginal ML (EM, quadrature). X: persons x items, 0/1/NaN (NaN = not administered).
    Latent mean fixed at 0, latent SD estimated. Returns b, seB, thetaEap, thetaPsd, latentSd, nIter."""
    nodes = np.linspace(-8, 8, nQuad); sig = 1.0
    prior = lambda s: (lambda w: w / w.sum())(np.exp(-nodes ** 2 / (2 * s ** 2)))
    obs = ~np.isnan(X); Xf = np.nan_to_num(X)
    p0 = np.clip(np.nanmean(X, 0), .01, .99); b = -np.log(p0 / (1 - p0)); w = prior(sig)
    for it in range(maxIter):
        P = 1 / (1 + np.exp(-(nodes[:, None] - b[None, :])))
        ll = Xf @ np.log(P).T + (obs * (1 - Xf)) @ np.log(1 - P).T
        post = np.exp(ll - ll.max(1, keepdims=True)) * w; post /= post.sum(1, keepdims=True)
        nJk, rJk = post.T @ obs, post.T @ Xf
        bNew = b - (rJk - nJk * P).sum(0) / (nJk * P * (1 - P)).sum(0)
        sig = math.sqrt((post @ nodes ** 2).mean()); w = prior(sig)
        done = np.max(abs(bNew - b)) < tol; b = bNew
        if done: break
    P = 1 / (1 + np.exp(-(nodes[:, None] - b[None, :])))
    theta = post @ nodes
    return b, 1 / np.sqrt((nJk * P * (1 - P)).sum(0)), theta, np.sqrt(np.clip(post @ nodes ** 2 - theta ** 2, 0, None)), sig, it


def raschItemFit(X, b, theta):
    """Infit / outfit mean-squares (0.7–1.3 = acceptable)."""
    P = 1 / (1 + np.exp(-(theta[:, None] - b[None, :]))); W = P * (1 - P); R2 = (X - P) ** 2
    return np.nansum(R2, 0) / np.nansum(np.where(np.isnan(X), np.nan, W), 0), np.nanmean(R2 / W, 0)


def runRasch(sessList, linkOf):
    """Calibrate all sessions on one scale. Questions shared between sessions (linkOf) anchor the
    scale; with no shared questions the sessions are linked only by assuming equivalent groups."""
    X = pd.concat([s['sc'].rename(columns={q: linkOf[(s['label'], q)] for q in s['sc'].columns}) for s in sessList], axis=0)
    p = X.mean()
    calib = [c for c in X.columns if 0 < p[c] < 1 and X[c].notna().sum() >= 5]
    Xm = X[calib].to_numpy(float)
    b, se, th, psd, sig, nIt = fitRasch(Xm)
    infit, outfit = raschItemFit(Xm, b, th)
    items = pd.DataFrame(dict(item=calib, b=b, se=se, infit=infit, outfit=outfit, p=p[calib].values))
    rel = 1 - np.mean(psd ** 2) / (np.var(th) + np.mean(psd ** 2))
    persons = {}
    i = 0
    for s in sessList:
        for sid in s['df'].sis_id:
            persons[(s['label'], int(sid))] = (th[i], psd[i]); i += 1
    notEst = [c for c in X.columns if c not in calib]
    return dict(items=items, persons=persons, reliability=rel, latentSd=sig, notEstimable=notEst, pAll=p)


# ───────────────────────────── workbook helpers ─────────────────────────────
def _fill(c): return PatternFill('solid', fgColor=c)


def _put(ws, ref, v, bg=None, fg='1A1A2E', b=False, sz=11, wrap=False, h='left', fmt=None, border=False):
    c = ws[ref]; c.value = v; c.font = Font(color=fg, bold=b, size=sz)
    if bg: c.fill = _fill(bg)
    c.alignment = Alignment(wrap_text=wrap, vertical='center', horizontal=h)
    if fmt: c.number_format = fmt
    if border: c.border = THIN
    return c


def _banner(ws, title, sub, width):
    ws.merge_cells(f'A1:{colL(width)}1'); _put(ws, 'A1', title, NAVY, 'FFFFFF', True, 14); ws.row_dimensions[1].height = 26
    if sub: ws.merge_cells(f'A2:{colL(width)}2'); _put(ws, 'A2', sub, BLUE, 'FFFFFF')


def _section(ws, r, text, width=8, bg=BLUE):
    _put(ws, f'A{r}', text, bg, 'FFFFFF', True); ws.merge_cells(f'A{r}:{colL(width)}{r}')


def _reviewNote(r):
    notes = []
    if r.pCorrect == 1: notes.append('All correct – no discrimination')
    if r.pCorrect < .6: notes.append('Hardest item – check key & teaching coverage')
    elif r.pCorrect < .75: notes.append('Moderately hard')
    if pd.notna(r.rPbis) and r.rPbis < 0: notes.append('Negative r_pb – review wording/key' if r.nIncorrect >= 3 else 'Negative r_pb (≤2 wrong – likely noise)')
    elif pd.notna(r.rPbis) and r.rPbis >= .2 and r.pCorrect < 1: notes.append('Discriminates well')
    return '; '.join(notes)


def autoFindings(sessList, stats, items, passMark, linkOf=None):
    """Plain-language key findings computed from the results."""
    multi = len(sessList) > 1; out = []
    pctAll = pd.concat([stats[s['label']]['pct'] for s in sessList]); mx = sessList[0]['maxScore']
    totAll = pd.concat([stats[s['label']]['tot'] for s in sessList])
    nFail = int((pctAll < passMark).sum())
    out.append(f'Cohort (N = {len(pctAll)}): mean {totAll.mean():.1f}/{mx} ({pctAll.mean():.1%}), median {totAll.median():.0f}/{mx}, '
               f'range {totAll.min():.0f}–{totAll.max():.0f}. ' + (f'{nFail} student(s) below {passMark:.0%}.' if nFail else f'No student below {passMark:.0%}.'))
    if multi:
        a, b = sessList[0]['label'], sessList[1]['label']
        t, p, d = welchTest(stats[a]['pct'], stats[b]['pct'])
        verdict = (f'{a} scored significantly {"higher" if t > 0 else "lower"} than {b} – check exam conditions / paper difficulty before comparing students across sessions.'
                   if p < .05 else 'no meaningful difference between sessions.')
        out.append(f'{a} vs {b}: {stats[a]["pctMean"]:.1%} vs {stats[b]["pctMean"]:.1%} (Welch p {fmtP(p)}, d = {d:.2f}) – {verdict}')
        if linkOf:
            pa = {linkOf[(a, r.q)]: r.pCorrect for r in items[a].itertuples()}; pb = {linkOf[(b, r.q)]: r.pCorrect for r in items[b].itertuples()}
            common = [k for k in pa if k in pb and k.startswith('LINK|')]
            if common:
                diffs = [pa[k] - pb[k] for k in common]; nLow = sum(x < 0 for x in diffs)
                same = f'All {len(common)}' if len(common) == len(pa) == len(pb) else f'{len(common)}'
                out.append(f'{same} questions are identical in both sessions (same stem and key; Canvas IDs differ). On these same questions {a} averaged '
                           f'{np.mean([pa[k] for k in common]):.1%} vs {b} {np.mean([pb[k] for k in common]):.1%} correct, {a} lower on {nLow} of {len(common)} '
                           f'(sign test p {fmtP(signTestP(diffs))}). ' + ('The gap is therefore a session effect, not harder questions.' if signTestP(diffs) < .05 else 'No consistent session effect.'))
            else:
                out.append(f'{a} and {b} papers share no questions (different stems) – they are parallel forms.')
    allIt = pd.concat([items[s['label']].assign(sess=s['label']) for s in sessList])
    tag = (lambda r: f'{r.sess} {r.q}') if multi else (lambda r: r.q)
    hard = allIt[allIt.pCorrect < .6].sort_values('pCorrect'); mod = allIt[(allIt.pCorrect >= .6) & (allIt.pCorrect < .75)].sort_values('pCorrect')
    out.append('Hardest items (<60%): ' + (', '.join(f'{tag(r)} ({r.pCorrect:.0%})' for r in hard.itertuples()) or 'none') +
               '. Moderate (60–74%): ' + (', '.join(f'{tag(r)} ({r.pCorrect:.0%})' for r in mod.itertuples()) or 'none') + '.')
    tw = [f'{tag(r)} – {r.topWrong[1]}/{r.nShown} chose "{r.topWrong[0][:70]}"' for r in hard.head(5).itertuples() if r.topWrong]
    if tw: out.append('Most common wrong answer on the hardest items: ' + '; '.join(tw) + '.')
    negs = allIt[(allIt.rPbis < 0) & (allIt.nIncorrect >= 3)]
    if len(negs): out.append('Review wording/key (stronger students got these wrong more often): ' + ', '.join(tag(r) for r in negs.itertuples()) + '.')
    pools = [s for s in sessList if s['isPool']]
    if pools:
        who = ' and '.join(s['label'] for s in pools) + ': ' if multi else ''
        out.append(f'{who}each student received {pools[0]["maxScore"]} of {pools[0]["nItems"]} questions drawn at random – item statistics use only students '
                   'who saw the item ("Shown to"); IRT θ (Student Scores) puts students who saw different questions on one scale.')
    for s in sessList:
        for c in s['corrections']:
            out.append(f'Scoring correction ({s["label"]}): {c[0]} ({c[1]}) {c[2]} "{str(c[3])[:60]}" {c[4]:.0f} → {c[5]:.0f} – {c[6]}.')
    return out


# ───────────────────────────── main builder ─────────────────────────────
def buildMcqAnalysisWorkbook(sessions, outPath, examTitle, examSubtitle='', passMark=0.5,
                             scoreOverrides=None, extraFindings=None, rescoreToKey=True, runIrt=True):
    """Build the coordinator MCQ workbook.
    sessions       : {'AM': 'path/or/glob.csv', 'PM': ...}  (one entry for a single-session exam)
    outPath        : .xlsx to write
    examTitle      : banner title, e.g. 'DENT90148 — Endodontics MCQ (DDS2)'
    examSubtitle   : banner line 2 (date, paper description); N is appended automatically
    passMark       : fraction used for the 'below pass' count in findings
    scoreOverrides : {'AM': {(studentNo, 'Q19'): 0}} manual item marks per session
    extraFindings  : list of extra bullet strings (e.g. Config.xlsx notes)
    Returns dict(sessions, stats, items, irt, findings, outPath)."""
    sessList = [loadCanvasQuizReport(p, lab, (scoreOverrides or {}).get(lab), rescoreToKey) for lab, p in sessions.items()]
    stats = {s['label']: sessionStatistics(s) for s in sessList}
    items = {s['label']: itemStatistics(s) for s in sessList}
    linkOf, linkName = matchItemsAcrossSessions(sessList)
    irt = runRasch(sessList, linkOf) if runIrt else None
    findings = autoFindings(sessList, stats, items, passMark, linkOf) + list(extraFindings or [])
    pByLink = {(s['label'], linkOf[(s['label'], r.q)]): (r.q, r.pCorrect) for s in sessList for r in items[s['label']].itertuples()}
    multi, mx = len(sessList) > 1, sessList[0]['maxScore']
    nTotal = sum(st['n'] for st in stats.values())

    wb = Workbook(); ws = wb.active; ws.title = 'Summary Dashboard'
    _banner(ws, examTitle, f'{examSubtitle}  |  N = {nTotal} students'.strip(' |'), 8)
    for i, w in enumerate([24, 14, 14, 14, 14, 14, 14, 50], 1): ws.column_dimensions[colL(i)].width = w

    # cohort performance
    r = 4; _section(ws, r, 'COHORT PERFORMANCE')
    for j, hname in enumerate(['Session', 'N', f'Mean /{mx}', 'Mean %', f'Median /{mx}', 'Std Dev', 'Range', 'KR-20 reliability'], 1):
        _put(ws, f'{colL(j)}{r + 1}', hname, HDR, 'FFFFFF', True, h='center')
    for i, s in enumerate(sessList):
        st = stats[s['label']]; rr = r + 2 + i
        vals = [s['label'], st['n'], st['mean'], st['pctMean'], st['median'], st['sd'], f"{st['min']:.0f} – {st['max']:.0f}",
                st['kr20'] if st['kr20'] is not None else 'n/a (random items)']
        for j, v in enumerate(vals, 1):
            _put(ws, f'{colL(j)}{rr}', v, 'F2F3F4' if i % 2 else 'FFFFFF', b=(j == 1), h='left' if j == 1 else 'center',
                 fmt={3: '0.00', 4: '0.0%', 6: '0.00', 8: '0.00'}.get(j))
    r = r + 2 + len(sessList) + 1

    # session comparison
    if multi:
        a, b = sessList[0]['label'], sessList[1]['label']
        t, p, d = welchTest(stats[a]['pct'], stats[b]['pct'])
        _section(ws, r, f'{a} vs {b} comparison', bg=HDR)
        diff = stats[a]['pctMean'] - stats[b]['pctMean']
        _put(ws, f'A{r + 1}', f'Mean difference {a} − {b} = {diff:+.1%};  Welch t = {t:.2f}, p {fmtP(p)};  Cohen\'s d = {d:.2f}  →  '
             + ('sessions differ significantly.' if p < .05 else 'the two sessions performed equivalently.'), wrap=True)
        ws.merge_cells(f'A{r + 1}:H{r + 1}'); ws.row_dimensions[r + 1].height = 32
        r += 3

    # difficulty distribution
    k = len(sessList); _section(ws, r, 'QUESTION DIFFICULTY DISTRIBUTION')
    qStart = 2 + k; spans = []
    free = list(range(qStart, 9))
    if k == 1: spans = [(free[0], free[-1])]
    else:
        lastCols = free[-(k - 1):]; spans = [(free[0], lastCols[0] - 1)] + [(c, c) for c in lastCols]
    hdrs = ['Tier'] + [f'{s["label"]} count' if multi else 'Count' for s in sessList]
    for j, hname in enumerate(hdrs, 1): _put(ws, f'{colL(j)}{r + 1}', hname, HDR, 'FFFFFF', True)
    for s, (c0, c1) in zip(sessList, spans):
        _put(ws, f'{colL(c0)}{r + 1}', f'{s["label"]} questions' if multi else 'Questions', HDR, 'FFFFFF', True)
        for c in range(c0 + 1, c1 + 1): ws[f'{colL(c)}{r + 1}'].fill = _fill(HDR)
        if c1 > c0: ws.merge_cells(f'{colL(c0)}{r + 1}:{colL(c1)}{r + 1}')
    for i, tier in enumerate(['High', 'Good', 'Moderate', 'Low']):
        bg, fg, lab = TIER_STYLE[tier]; rr = r + 2 + i
        _put(ws, f'A{rr}', lab, bg, fg, True)
        for j, s in enumerate(sessList):
            it = items[s['label']]; qs = it[it.pCorrect.map(difficultyFlag) == tier].q.tolist()
            _put(ws, f'{colL(2 + j)}{rr}', len(qs), bg, fg, h='center')
            c0, c1 = spans[j]; _put(ws, f'{colL(c0)}{rr}', ', '.join(qs) or '—', bg, fg, wrap=True)
            for c in range(c0 + 1, c1 + 1): ws[f'{colL(c)}{rr}'].fill = _fill(bg)
            if c1 > c0: ws.merge_cells(f'{colL(c0)}{rr}:{colL(c1)}{rr}')
        ws.row_dimensions[rr].height = 30
    r += 7

    # score distribution
    _section(ws, r, f'SCORE DISTRIBUTION (out of {mx})')
    hdrs = ['Score'] + ([f'{s["label"]} students' for s in sessList] + ['Total'] if multi else ['Students']) + ['% of cohort', 'Cumulative %']
    for j, hname in enumerate(hdrs, 1): _put(ws, f'{colL(j)}{r + 1}', hname, HDR, 'FFFFFF', True, h='center')
    lo = int(min(st['min'] for st in stats.values())); hi = int(max(st['max'] for st in stats.values())); cum = 0
    for i, sv in enumerate(range(lo, hi + 1)):
        cnts = [int((stats[s['label']]['tot'] == sv).sum()) for s in sessList]; tot = sum(cnts); cum += tot
        vals = [f'{sv}/{mx}  ({sv / mx:.0%})'] + (cnts + [tot] if multi else cnts) + [tot / nTotal, cum / nTotal]
        for j, v in enumerate(vals, 1):
            _put(ws, f'{colL(j)}{r + 2 + i}', v, 'F2F3F4' if i % 2 else 'FFFFFF', h='center', fmt='0.0%' if j >= len(vals) - 1 else None)
    r += hi - lo + 4

    # key findings
    _section(ws, r, 'KEY FINDINGS')
    for i, f in enumerate(findings):
        rr = r + 1 + i; _put(ws, f'A{rr}', f'• {f}', wrap=True); ws.merge_cells(f'A{rr}:H{rr}')
        ws.row_dimensions[rr].height = 16 * max(2, math.ceil(len(f) / 150))

    # question sheets
    for s in sessList:
        it = items[s['label']]; pool = s['isPool']
        others = [o['label'] for o in sessList if o['label'] != s['label']]
        twin = lambda q: '; '.join(f'{o} {pByLink[(o, linkOf[(s["label"], q)])][0]} ({pByLink[(o, linkOf[(s["label"], q)])][1]:.0%})'
                                   for o in others if (o, linkOf[(s['label'], q)]) in pByLink)
        hasTwin = any(twin(q) for q in it.q)
        cols = [('Q#', 6), ('Question Stem', 60), ('Correct Answer', 40)] + ([('Same question in other session (% correct)', 18)] if hasTwin else []) + \
               ([('Shown to (#)', 9)] if pool else []) + \
               [('Correct (#)', 10), ('Incorrect (#)', 11), ('% Correct', 10), ('Std Dev', 9), ('Discrimination D (top–bottom 27%)', 15),
                ('Item-rest r_pb', 11), ('Difficulty Flag', 16), ('Wrong answers chosen (count)', 60), ('Review note', 38)]
        q = wb.create_sheet(f'{s["label"]} Question Analysis'[:31] if multi else 'Question Analysis')
        _banner(q, f'{s["label"]} Session — Question-Level Analysis' if multi else 'Question-Level Analysis',
                f'N = {stats[s["label"]]["n"]}  |  1 mark per item' + (f'  |  {s["maxScore"]} of {s["nItems"]} items drawn per student' if pool else '') +
                '  |  Green ≥90% · Blue 75–89% · Amber 60–74% · Red <60%  |  D and r_pb: ≥0.20 good, <0 review', len(cols))
        for j, (hname, w) in enumerate(cols, 1):
            _put(q, f'{colL(j)}4', hname, NAVY, 'FFFFFF', True, wrap=True, h='center'); q.column_dimensions[colL(j)].width = w
        q.row_dimensions[4].height = 45
        for i, row in it.iterrows():
            bg, fg, lab = TIER_STYLE[difficultyFlag(row.pCorrect)]; rr = 5 + i
            vals = [row.q, row.stem, row.key] + ([twin(row.q) or '—'] if hasTwin else []) + ([row.nShown] if pool else []) + \
                   [row.nCorrect, row.nIncorrect, row.pCorrect, row.sd, None if pd.isna(row.discD) else row.discD,
                    None if pd.isna(row.rPbis) else row.rPbis, lab, row.wrongAnswers or '—', _reviewNote(row)]
            nC = len(vals)
            for j, v in enumerate(vals, 1):
                fmt = {nC - 6: '0.0%', nC - 5: '0.00', nC - 4: '0.00', nC - 3: '0.00'}.get(j)
                _put(q, f'{colL(j)}{rr}', v, bg, fg if j in (1, nC - 2) else '1A1A2E', b=(j == 1), wrap=j in (2, 3, nC - 1, nC),
                     h='center' if 4 <= j <= nC - 2 else 'left', fmt=fmt, border=True)
            if pd.notna(row.rPbis) and row.rPbis < 0: q[f'{colL(nC - 3)}{rr}'].font = Font(color='C0392B', bold=True)
        q.freeze_panes = 'C5'

    # student scores
    sw = wb.create_sheet('Student Scores')
    _banner(sw, 'Student Scores', 'Sorted by score (ascending) – lowest performers first  |  θ = Rasch ability (EAP, logits; cohort mean 0)', 9)
    for j, (hname, w) in enumerate([('Name', 28), ('Student Number', 15), ('Session', 12), ('Submitted (Melbourne)', 20), (f'Score /{mx}', 10),
                                    ('%', 9), ('Items wrong', 40), ('IRT θ (EAP)', 12), ('θ SE', 9)], 1):
        _put(sw, f'{colL(j)}4', hname, NAVY, 'FFFFFF', True, h='center'); sw.column_dimensions[colL(j)].width = w
    rows = []
    for s in sessList:
        d, sc = s['df'], s['sc']
        for i in d.index:
            sub = pd.to_datetime(str(d.submitted[i]).replace(' UTC', ''), utc=True).tz_convert('Australia/Melbourne').strftime('%d/%m/%Y %H:%M')
            th = irt['persons'][(s['label'], int(d.sis_id[i]))] if irt else (None, None)
            rows.append([d.at[i, 'name'], int(d.sis_id[i]), s['label'], sub, int(sc.loc[i].sum()), sc.loc[i].sum() / sc.loc[i].notna().sum(),
                         ', '.join(sc.columns[sc.loc[i] == 0]) or '—', *th])
    rows.sort(key=lambda x: (x[5], x[7] if x[7] is not None else 0, x[0]))
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            _put(sw, f'{colL(j)}{5 + i}', v, 'F2F3F4' if i % 2 else 'FFFFFF', h='left' if j in (1, 7) else 'center',
                 fmt={6: '0%', 8: '0.00', 9: '0.00'}.get(j))
        if row[5] < passMark: sw[f'E{5 + i}'].font = Font(color='C0392B', bold=True)
    sw.freeze_panes = 'A5'; sw.auto_filter.ref = f'A4:I{4 + len(rows)}'

    # performance chart
    c = wb.create_sheet('Performance Chart')
    _put(c, 'A1', 'Chart Data — % Correct by Question', NAVY, 'FFFFFF', True); c.merge_cells(f'A1:{colL(1 + k)}1')
    _put(c, 'A2', 'Question', HDR, 'FFFFFF', True)
    nRows = max(s['nItems'] for s in sessList)
    for j, s in enumerate(sessList):
        _put(c, f'{colL(2 + j)}2', f'{s["label"]} % Correct' if multi else '% Correct', HDR, 'FFFFFF', True)
        c.column_dimensions[colL(2 + j)].width = 14
        for i, pv in enumerate(items[s['label']].pCorrect): c.cell(3 + i, 2 + j, pv).number_format = '0%'
    for i in range(nRows): c.cell(3 + i, 1, f'Q{i + 1}')
    palette = [UOM_BLUE, UOM_ORANGE, '1E8449']
    for j, s in enumerate(sessList):
        ch = BarChart(); ch.type = 'col'; ch.title = f'{s["label"]} — % correct by question' if multi else '% correct by question'
        ch.y_axis.title = '% correct'; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.y_axis.number_format = '0%'
        ch.y_axis.delete = False; ch.x_axis.delete = False
        ch.add_data(Reference(c, min_col=2 + j, min_row=2, max_row=2 + s['nItems']), titles_from_data=True)
        ch.set_categories(Reference(c, min_col=1, min_row=3, max_row=2 + s['nItems']))
        ch.series[0].graphicalProperties.solidFill = palette[j % 3]; ch.legend = None
        ch.width, ch.height = max(22, s['nItems'] * 0.9), 8
        c.add_chart(ch, f'{colL(3 + k)}{2 + j * 17}')

    # IRT item calibration
    if irt:
        ir = wb.create_sheet('IRT Item Calibration')
        _banner(ir, 'IRT — Rasch (1PL) Item Calibration' + (', all sessions combined' if multi else ''), None, 10)
        irCols = [('Item', 18), ('Session', 12), ('Question Stem', 60), ('% Correct', 10), ('b (difficulty)', 12), ('SE', 8),
                  ('Infit MNSQ', 10), ('Outfit MNSQ', 10), ('Fit', 10), ('Difficulty band', 24)]
        for j, (hname, w) in enumerate(irCols, 1):
            _put(ir, f'{colL(j)}3', hname, NAVY, 'FFFFFF', True, wrap=True, h='center'); ir.column_dimensions[colL(j)].width = w
        stemOf = {}
        for s_ in sessList:
            for qq, st in zip(items[s_['label']].q, items[s_['label']].stem): stemOf.setdefault(linkOf[(s_['label'], qq)], st)
        recs = irt['items'].sort_values('b', ascending=False).to_dict('records') + \
               [dict(item=x, b=None, se=None, infit=None, outfit=None, p=irt['pAll'][x]) for x in irt['notEstimable']]
        for i, rec in enumerate(recs):
            if rec['item'].startswith('LINK|'): sLab, qq = 'Both', linkName[rec['item']]
            else: sLab, qq = rec['item'].split('|')
            rr = 4 + i
            if rec['b'] is None:
                band, fit, bg = ('All correct – not estimable' if rec['p'] == 1 else 'All wrong – not estimable' if rec['p'] == 0 else 'Too few responses'), '', 'F2F3F4'
            else:
                band = 'Hard for cohort (b > −1)' if rec['b'] > -1 else 'Moderate (−2.5 < b ≤ −1)' if rec['b'] > -2.5 else 'Easy (b ≤ −2.5)'
                fit = 'OK' if 0.7 <= rec['infit'] <= 1.3 else 'Misfit'
                bg = 'FADBD8' if band.startswith('Hard') else 'FEF9E7' if band.startswith('Moderate') else 'D5F5E3'
            for j, v in enumerate([qq, sLab, stemOf[rec['item']], rec['p'], rec['b'], rec['se'], rec['infit'], rec['outfit'], fit, band], 1):
                _put(ir, f'{colL(j)}{rr}', v, bg, wrap=(j == 3), h='left' if j in (3, 10) else 'center',
                     fmt={4: '0.0%', 5: '0.00', 6: '0.00', 7: '0.00', 8: '0.00'}.get(j), border=True)
        ir.freeze_panes = 'D4'

    wb.save(outPath)
    return dict(sessions=sessList, stats=stats, items=items, irt=irt, findings=findings, outPath=outPath)
