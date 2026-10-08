"""
oral_pres_utils.py — Oral Presentation (DASH `oral-presentation` forms) → Excel report.

Built for BOH3 2026 (checklist BOH3-OHTR, "Oral Health Therapy Research"), double-marked.
Called from main.ipynb section "## Oral Presentation" via `from oral_pres_utils import *`.

Public functions (everything else is private / not exported by `import *`):
    fetchOralPresentationJson(bearerToken, outPath, year, cohorts)   pull from DASH (ordering=id) + dedupe + check
    checkOralExport(jsonPath)                                         integrity check of a pulled JSON
    loadOralPresentation(jsonPath, rosterPath, checklistKey, cohort)  parse + score + all statistics (dict)
    buildOralPresentationReport(jsonPath, outPath, ...)              full workbook from scratch (current layout)
    refreshOralPresentationReport(jsonPath, workbookPath, ...)       rewrite values into an EXISTING workbook,
                                                                      keeping manual edits (layout, widths, sorts,
                                                                      removed columns/sections, charts)
Scoring: Excellent 4 · Good 3 · Satisfactory 2 · Poor 1 · Missing 0 (an unanswered item counts as Missing).
Examiner % = total / (4 × n items) × 100; Final % = mean of the examiners' %.
Examiner 1 / Examiner 2 = order of form id (submission order), not a role.
"""
import os
import re
import csv
import json
import copy as _copy
import shutil as _shutil
import zipfile as _zipfile
import datetime as _dt
import statistics as _st
import collections as _collections

import numpy as _np
from scipy import stats as _scipyStats
import openpyxl as _openpyxl
from openpyxl import Workbook as _XlWorkbook
from openpyxl.styles import Font as _XlFont, PatternFill as _XlPatternFill, Alignment as _XlAlignment, \
    Border as _XlBorder, Side as _XlSide
from openpyxl.utils import get_column_letter as _xlColLetter
from openpyxl.chart import BarChart as _XlBarChart, Reference as _XlReference
from openpyxl.chart.label import DataLabelList as _XlDataLabelList
from openpyxl.formatting.rule import ColorScaleRule as _XlColorScaleRule

__all__ = ["ORAL_LEVEL_POINTS", "fetchOralPresentationJson", "checkOralExport", "loadOralPresentation",
           "buildOralPresentationReport", "refreshOralPresentationReport"]

# ---------------------------------------------------------------- config
ORAL_LEVEL_POINTS = {"Excellent": 4, "Good": 3, "Satisfactory": 2, "Poor": 1, "Missing": 0}
_LEVELS = ["Excellent", "Good", "Satisfactory", "Poor", "Missing"]
_MAX_PTS = 4
_LOW_PCT = 50.0          # all-Satisfactory average
_HIGH_PCT = 75.0         # Good average
_DISAGREE_PCT = 20.0     # examiner gap flag (percentage points)
_BASE_URL = "https://api.unimelb-dash.com"

# palette (same as the BOH3 viva report)
_NAVY, _BLUE, _BLUE2, _GREEN, _OCHRE = "1F3864", "2E5496", "3A5FA0", "548235", "7F6000"
_SUBHDR, _INK = "D6E0F0", "1E293B"
_LEVEL_FILL = {4: ("C6EFCE", "006100"), 3: ("E2EFDA", "375623"), 2: ("FFEB9C", "9C5700"),
               1: ("FFC7CE", "9C0006"), 0: ("D9D9D9", "000000")}
_CHART_COL = {"Excellent": "548235", "Good": "A9D08E", "Satisfactory": "FFC000", "Poor": "C00000", "Missing": "7F7F7F"}
_THIN = _XlSide(style="thin", color="BFBFBF")
_BORDER = _XlBorder(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BANDS = [("0–39", 0, 40), ("40–49", 40, 50), ("50–59", 50, 60), ("60–69", 60, 70),
          ("70–79", 70, 80), ("80–89", 80, 90), ("90–100", 90, 101)]


# ================================================================ fetch + integrity
def fetchOralPresentationJson(bearerToken, outPath, year=2026, cohorts="BOH3", baseUrl=_BASE_URL):
    """Pull every oral-presentation form from DASH and save it as JSON.

    Uses ordering=id + page_size=max. (The generic pull cell orders by datetime; most forms share one
    datetime, so paged results were unstable — some forms came back twice and others were dropped.)
    Any repeated id is removed before saving. Returns the list of records.
    """
    import requests
    url = f"{baseUrl}/assessment/oral-presentation/get?page_size=max&page=1&cohort={cohorts}&year={year}&ordering=id"
    headers = {"Authorization": f"Token {bearerToken}", "Accept": "application/json"}
    records = []
    with requests.Session() as session:
        while url:
            resp = session.get(url, headers=headers, timeout=60)
            resp.raise_for_status()
            payload = resp.json()
            page = payload.get("results")
            if not isinstance(page, list):
                raise ValueError("Unexpected response: missing 'results' list")
            records.extend(page)
            url = payload.get("next")
    byId, dups = {}, _collections.Counter()
    for r in records:
        if r["id"] in byId: dups[r["id"]] += 1
        byId[r["id"]] = r
    if dups: print(f"WARNING: DASH returned {sum(dups.values())} repeated row(s) {sorted(dups)} — removed.")
    records = sorted(byId.values(), key=lambda r: r["id"])
    with open(outPath, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    print(f"Saved {len(records)} unique records -> {outPath}")
    checkOralExport(outPath)
    return records


def checkOralExport(jsonPath):
    """Print and return integrity facts: repeated ids, submitted count, students without exactly 2 forms."""
    records = json.load(open(jsonPath, encoding="utf-8"))
    ids = _collections.Counter(r["id"] for r in records)
    uniq = {r["id"]: r for r in records}
    submitted = [r for r in uniq.values() if r.get("submitted")]
    real = [r for r in submitted if not _isTestName(r["student"])]
    perStudent = _collections.Counter(r["student"].strip().lower() for r in real)
    sameExaminer = [k for k, v in _collections.Counter((r["student"].strip().lower(), r["assessor"]) for r in real).items() if v > 1]
    out = {"records": len(records), "unique": len(uniq), "repeatedIds": sorted(i for i, n in ids.items() if n > 1),
           "submitted": len(submitted), "submittedReal": len(real), "students": len(perStudent),
           "notTwoForms": {k: v for k, v in perStudent.items() if v != 2}, "sameExaminerTwice": sameExaminer}
    print(f"{out['records']} records · {out['unique']} unique · {out['submittedReal']} submitted (real students) · "
          f"{out['students']} students")
    if out["repeatedIds"]: print("  repeated ids (export problem — re-pull):", out["repeatedIds"])
    if out["notTwoForms"]: print("  students without exactly 2 submitted forms:", out["notTwoForms"])
    if sameExaminer: print("  same examiner twice for a student:", sameExaminer)
    if not (out["repeatedIds"] or out["notTwoForms"] or sameExaminer): print("  OK — no duplicates, every student double-marked.")
    return out


def _isTestName(name):
    n = (name or "").lower()
    return "dummy" in n or "test" in n


# ================================================================ load + score + stats
def _surnameKey(n):
    return (n.split()[-1].lower(), n.lower())


def _cleanName(n):
    n = n.strip()
    return n.title() if n.islower() else n


def _cronbachAlpha(M):
    k = M.shape[1]
    return k / (k - 1) * (1 - M.var(axis=0, ddof=1).sum() / M.sum(axis=1).var(ddof=1))


def _weightedKappa(a, b, cats=(0, 1, 2, 3, 4)):
    idx = {c: i for i, c in enumerate(cats)}
    k = len(cats)
    O = _np.zeros((k, k))
    for x, y in zip(a, b): O[idx[x], idx[y]] += 1
    O /= O.sum()
    E = _np.outer(O.sum(1), O.sum(0))
    W = _np.array([[((i - j) / (k - 1)) ** 2 for j in range(k)] for i in range(k)])
    den = (W * E).sum()
    return float(1 - (W * O).sum() / den) if den else None


def loadOralPresentation(jsonPath, rosterPath="studentEmailList.csv", checklistKey="BOH3-OHTR", cohort="BOH3"):
    """Parse the DASH JSON, score every submitted form and compute all report statistics.

    Returns a dict: forms, excluded, students (sorted by surname), mcs, items, rubric, itemStats, alpha,
    examiners, pairs, cohortMean, formMean, notPresented, rosterCount, dates.
    Excluded: repeated ids, drafts, dummy/test students, forms without checklist answers.
    """
    raw = json.load(open(jsonPath, encoding="utf-8"))
    roster = {r["student_name"].strip().lower(): r for r in csv.DictReader(open(rosterPath, encoding="utf-8"))}
    rosterCohort = {k: v for k, v in roster.items() if v.get("cohort") == cohort}
    ref = next(r for r in raw if checklistKey in r["form"]["checklists"])
    items = ref["form"]["checklists"][checklistKey]["fields"]
    mcs = list(items)
    rubric = ref["form"]["checklists"][checklistKey]["extra_config"]["rubric"]

    excluded, seen, forms = [], set(), []
    for r in raw:
        if r.get("cohort") and r["cohort"] != cohort: continue
        ans = (r["form"]["data"] or {}).get("assessor", {}) or {}
        reason = None
        if r["id"] in seen: reason = "Repeated form id in export"
        elif not r.get("submitted"): reason = "Not submitted (draft)"
        elif _isTestName(r["student"]): reason = "Dummy/test student"
        elif not ans.get(checklistKey): reason = "No checklist answers"
        seen.add(r["id"])
        if reason:
            excluded.append(dict(id=r["id"], student=r["student"], assessor=r["assessor"], date=r["datetime"][:10], reason=reason))
            continue
        ros = roster.get(r["student"].strip().lower(), {})
        labels = {mc: ans[checklistKey].get(mc) or "Missing" for mc in mcs}
        pts = {mc: ORAL_LEVEL_POINTS[labels[mc]] for mc in mcs}
        total = sum(pts.values())
        stNo = r.get("student_number") or (int(ros["student_number"]) if ros else None)
        forms.append(dict(id=r["id"], date=_dt.date.fromisoformat(r["datetime"][:10]),
                          student=_cleanName(ros.get("student_name", r["student"])), studentNo=stNo,
                          subject=r.get("subject"), assessor=r["assessor"], labels=labels, pts=pts, total=total,
                          pct=100 * total / (_MAX_PTS * len(mcs)), nPoor=sum(p <= 1 for p in pts.values()),
                          comments=(ans.get("comments") or "").strip()))

    byStudent = _collections.defaultdict(list)
    for f in forms: byStudent[f["student"]].append(f)
    students = []
    for name, fs in byStudent.items():
        fs.sort(key=lambda f: f["id"])
        e1, e2 = fs[0], (fs[1] if len(fs) > 1 else None)
        meanPct = _st.mean(f["pct"] for f in fs)
        gap = abs(e1["pct"] - e2["pct"]) if e2 else None
        flags = []
        if not e2: flags.append("Single-marked")
        if meanPct < _LOW_PCT: flags.append(f"Mean < {_LOW_PCT:.0f}%")
        if any(f["nPoor"] for f in fs): flags.append("Poor/Missing rating")
        if gap is not None and gap >= _DISAGREE_PCT: flags.append(f"Examiner gap ≥ {_DISAGREE_PCT:.0f} pts")
        if len(fs) > 2: flags.append(f"{len(fs)} forms")
        students.append(dict(name=name, no=e1["studentNo"], subject=e1["subject"], date=e1["date"], e1=e1, e2=e2,
                             forms=fs, meanPct=meanPct, gap=gap, flags=flags, nPoor=sum(f["nPoor"] for f in fs),
                             itemMean={mc: _st.mean(f["pts"][mc] for f in fs) for mc in mcs}))
    students.sort(key=lambda s: _surnameKey(s["name"]))

    X = _np.array([[s["itemMean"][mc] for mc in mcs] for s in students], float)
    alpha = float(_cronbachAlpha(X))
    itemStats = []
    for j, mc in enumerate(mcs):
        rest = _np.delete(X, j, axis=1)
        allPts = [f["pts"][mc] for f in forms]
        cnt = _collections.Counter(f["labels"][mc] for f in forms)
        pr = [(s["e1"]["pts"][mc], s["e2"]["pts"][mc]) for s in students if s["e2"]]
        itemStats.append(dict(mc=mc, desc=items[mc], n=len(allPts), dist={l: cnt.get(l, 0) for l in _LEVELS},
                              meanPts=float(_np.mean(allPts)), pct=float(100 * _np.mean(allPts) / _MAX_PTS),
                              sd=float(_np.std(allPts, ddof=1)),
                              ritc=float(_np.corrcoef(X[:, j], rest.sum(axis=1))[0, 1]),
                              alphaDel=float(_cronbachAlpha(rest)),
                              exact=float(100 * _np.mean([a == b for a, b in pr])) if pr else None,
                              adj=float(100 * _np.mean([abs(a - b) <= 1 for a, b in pr])) if pr else None))

    formMean = float(_np.mean([f["pct"] for f in forms]))
    byEx = _collections.defaultdict(list)
    for f in forms: byEx[f["assessor"]].append(f)
    examiners = []
    for ex, fs in byEx.items():
        p = [f["pct"] for f in fs]
        labs = [f["labels"][mc] for f in fs for mc in mcs]
        examiners.append(dict(name=ex, n=len(fs), mean=float(_np.mean(p)), median=float(_np.median(p)),
                              sd=float(_np.std(p, ddof=1)) if len(p) > 1 else None, delta=float(_np.mean(p)) - formMean,
                              mix={l: 100 * labs.count(l) / len(labs) for l in _LEVELS},
                              itemMean={mc: float(_np.mean([f["pts"][mc] for f in fs])) for mc in mcs}))
    examiners.sort(key=lambda e: -e["mean"])

    pairMap = _collections.defaultdict(list)
    for s in students:
        if s["e2"]:
            a, b = sorted([s["e1"], s["e2"]], key=lambda f: f["assessor"])
            pairMap[(a["assessor"], b["assessor"])].append((a, b))
    pairs = []
    for (A, B), lst in sorted(pairMap.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        ap, bp = [a["pct"] for a, b in lst], [b["pct"] for a, b in lst]
        ai = [a["pts"][mc] for a, b in lst for mc in mcs]
        bi = [b["pts"][mc] for a, b in lst for mc in mcs]
        pval = float(_scipyStats.ttest_rel(ap, bp).pvalue) if len(lst) >= 3 else None
        pairs.append(dict(label=f"{A} | {B}", n=len(lst), aMean=float(_np.mean(ap)), bMean=float(_np.mean(bp)),
                          diff=float(_np.mean(ap) - _np.mean(bp)), absGap=float(_np.mean(_np.abs(_np.subtract(ap, bp)))),
                          p=None if pval is None or _np.isnan(pval) else pval,
                          exact=100 * float(_np.mean(_np.equal(ai, bi))),
                          adj=100 * float(_np.mean([abs(x - y) <= 1 for x, y in zip(ai, bi)])),
                          kappa=_weightedKappa(ai, bi) if len(lst) >= 3 else None,
                          nGap=sum(abs(x - y) >= _DISAGREE_PCT for x, y in zip(ap, bp))))

    presented = {s["name"].lower() for s in students}
    notPresented = sorted([v for k, v in rosterCohort.items() if k not in presented],
                          key=lambda v: _surnameKey(v["student_name"]))
    return dict(forms=forms, excluded=excluded, students=students, mcs=mcs, items=items, rubric=rubric,
                itemStats=itemStats, alpha=alpha, examiners=examiners, pairs=pairs, formMean=formMean,
                cohortMean=float(_np.mean([s["meanPct"] for s in students])), notPresented=notPresented,
                rosterCount=len(rosterCohort), dates=sorted({f["date"] for f in forms}),
                checklistKey=checklistKey, checklistName=ref["form"]["checklists"][checklistKey]["name"], cohort=cohort)


def _kpis(m):
    st, forms, mcs = m["students"], m["forms"], m["mcs"]
    means = [s["meanPct"] for s in st]
    gaps = [s["gap"] for s in st if s["gap"] is not None]
    return {"Students assessed": len(st), "Double-marked": sum(1 for s in st if s["e2"]),
            "Single-marked": sum(1 for s in st if not s["e2"]), "Submitted forms used": len(forms),
            "Examiners": len(m["examiners"]), "Mean final %": round(float(_np.mean(means)), 1),
            "Median final %": round(float(_np.median(means)), 1), "SD": round(float(_np.std(means, ddof=1)), 1),
            "Min / Max": f"{min(means):.1f} / {max(means):.1f}",
            f"Students with mean < {_LOW_PCT:.0f}%": sum(v < _LOW_PCT for v in means),
            "Students with ≥1 Poor/Missing rating": sum(1 for s in st if s["nPoor"]),
            f"Examiner gap ≥ {_DISAGREE_PCT:.0f} pts": sum(g >= _DISAGREE_PCT for g in gaps),
            "Mean examiner gap (pts)": round(float(_np.mean(gaps)), 1) if gaps else None,
            "Cronbach's α (6 criteria)": round(m["alpha"], 2),
            "Overall rating mix": " · ".join(
                f"{l} {100 * sum(1 for f in forms for mc in mcs if f['labels'][mc] == l) / (len(forms) * len(mcs)):.0f}%"
                for l in _LEVELS)}


def _descriptionValues(m):
    subj = _collections.Counter(s["subject"] for s in m["students"])
    return {"Generated": _dt.datetime.now().strftime("%d %b %Y %H:%M"),
            "Cohort / subject(s)": ", ".join(f"{k} ({v})" for k, v in subj.items()),
            "Assessment date(s)": ", ".join(d.strftime("%d %b %Y") for d in m["dates"]),
            "Students assessed": f"{len(m['students'])} of {m['rosterCount']} {m['cohort']} on roster "
                                 f"({len(m['notPresented'])} with no submitted form — see Data Quality)",
            "Forms used": f"{len(m['forms'])} submitted forms"}


# ================================================================ cell helpers
def _fill(h):
    return _XlPatternFill("solid", fgColor=h)


_NOFILL = _XlPatternFill(fill_type=None)


def _hdr(c, text, bg=_NAVY, fg="FFFFFF", sz=11):
    c.value = text
    c.font = _XlFont(name="Calibri", sz=sz, b=True, color=fg)
    c.fill = _fill(bg)
    c.alignment = _XlAlignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = _BORDER


def _sub(c, text):
    _hdr(c, text, bg=_SUBHDR, fg=_NAVY, sz=9)


def _cell(ws, r, col, v, bold=False, numFmt=None, bg=None, fg=_INK, align="center", wrap=False, sz=9, vertical="center"):
    c = ws.cell(r, col, v)
    c.font = _XlFont(name="Calibri", sz=sz, b=bold, color=fg)
    c.border = _BORDER
    c.alignment = _XlAlignment(horizontal=align, vertical=vertical, wrap_text=wrap)
    if numFmt: c.number_format = numFmt
    if bg: c.fill = _fill(bg)
    return c


def _restyle(c, v, bg=None, fg=_INK, bold=False):
    """Set value + value-dependent fill/font colour; keep the cell's alignment, border, number format and size."""
    c.value = v
    c.fill = _fill(bg) if bg else _NOFILL
    f = c.font
    c.font = _XlFont(name=f.name or "Calibri", sz=f.sz or 9, b=bold, i=f.i, color=fg)


def _levelCell(c, pts):
    if pts is None: _restyle(c, None, "F2F2F2")
    else:
        bg, fg = _LEVEL_FILL[pts]
        _restyle(c, pts, bg, fg, bold=pts <= 1)


def _pctCell(c, v, boldNormal=False):
    if v is None: _restyle(c, None)
    elif v < _LOW_PCT: _restyle(c, v, "FFC7CE", "9C0006", True)
    elif v >= _HIGH_PCT: _restyle(c, v, "C6EFCE", "006100", boldNormal)
    else: _restyle(c, v, None, _INK, boldNormal)


def _ritcFill(v):
    return "C6EFCE" if v >= 0.3 else "FFEB9C" if v >= 0.2 else "FFC7CE"


def _kappaFill(v):
    return None if v is None else "C6EFCE" if v >= 0.6 else "FFEB9C" if v >= 0.4 else "FFC7CE"


def _title(ws, text, note=None, width=10):
    _hdr(ws.cell(2, 2), text, sz=13)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=width)
    ws.row_dimensions[2].height = 24
    if note:
        c = ws.cell(3, 2, note)
        c.font = _XlFont(sz=9, i=True, color="4A235A")
        c.alignment = _XlAlignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=width)
        ws.row_dimensions[3].height = 30
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3


def _byFinal(m):
    return sorted(m["students"], key=lambda s: (-s["meanPct"], _surnameKey(s["name"])))


# ================================================================ full build
def buildOralPresentationReport(jsonPath, outPath, rosterPath="studentEmailList.csv", checklistKey="BOH3-OHTR",
                                cohort="BOH3", year=2026):
    """Build the whole workbook from scratch (layout = the reviewed 2026 version).

    Sheets: Description · Student Results (sorted by Final %, both examiners' comments at the end, unwrapped)
    · Cohort Summary (+2 charts) · Item Analysis · Examiner Analysis · Comments · Data Quality · Raw Forms.
    Use refreshOralPresentationReport instead when the workbook has manual edits worth keeping.
    """
    m = loadOralPresentation(jsonPath, rosterPath, checklistKey, cohort)
    mcs, items = m["mcs"], m["items"]
    nItems = len(mcs)
    wb = _XlWorkbook()

    # ---- Description
    ws = wb.active
    ws.title = "Description"
    _title(ws, f"{cohort} Oral Presentation {year} — {m['checklistName']} ({checklistKey})", width=4)
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 95
    d = _descriptionValues(m)
    lines = [("Generated", d["Generated"]), ("Cohort / subject(s)", d["Cohort / subject(s)"]),
             ("Assessment date(s)", d["Assessment date(s)"]), ("Students assessed", d["Students assessed"]),
             ("Forms used", d["Forms used"]),
             ("Checklist", f"{nItems} criteria (MC1–MC{nItems}), 4-level rubric; full descriptors on 'Item Analysis'."),
             ("Scoring", f"Excellent = 4 · Good = 3 · Satisfactory = 2 · Poor = 1 · Missing = 0 (an unanswered item also counts as Missing). Max {4 * nItems} per examiner."),
             ("Examiner %", f"Examiner total ÷ {4 * nItems} × 100 (all {nItems} criteria always counted)."),
             ("Final %", "Mean of the two examiners' %. Single-marked students use the one available form."),
             ("Reference line", f"{_LOW_PCT:.0f}% = all criteria Satisfactory. No cut score set in DASH — flags are for review, not a pass/fail decision."),
             ("Flags", f"Mean < {_LOW_PCT:.0f}% · any Poor/Missing rating · examiner gap ≥ {_DISAGREE_PCT:.0f} percentage points · single-marked."),
             ("Reliability", "Cronbach's α and item–rest correlations computed on student-level item means (average of examiners)."),
             ("Agreement", "Exact = same level; Adjacent = within one level. Weighted κ = quadratic over the 5 levels.")]
    r = 5
    for k, v in lines:
        _cell(ws, r, 2, k, bold=True, align="left", sz=10)
        _cell(ws, r, 3, v, align="left", wrap=True, sz=10)
        r += 1
    r += 1
    _cell(ws, r, 2, "Colour key", bold=True, align="left", sz=10); _cell(ws, r, 3, "", align="left", sz=10)
    r += 1
    for pts, lab in [(4, "Excellent (4)"), (3, "Good (3)"), (2, "Satisfactory (2)"), (1, "Poor (1)"), (0, "Missing (0)")]:
        c = _cell(ws, r, 2, None)
        _levelCell(c, pts); c.value = lab
        _cell(ws, r, 3, "Item rating cell", align="left")
        r += 1
    _pctCell(_cell(ws, r, 2, None), 0); ws.cell(r, 2).value = f"< {_LOW_PCT:.0f}%"
    _cell(ws, r, 3, "Percentage below the all-Satisfactory line", align="left"); r += 1
    _pctCell(_cell(ws, r, 2, None), 80); ws.cell(r, 2).value = f"≥ {_HIGH_PCT:.0f}%"
    _cell(ws, r, 3, f"Percentage at/above {_HIGH_PCT:.0f}% (Good average)", align="left"); r += 2
    _cell(ws, r, 2, "Sheets", bold=True, align="left", sz=10); r += 1
    for sname, dsc in [("Student Results", "One row per student (sorted by Final %): both examiners' item ratings, totals, final %, flags and comments"),
                       ("Cohort Summary", "Headline numbers, % band distribution and rating mix per criterion (charts)"),
                       ("Item Analysis", "Per-criterion distribution, mean, discrimination, α-if-deleted, examiner agreement + rubric"),
                       ("Examiner Analysis", "Examiner leniency/harshness, pair agreement (κ, paired t) and examiner × item means"),
                       ("Comments", "Examiner written feedback per student"),
                       ("Data Quality", "Single-marked students, roster students without a form"),
                       ("Raw Forms", "One row per submitted form with labels and points")]:
        _cell(ws, r, 2, sname, align="left"); _cell(ws, r, 3, dsc, align="left"); r += 1

    # ---- Student Results
    ws = wb.create_sheet("Student Results")
    ws.sheet_view.showGridLines = False
    info = ["Student No", "Student", "Date", "Examiner 1", "Examiner 2"]
    c0 = len(info) + 1
    resCols = [f"E1 Total\n/{4 * nItems}", f"E2 Total\n/{4 * nItems}", "E1 %", "E2 %", "Final %", "Examiner\nGap (pts)", "Poor/Missing\nRatings"]
    resStart = c0 + 2 * nItems
    flagCol = resStart + len(resCols)
    groups = [("Student Information", 1, len(info), _NAVY)]
    for i, mc in enumerate(mcs):
        groups.append((f"{mc} · {items[mc]}", c0 + 2 * i, c0 + 2 * i + 1, _BLUE if i % 2 == 0 else _BLUE2))
    groups += [("Result", resStart, flagCol - 1, _GREEN), ("Flags", flagCol, flagCol, _OCHRE),
               ("Comments", flagCol + 1, flagCol + 2, _OCHRE)]
    for text, a, b, bg in groups:
        for cc in range(a, b + 1): _hdr(ws.cell(1, cc), None, bg=bg, sz=9)
        _hdr(ws.cell(1, a), text, bg=bg, sz=9)
        if b > a: ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
    ws.row_dimensions[1].height = 42
    ws.row_dimensions[2].height = 30
    for i, h in enumerate(info): _sub(ws.cell(2, i + 1), h)
    for i in range(nItems):
        _sub(ws.cell(2, c0 + 2 * i), "E1"); _sub(ws.cell(2, c0 + 2 * i + 1), "E2")
    for i, h in enumerate(resCols): _sub(ws.cell(2, resStart + i), h)
    _sub(ws.cell(2, flagCol), "Flags")
    _sub(ws.cell(2, flagCol + 1), "Examiner 1 Comments"); _sub(ws.cell(2, flagCol + 2), "Examiner 2 Comments")
    order = _byFinal(m)
    for r, s in enumerate(order, start=3):
        for i in range(len(info)):
            _cell(ws, r, i + 1, None, align="left" if i in (1, 3, 4) else "center", numFmt="dd mmm yyyy" if i == 2 else None)
        for i in range(2 * nItems): _cell(ws, r, c0 + i, None)
        for i in range(len(resCols)): _cell(ws, r, resStart + i, None, numFmt="0.0" if 2 <= i <= 5 else None)
        _cell(ws, r, flagCol, None, align="left")
        _cell(ws, r, flagCol + 1, None, align="left"); _cell(ws, r, flagCol + 2, None, align="left")
    last = 2 + len(order)
    meanRow = last + 1
    _cell(ws, meanRow, 2, "Cohort mean", bold=True, align="left", bg=_SUBHDR)
    for col in list(range(c0, c0 + 2 * nItems)): _cell(ws, meanRow, col, None, numFmt="0.00", bold=True, bg=_SUBHDR)
    for col in (resStart + 2, resStart + 3, resStart + 4): _cell(ws, meanRow, col, None, numFmt="0.0", bold=True, bg=_SUBHDR)
    _writeStudentResults(ws, m)
    for k, w in {1: 11, 2: 24, 3: 11, 4: 17, 5: 17}.items(): ws.column_dimensions[_xlColLetter(k)].width = w
    for cc in range(c0, resStart): ws.column_dimensions[_xlColLetter(cc)].width = 5.2
    for cc in range(resStart, flagCol): ws.column_dimensions[_xlColLetter(cc)].width = 8.5
    ws.column_dimensions[_xlColLetter(flagCol - 1)].width = 14.9
    ws.column_dimensions[_xlColLetter(flagCol)].width = 42
    ws.column_dimensions[_xlColLetter(flagCol + 1)].width = 60
    ws.column_dimensions[_xlColLetter(flagCol + 2)].width = 60
    ws.freeze_panes = ws.cell(3, c0)
    ws.auto_filter.ref = f"A2:{_xlColLetter(flagCol + 2)}{last}"

    # ---- Cohort Summary
    ws = wb.create_sheet("Cohort Summary")
    _title(ws, f"{cohort} Oral Presentation {year} — Cohort Summary", width=13)
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 18
    kp = _kpis(m)
    for i, k in enumerate(kp):
        _cell(ws, 5 + i, 2, k, bold=True, align="left", sz=10)
        _cell(ws, 5 + i, 3, None, sz=10, bold=True, align="left" if isinstance(kp[k], str) else "center")
    _hdr(ws.cell(4, 5), "Final % band", bg=_BLUE, sz=10); _hdr(ws.cell(4, 6), "Students", bg=_BLUE, sz=10)
    for i, (lab, a, b) in enumerate(_BANDS):
        _cell(ws, 5 + i, 5, lab, sz=10); _cell(ws, 5 + i, 6, None, sz=10)
    ws.column_dimensions["E"].width = 13
    ws.column_dimensions["F"].width = 10
    ch = _XlBarChart(); ch.type = "col"; ch.title = "Distribution of final %"; ch.style = 10
    ch.add_data(_XlReference(ws, min_col=6, min_row=4, max_row=4 + len(_BANDS)), titles_from_data=True)
    ch.set_categories(_XlReference(ws, min_col=5, min_row=5, max_row=4 + len(_BANDS)))
    ch.legend = None; ch.y_axis.title = "Students"; ch.x_axis.title = "Final %"; ch.height, ch.width = 7, 13
    ch.series[0].graphicalProperties.solidFill = _BLUE
    ch.dataLabels = _XlDataLabelList(); ch.dataLabels.showVal = True
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ws.add_chart(ch, "H4")
    r0 = 22
    _hdr(ws.cell(r0, 2), "Criterion", bg=_BLUE, sz=10)
    for i, l in enumerate(_LEVELS): _hdr(ws.cell(r0, 3 + i), f"% {l}", bg=_BLUE, sz=10)
    _hdr(ws.cell(r0, 3 + len(_LEVELS)), "Mean %", bg=_BLUE, sz=10)
    for i, it in enumerate(m["itemStats"]):
        _cell(ws, r0 + 1 + i, 2, f"{it['mc']} · {it['desc']}", align="left")
        for k in range(len(_LEVELS)): _cell(ws, r0 + 1 + i, 3 + k, None, numFmt="0.0")
        _cell(ws, r0 + 1 + i, 3 + len(_LEVELS), None, numFmt="0.0", bold=True)
    for cc in "DEFGH": ws.column_dimensions[cc].width = 13
    ch = _XlBarChart(); ch.type = "bar"; ch.grouping = "percentStacked"; ch.overlap = 100
    ch.title = "Rating mix by criterion"
    ch.add_data(_XlReference(ws, min_col=3, max_col=2 + len(_LEVELS), min_row=r0, max_row=r0 + nItems), titles_from_data=True)
    ch.set_categories(_XlReference(ws, min_col=2, min_row=r0 + 1, max_row=r0 + nItems))
    for sser, l in zip(ch.series, _LEVELS): sser.graphicalProperties.solidFill = _CHART_COL[l]
    ch.x_axis.scaling.orientation = "maxMin"; ch.height, ch.width = 8, 18
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.legend.position = "b"
    ws.add_chart(ch, f"B{r0 + nItems + 3}")
    _writeCohortSummary(ws, m)

    # ---- Item Analysis
    ws = wb.create_sheet("Item Analysis")
    _title(ws, "Item Analysis — per criterion",
           f"Cronbach's α ({nItems} criteria, student-level means) = {m['alpha']:.2f}. Item–rest r ≥ 0.30 good discrimination; "
           "α-if-deleted above overall α suggests the item adds little.", width=17)
    cols = ["Item", "Criterion", "Ratings"] + _LEVELS + ["Mean pts\n(/4)", "SD", "Mean %", "Item–rest r", "α if deleted",
                                                        "Exact\nagreement %", "Adjacent\nagreement %"]
    for i, h in enumerate(cols): _hdr(ws.cell(5, 2 + i), h, bg=_BLUE, sz=9)
    ws.row_dimensions[5].height = 30
    fmts = [None, None, None] + [None] * len(_LEVELS) + ["0.00", "0.00", "0.0", "0.00", "0.00", "0.0", "0.0"]
    for i, it in enumerate(m["itemStats"]):
        for k, fm in enumerate(fmts): _cell(ws, 6 + i, 2 + k, None, numFmt=fm, align="left" if k == 1 else "center")
        ws.cell(6 + i, 2).value = it["mc"]; ws.cell(6 + i, 3).value = it["desc"]
    for i, w in enumerate([7, 44, 8, 10, 8, 12, 8, 8, 9, 7, 9, 10, 10, 11, 11]): ws.column_dimensions[_xlColLetter(2 + i)].width = w
    _writeItemAnalysis(ws, m)
    r = 8 + nItems
    for cc in range(2, 16): _hdr(ws.cell(r, cc), None, bg=_NAVY, sz=10)
    _hdr(ws.cell(r, 2), "Rubric descriptors", bg=_NAVY, sz=10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=15)
    r += 1
    _hdr(ws.cell(r, 2), "Item", bg=_BLUE, sz=9); _hdr(ws.cell(r, 3), "Criterion", bg=_BLUE, sz=9)
    spans = [(4, 6, "Excellent (4)", "O1"), (7, 9, "Good (3)", "O2"), (10, 12, "Satisfactory (2)", "O3"), (13, 15, "Poor (1)", "O4")]
    for a, b, t, _o in spans:
        for cc in range(a, b + 1): _hdr(ws.cell(r, cc), None, bg=_BLUE, sz=9)
        _hdr(ws.cell(r, a), t, bg=_BLUE, sz=9)
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    r += 1
    for mc in mcs:
        _cell(ws, r, 2, mc); _cell(ws, r, 3, items[mc], align="left", wrap=True)
        for a, b, t, o in spans:
            for cc in range(a, b + 1): _cell(ws, r, cc, None)
            _cell(ws, r, a, m["rubric"][mc].get(o, ""), align="left", wrap=True, sz=8, vertical="top")
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        ws.row_dimensions[r].height = 70
        r += 1

    # ---- Examiner Analysis
    ws = wb.create_sheet("Examiner Analysis")
    _title(ws, "Examiner Analysis", f"Δ vs cohort = examiner mean form % − mean of all forms ({m['formMean']:.1f}%). "
                                    "Positive = more lenient. Pairs compared on the same students.", width=13)
    cols = ["Examiner", "Forms", "Mean %", "Median %", "SD", "Δ vs cohort"] + [f"% {l}" for l in _LEVELS]
    for i, h in enumerate(cols): _hdr(ws.cell(5, 2 + i), h, bg=_BLUE, sz=9)
    for i in range(len(m["examiners"])):
        for k in range(len(cols)): _cell(ws, 6 + i, 2 + k, None, numFmt=None if k < 2 else "0.0", align="left" if k == 0 else "center")
        ws.cell(6 + i, 2).value = m["examiners"][i]["name"]
    r = 6 + len(m["examiners"]) + 2
    for cc in range(2, 13): _hdr(ws.cell(r, cc), None, bg=_NAVY, sz=10)
    _hdr(ws.cell(r, 2), "Examiner pairs — agreement on the same students", bg=_NAVY, sz=10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
    r += 1
    cols = ["Pair (A | B)", "Students", "A mean %", "B mean %", "Mean diff\n(A−B)", "Mean |gap|", "Paired t p",
            "Exact item\nagreement %", "Adjacent\nagreement %", "Weighted κ", "Gap ≥ 20"]
    for i, h in enumerate(cols): _hdr(ws.cell(r, 2 + i), h, bg=_BLUE, sz=9)
    ws.row_dimensions[r].height = 30
    fmts = [None, None, "0.0", "0.0", "0.0", "0.0", "0.000", "0.0", "0.0", "0.00", None]
    for i in range(len(m["pairs"])):
        for k, fm in enumerate(fmts): _cell(ws, r + 1 + i, 2 + k, None, numFmt=fm, align="left" if k == 0 else "center")
        ws.cell(r + 1 + i, 2).value = m["pairs"][i]["label"]
    r += 1 + len(m["pairs"])
    _cell(ws, r, 2, "κ guide: ≥0.60 substantial · 0.40–0.59 moderate · <0.40 fair/poor. Paired t p < 0.05 = systematic "
                    "difference between the two examiners.", align="left", fg="7F7F7F")
    r += 3
    for cc in range(2, 13): _hdr(ws.cell(r, cc), None, bg=_NAVY, sz=10)
    _hdr(ws.cell(r, 2), "Examiner × criterion — mean points (/4)", bg=_NAVY, sz=10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
    r += 1
    _hdr(ws.cell(r, 2), "Examiner", bg=_BLUE, sz=9)
    for i, mc in enumerate(mcs): _hdr(ws.cell(r, 3 + i), mc, bg=_BLUE, sz=9)
    _hdr(ws.cell(r, 3 + nItems), "Forms", bg=_BLUE, sz=9)
    heat0 = r + 1
    for i, e in enumerate(sorted(m["examiners"], key=lambda e: e["name"])):
        _cell(ws, heat0 + i, 2, e["name"], align="left")
        for k in range(nItems): _cell(ws, heat0 + i, 3 + k, None, numFmt="0.00")
        _cell(ws, heat0 + i, 3 + nItems, None)
    allRow = heat0 + len(m["examiners"])
    _cell(ws, allRow, 2, "All forms", bold=True, align="left", bg=_SUBHDR)
    for k in range(nItems): _cell(ws, allRow, 3 + k, None, numFmt="0.00", bold=True, bg=_SUBHDR)
    ws.conditional_formatting.add(f"C{heat0}:{_xlColLetter(2 + nItems)}{allRow - 1}",
                                  _XlColorScaleRule(start_type="num", start_value=2, start_color="F8696B", mid_type="num",
                                                    mid_value=3, mid_color="FFEB84", end_type="num", end_value=4, end_color="63BE7B"))
    ws.column_dimensions["B"].width = 36
    for cc in range(3, 13): ws.column_dimensions[_xlColLetter(cc)].width = 11
    _writeExaminerAnalysis(ws, m)

    # ---- Comments
    ws = wb.create_sheet("Comments")
    ws.sheet_view.showGridLines = False
    for i, h in enumerate(["Student No", "Student", "Examiner", "Final %", "Form %", "Comments"]): _hdr(ws.cell(1, 1 + i), h, bg=_NAVY, sz=10)
    nRows = sum(len(s["forms"]) for s in m["students"])
    for r in range(2, 2 + nRows):
        for col in range(1, 7):
            _cell(ws, r, col, None, numFmt="0.0" if col in (4, 5) else None, align="left" if col in (2, 3, 6) else "center",
                  vertical="top" if col == 6 else "center")
    for r in range(1, 2 + nRows): ws.row_dimensions[r].height = 15.75
    for k, w in zip("ABCDEF", [11, 24, 18, 10.9, 13, 120]): ws.column_dimensions[k].width = w
    ws.freeze_panes = "C2"
    _writeComments(ws, m)

    # ---- Data Quality
    ws = wb.create_sheet("Data Quality")
    _title(ws, "Data Quality", width=8)
    single = [s for s in m["students"] if not s["e2"]]
    r = 4
    for cc in range(2, 8): _hdr(ws.cell(r, cc), None, bg=_BLUE, sz=10)
    _hdr(ws.cell(r, 2), f"Single-marked students ({len(single)})", bg=_BLUE, sz=10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    r += 1
    for i, h in enumerate(["Student No", "Student", "Examiner", "Date", "Form %", ""]): _sub(ws.cell(r, 2 + i), h)
    r += 1
    for s in single or [None]:
        vals = [s["no"], s["name"], s["e1"]["assessor"], s["e1"]["date"].strftime("%d %b %Y"), round(s["e1"]["pct"], 1), ""] if s \
            else ["", "None — every student has two submitted forms", "", "", "", ""]
        for i, v in enumerate(vals): _cell(ws, r, 2 + i, v, align="left")
        r += 1
    r += 1
    for cc in range(2, 8): _hdr(ws.cell(r, cc), None, bg=_BLUE, sz=10)
    _hdr(ws.cell(r, 2), f"{cohort} roster students with no submitted form ({len(m['notPresented'])})", bg=_BLUE, sz=10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    r += 1
    for i, h in enumerate(["Student No", "Student", "Email", "", "", ""]): _sub(ws.cell(r, 2 + i), h)
    r += 1
    for v in m["notPresented"]:
        for i, x in enumerate([int(v["student_number"]), v["student_name"], v["student_email"], "", "", ""]): _cell(ws, r, 2 + i, x, align="left")
        r += 1
    for k, w in zip("BCDEFG", [11, 26, 40, 13, 10, 36]): ws.column_dimensions[k].width = w

    # ---- Raw Forms
    ws = wb.create_sheet("Raw Forms")
    heads = ["Form id", "Date", "Student No", "Student", "Subject", "Examiner"] + [f"{mc} label" for mc in mcs] + \
            [f"{mc} pts" for mc in mcs] + [f"Total /{4 * nItems}", "Form %", "Poor/Missing ratings"]
    for i, h in enumerate(heads): _hdr(ws.cell(1, 1 + i), h, bg=_NAVY, sz=9)
    ws.row_dimensions[1].height = 24
    for r in range(2, 2 + len(m["forms"])):
        for i in range(len(heads)):
            _cell(ws, r, 1 + i, None, numFmt="dd mmm yyyy" if i == 1 else ("0.0" if i == len(heads) - 2 else None),
                  align="left" if i in (3, 5) else "center")
    for i in range(len(heads)): ws.column_dimensions[_xlColLetter(1 + i)].width = 22 if i in (3, 5) else 11
    ws.freeze_panes = "E2"
    _writeRawForms(ws, m)

    os.makedirs(os.path.dirname(os.path.abspath(outPath)), exist_ok=True)
    wb.save(outPath)
    print(f"Built {outPath}: {len(m['students'])} students · {len(m['forms'])} forms · α {m['alpha']:.2f} · "
          f"mean {m['cohortMean']:.1f}% · excluded {len(m['excluded'])}")
    return outPath


# ================================================================ writers (shared by build + refresh)
def _findRow(ws, col, pred, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if pred(v): return r
    return None


def _studentResultsLayout(ws):
    """Locate columns by header text so user-deleted columns (Subject, Rank, …) are tolerated."""
    hdr = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(2, c).value
        if v is not None: hdr.setdefault(str(v).replace("\n", " ").strip(), []).append(c)
    e1Cols = hdr.get("E1", [])
    return hdr, e1Cols


def _writeStudentResults(ws, m):
    hdr, e1Cols = _studentResultsLayout(ws)
    mcs = m["mcs"]
    if len(e1Cols) != len(mcs):
        raise ValueError(f"Student Results: found {len(e1Cols)} E1 item columns, data has {len(mcs)} criteria — rebuild the workbook.")
    col = lambda name: hdr[name][0] if name in hdr else None
    order = _byFinal(m)
    meanRow = _findRow(ws, 2, lambda v: v == "Cohort mean", 3)
    dataRows = list(range(3, meanRow))
    # grow / shrink the student block (rows above the cohort-mean row)
    if len(order) > len(dataRows):
        extra = len(order) - len(dataRows)
        ws.insert_rows(meanRow, extra)
        for r in range(meanRow, meanRow + extra):
            for c in range(1, ws.max_column + 1):
                if ws.cell(meanRow - 1, c).has_style: ws.cell(r, c)._style = _copy.copy(ws.cell(meanRow - 1, c)._style)
        meanRow += extra
    elif len(order) < len(dataRows):
        ws.delete_rows(3 + len(order), len(dataRows) - len(order))
        meanRow -= len(dataRows) - len(order)
    dataRows = list(range(3, meanRow))
    nT = next((k for k in hdr if k.startswith("E1 Total")), None)
    nT2 = next((k for k in hdr if k.startswith("E2 Total")), None)
    for r, s in zip(dataRows, order):
        e1, e2 = s["e1"], s["e2"]
        for name, v in [("Student No", s["no"]), ("Student", s["name"]), ("Subject", s["subject"]),
                        ("Date", _dt.datetime.combine(s["date"], _dt.time())), ("Examiner 1", e1["assessor"]),
                        ("Examiner 2", e2["assessor"] if e2 else "—")]:
            if col(name): ws.cell(r, col(name)).value = v
        for i, mc in enumerate(mcs):
            c1 = e1Cols[i]
            _levelCell(ws.cell(r, c1), e1["pts"][mc])
            _levelCell(ws.cell(r, c1 + 1), e2["pts"][mc] if e2 else None)
        if nT: _restyle(ws.cell(r, col(nT)), e1["total"])
        if nT2: _restyle(ws.cell(r, col(nT2)), e2["total"] if e2 else None)
        if col("E1 %"): _pctCell(ws.cell(r, col("E1 %")), e1["pct"])
        if col("E2 %"): _pctCell(ws.cell(r, col("E2 %")), e2["pct"] if e2 else None)
        if col("Final %"): _pctCell(ws.cell(r, col("Final %")), s["meanPct"], True)
        g = s["gap"]
        if col("Examiner Gap (pts)"): _restyle(ws.cell(r, col("Examiner Gap (pts)")), g, "FFEB9C" if g is not None and g >= _DISAGREE_PCT else None)
        if col("Max Item Gap (lvls)"):
            ws.cell(r, col("Max Item Gap (lvls)")).value = max(abs(e1["pts"][mc] - e2["pts"][mc]) for mc in mcs) if e2 else None
        n = s["nPoor"]
        if col("Poor/Missing Ratings"): _restyle(ws.cell(r, col("Poor/Missing Ratings")), n, "FFC7CE" if n else None, "9C0006" if n else _INK, bool(n))
        if col("Rank"): ws.cell(r, col("Rank")).value = order.index(s) + 1
        if col("Flags"): _restyle(ws.cell(r, col("Flags")), "; ".join(s["flags"]) or None, "FFF2CC" if s["flags"] else None)
        if col("Examiner 1 Comments"): ws.cell(r, col("Examiner 1 Comments")).value = e1["comments"] or None
        if col("Examiner 2 Comments"): ws.cell(r, col("Examiner 2 Comments")).value = (e2["comments"] or None) if e2 else None
    for i, mc in enumerate(mcs):
        for k, key in enumerate(("e1", "e2")):
            v = [s[key]["pts"][mc] for s in m["students"] if s[key]]
            ws.cell(meanRow, e1Cols[i] + k).value = round(float(_np.mean(v)), 2) if v else None
    for name, key in (("E1 %", "e1"), ("E2 %", "e2")):
        v = [s[key]["pct"] for s in m["students"] if s[key]]
        if col(name): ws.cell(meanRow, col(name)).value = float(_np.mean(v)) if v else None
    if col("Final %"): ws.cell(meanRow, col("Final %")).value = m["cohortMean"]
    if ws.auto_filter.ref:
        ws.auto_filter.ref = re.sub(r"\d+$", str(dataRows[-1]), ws.auto_filter.ref)


def _writeCohortSummary(ws, m):
    kp = _kpis(m)
    bandMap = {lab: (a, b) for lab, a, b in _BANDS}
    means = [s["meanPct"] for s in m["students"]]
    itemBy = {it["mc"]: it for it in m["itemStats"]}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 2).value
        if k in kp: ws.cell(r, 3).value = kp[k]
        b = ws.cell(r, 5).value
        if b in bandMap:
            lo, hi = bandMap[b]
            ws.cell(r, 6).value = sum(lo <= v < hi for v in means)
        if isinstance(k, str) and k.split(" ·")[0] in itemBy:
            it = itemBy[k.split(" ·")[0]]
            for j, l in enumerate(_LEVELS):
                _restyle(ws.cell(r, 3 + j), round(100 * it["dist"][l] / it["n"], 1), _LEVEL_FILL[ORAL_LEVEL_POINTS[l]][0])
            _pctCell(ws.cell(r, 3 + len(_LEVELS)), it["pct"], True)


def _writeItemAnalysis(ws, m):
    if isinstance(ws["B3"].value, str):
        ws["B3"].value = re.sub(r"(student-level means\) = )[0-9.]+?(\.(\s|$))",
                                lambda x: f"{x.group(1)}{m['alpha']:.2f}{x.group(2)}", ws["B3"].value)
    hdrRow = _findRow(ws, 2, lambda v: v == "Item")
    hdr = {str(ws.cell(hdrRow, c).value).strip(): c for c in range(2, ws.max_column + 1) if ws.cell(hdrRow, c).value}
    itemBy = {it["mc"]: it for it in m["itemStats"]}
    r = hdrRow + 1
    while ws.cell(r, 2).value in itemBy:
        it = itemBy[ws.cell(r, 2).value]
        if hdr.get("Ratings"): ws.cell(r, hdr["Ratings"]).value = it["n"]
        for l in _LEVELS:
            if hdr.get(l): _restyle(ws.cell(r, hdr[l]), it["dist"][l], _LEVEL_FILL[ORAL_LEVEL_POINTS[l]][0] if it["dist"][l] else None)
        for name, key in (("Mean pts\n(/4)", "meanPts"), ("SD", "sd"), ("Exact\nagreement %", "exact"), ("Adjacent\nagreement %", "adj")):
            if hdr.get(name): ws.cell(r, hdr[name]).value = it[key]
        if hdr.get("Mean %"): _pctCell(ws.cell(r, hdr["Mean %"]), it["pct"])
        if hdr.get("Item–rest r"): _restyle(ws.cell(r, hdr["Item–rest r"]), it["ritc"], _ritcFill(it["ritc"]))
        if hdr.get("α if deleted"): _restyle(ws.cell(r, hdr["α if deleted"]), it["alphaDel"], "FFEB9C" if it["alphaDel"] > m["alpha"] else None)
        r += 1


def _writeExaminerAnalysis(ws, m):
    if isinstance(ws["B3"].value, str):
        ws["B3"].value = re.sub(r"\(\d+(\.\d+)?%\)", f"({m['formMean']:.1f}%)", ws["B3"].value, count=1)
    # block 1: examiners
    h1 = _findRow(ws, 2, lambda v: v == "Examiner")
    rows = []
    r = h1 + 1
    while ws.cell(r, 2).value not in (None, ""): rows.append(r); r += 1
    if len(rows) != len(m["examiners"]):
        raise ValueError(f"Examiner Analysis: sheet has {len(rows)} examiner rows, data has {len(m['examiners'])} — rebuild the workbook.")
    nLv = len(_LEVELS)
    for r, e in zip(rows, m["examiners"]):
        for k, v in enumerate([e["name"], e["n"], e["mean"], e["median"], e["sd"]]): ws.cell(r, 2 + k).value = v
        _restyle(ws.cell(r, 7), e["delta"], "C6EFCE" if e["delta"] >= 5 else "FFC7CE" if e["delta"] <= -5 else None)
        for j, l in enumerate(_LEVELS): ws.cell(r, 8 + j).value = e["mix"][l]
        note = ws.cell(r, 8 + nLv)
        if e["n"] < 5: note.value = "Small n — interpret with caution"
        elif note.value == "Small n — interpret with caution": note.value = None
    # block 2: pairs
    h2 = _findRow(ws, 2, lambda v: v == "Pair (A | B)")
    rows = []
    r = h2 + 1
    while isinstance(ws.cell(r, 2).value, str) and " | " in ws.cell(r, 2).value: rows.append(r); r += 1
    if len(rows) != len(m["pairs"]):
        raise ValueError(f"Examiner Analysis: sheet has {len(rows)} pair rows, data has {len(m['pairs'])} — rebuild the workbook.")
    for r, p in zip(rows, m["pairs"]):
        for k, key in enumerate(["label", "n", "aMean", "bMean", "diff", "absGap"]): ws.cell(r, 2 + k).value = p[key]
        _restyle(ws.cell(r, 8), p["p"], "FFC7CE" if p["p"] is not None and p["p"] < 0.05 else None)
        ws.cell(r, 9).value = p["exact"]; ws.cell(r, 10).value = p["adj"]
        _restyle(ws.cell(r, 11), p["kappa"], _kappaFill(p["kappa"]))
        ws.cell(r, 12).value = p["nGap"]
    # block 3: examiner × criterion
    h3 = _findRow(ws, 2, lambda v: v == "Examiner", h2)
    exBy = {e["name"]: e for e in m["examiners"]}
    r = h3 + 1
    while ws.cell(r, 2).value in exBy:
        e = exBy[ws.cell(r, 2).value]
        for i, mc in enumerate(m["mcs"]): ws.cell(r, 3 + i).value = e["itemMean"][mc]
        ws.cell(r, 3 + len(m["mcs"])).value = e["n"]
        r += 1
    if ws.cell(r, 2).value == "All forms":
        for i, it in enumerate(m["itemStats"]): ws.cell(r, 3 + i).value = it["meanPts"]


def _writeComments(ws, m):
    rows = [(s, f) for s in _byFinal(m) for f in s["forms"]]
    last0 = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value is not None or r == 1)
    for r in range(last0 + 1, 2 + len(rows)):
        for c in range(1, 7):
            if ws.cell(last0, c).has_style: ws.cell(r, c)._style = _copy.copy(ws.cell(last0, c)._style)
        if ws.row_dimensions[last0].height: ws.row_dimensions[r].height = ws.row_dimensions[last0].height
    for r in range(2 + len(rows), last0 + 1):
        for c in range(1, 7): ws.cell(r, c).value = None
    for r, (s, f) in enumerate(rows, start=2):
        ws.cell(r, 1).value = s["no"]; ws.cell(r, 2).value = s["name"]; ws.cell(r, 3).value = f["assessor"]
        _pctCell(ws.cell(r, 4), s["meanPct"]); _pctCell(ws.cell(r, 5), f["pct"])
        ws.cell(r, 6).value = f["comments"] or "—"
    last = 1 + len(rows)
    ws.auto_filter.ref = f"A1:F{last}"
    ss = ws.auto_filter.sortState
    if ss is not None:
        ss.ref = f"A2:F{last}"
        for sc in ss.sortCondition: sc.ref = re.sub(r"\d+$", str(last), sc.ref)
    else:
        ws.auto_filter.add_sort_condition(f"D2:D{last}", descending=True)


def _writeDataQuality(ws, m):
    h1 = _findRow(ws, 2, lambda v: isinstance(v, str) and v.startswith("Single-marked"))
    h2 = _findRow(ws, 2, lambda v: isinstance(v, str) and "roster students with no submitted form" in v)
    single = [s for s in m["students"] if not s["e2"]]
    if h1:
        slots = list(range(h1 + 2, (h2 - 1) if h2 else ws.max_row + 1))
        if len(single) > len(slots):
            raise ValueError("Data Quality: more single-marked students than rows available — rebuild the workbook.")
        ws.cell(h1, 2).value = f"Single-marked students ({len(single)})"
        for i, r in enumerate(slots):
            for c in range(2, 8): ws.cell(r, c).value = None
            if i < len(single):
                s = single[i]
                for c, v in zip(range(2, 7), [s["no"], s["name"], s["e1"]["assessor"], s["e1"]["date"].strftime("%d %b %Y"), round(s["e1"]["pct"], 1)]):
                    ws.cell(r, c).value = v
        if not single: ws.cell(slots[0], 3).value = "None — every student has two submitted forms"
    if h2:
        ws.cell(h2, 2).value = re.sub(r"\(\d+\)", f"({len(m['notPresented'])})", ws.cell(h2, 2).value)
        start = h2 + 2
        old = [r for r in range(start, ws.max_row + 1) if ws.cell(r, 2).value is not None]
        for i, v in enumerate(m["notPresented"]):
            r = start + i
            if old and r > old[-1]:
                for c in range(2, 8):
                    if ws.cell(old[-1], c).has_style: ws.cell(r, c)._style = _copy.copy(ws.cell(old[-1], c)._style)
            for c, x in zip(range(2, 5), [int(v["student_number"]), v["student_name"], v["student_email"]]): ws.cell(r, c).value = x
        for r in range(start + len(m["notPresented"]), (old[-1] + 1) if old else start):
            for c in range(2, 8): ws.cell(r, c).value = None


def _writeRawForms(ws, m):
    mcs = m["mcs"]
    fs = sorted(m["forms"], key=lambda f: (_surnameKey(f["student"]), f["id"]))
    last0 = max(ws.max_row, 2)
    for r in range(last0 + 1, 2 + len(fs)):
        for c in range(1, ws.max_column + 1):
            if ws.cell(last0, c).has_style: ws.cell(r, c)._style = _copy.copy(ws.cell(last0, c)._style)
    for r in range(2 + len(fs), last0 + 1):
        for c in range(1, ws.max_column + 1): ws.cell(r, c).value = None
    n = len(mcs)
    for r, f in enumerate(fs, start=2):
        vals = [f["id"], _dt.datetime.combine(f["date"], _dt.time()), f["studentNo"], f["student"], f["subject"], f["assessor"]] + \
               [f["labels"][mc] for mc in mcs]
        for i, v in enumerate(vals): ws.cell(r, 1 + i).value = v
        for i, mc in enumerate(mcs): _levelCell(ws.cell(r, 7 + n + i), f["pts"][mc])
        ws.cell(r, 7 + 2 * n).value = f["total"]; ws.cell(r, 8 + 2 * n).value = f["pct"]; ws.cell(r, 9 + 2 * n).value = f["nPoor"]
    ws.auto_filter.ref = f"A1:{_xlColLetter(9 + 2 * n)}{1 + len(fs)}"


def _writeDescription(ws, m):
    vals = _descriptionValues(m)
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 2).value
        if k in vals: ws.cell(r, 3).value = vals[k]


# ================================================================ refresh (keeps manual edits)
def _restoreCharts(srcPath, dstPath):
    """openpyxl re-serialises charts and can lose formatting; put the original chart XML back.
    Charts are matched on their cell references, so only charts that still point at the same ranges are restored."""
    refs = lambda xml: tuple(sorted(set(re.findall(r"<(?:c:)?f>(.*?)</(?:c:)?f>", xml))))
    with _zipfile.ZipFile(srcPath) as zs:
        srcCharts = {refs(zs.read(n).decode("utf-8")): zs.read(n) for n in zs.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n)}
    tmp = dstPath + ".tmp"
    restored = 0
    with _zipfile.ZipFile(dstPath) as zd, _zipfile.ZipFile(tmp, "w", _zipfile.ZIP_DEFLATED) as zo:
        for item in zd.infolist():
            data = zd.read(item.filename)
            if re.match(r"xl/charts/chart\d+\.xml$", item.filename):
                key = refs(data.decode("utf-8"))
                if key in srcCharts: data = srcCharts[key]; restored += 1
            zo.writestr(item, data)
    os.replace(tmp, dstPath)
    return restored


def refreshOralPresentationReport(jsonPath, workbookPath, rosterPath="studentEmailList.csv", checklistKey="BOH3-OHTR",
                                  cohort="BOH3", outPath=None, backupDir="_bak"):
    """Re-score a new JSON pull and write the numbers into an EXISTING report workbook.

    Keeps everything the user changed by hand: deleted columns / sections / tables, column widths, row heights,
    unwrapped text, sheet order, sorts and charts. Only values and value-dependent colours are rewritten.
    Columns are found by header text, so removed columns are simply skipped. Rows are added/removed only in the
    per-student / per-form lists (Student Results, Comments, Raw Forms, Data Quality lists).
    If the examiner or examiner-pair count changes, it stops and asks for a full rebuild.
    The original is copied to `backupDir` first (set backupDir=None to skip). outPath defaults to workbookPath.
    """
    m = loadOralPresentation(jsonPath, rosterPath, checklistKey, cohort)
    outPath = outPath or workbookPath
    if backupDir and os.path.abspath(outPath) == os.path.abspath(workbookPath):
        os.makedirs(backupDir, exist_ok=True)
        stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        _shutil.copy2(workbookPath, os.path.join(backupDir, f"{os.path.basename(workbookPath)}.bak_{stamp}"))
    wb = _openpyxl.load_workbook(workbookPath)
    writers = {"Description": _writeDescription, "Student Results": _writeStudentResults,
               "Cohort Summary": _writeCohortSummary, "Item Analysis": _writeItemAnalysis,
               "Examiner Analysis": _writeExaminerAnalysis, "Comments": _writeComments,
               "Data Quality": _writeDataQuality, "Raw Forms": _writeRawForms}
    for name, fn in writers.items():
        if name in wb.sheetnames: fn(wb[name], m)
    src = workbookPath
    if os.path.abspath(outPath) == os.path.abspath(workbookPath):
        src = outPath + ".src.tmp"
        _shutil.copy2(workbookPath, src)
    wb.save(outPath)
    nCharts = _restoreCharts(src, outPath)
    if src != workbookPath: os.remove(src)
    print(f"Refreshed {outPath}: {len(m['students'])} students · {len(m['forms'])} forms · α {m['alpha']:.2f} · "
          f"mean {m['cohortMean']:.1f}% · single-marked {sum(1 for s in m['students'] if not s['e2'])} · "
          f"excluded {len(m['excluded'])} · charts kept {nCharts}")
    return outPath
