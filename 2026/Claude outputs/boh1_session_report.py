"""
boh1_session_report.py — BOH1 single-session + item-history Excel reports from the CAF API dump.

Scoring: five-level MCs score 0/2/3/4/5 points (Done well 5, Done 4, Mostly done 3, Sometimes done 2, Not done 0);
Score = sum of points, Max = 5 x applicable MCs, % = Score / Max. O6 (N/A) is excluded from score and max.
checkFiveLevels() verifies every MC of the item uses exactly the five levels + N/A before scoring.
Only assessor-submitted forms are scored; unsubmitted forms are listed on the "Pending & Absent" sheet.

Usage (from project root):
    from importlib import import_module  # or run:  python "Claude outputs/boh1_session_report.py"
    buildSessionReports(sessionDate="2026-09-29", itemCode="532")
"""
import json, re, statistics
from collections import Counter, defaultdict
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule, ColorScaleRule
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

# ── Scoring (same as boh1_utils) ──────────────────────────────────────────────
POINTS = {"O1": 5, "O2": 4, "O3": 3, "O4": 2, "O5": 0}          # 0/2/3/4/5 points per five-level MC
MAX_POINTS = 5
SCORE_MAP = {k: v / MAX_POINTS for k, v in POINTS.items()}      # fraction of max, used for all % means
FIVE_LEVELS = {"O1": "Done well", "O2": "Done", "O3": "Mostly done", "O4": "Sometimes done", "O5": "Not done", "O6": "Not applicable"}
OPT_SHORT = {"O1": "DW", "O2": "D", "O3": "MD", "O4": "SD", "O5": "ND", "O6": "N/A"}
OPT_NUM = POINTS   # MC Matrix shows the 0/2/3/4/5 points; O6 (N/A) left blank
OPT_LABEL = {"O1": "Done well", "O2": "Done", "O3": "Mostly done", "O4": "Sometimes done", "O5": "Not done", "O6": "Not applicable"}
GR_LABEL = {1: "Unsatisfactory", 2: "Borderline", 3: "Satisfactory", 4: "Good", 5: "Excellent"}
PR_LABEL = {1: "L1 – Not ready", 2: "L2 – Continuous direct supv.", 3: "L3 – Periodic direct supv.", 4: "L4 – Indirect supv."}

# ── Palette: UniMelb blue + orange, colour-blind-safe diverging ───────────────
NAVY, BLUE, LBLUE, PALE = "094183", "4A7EBB", "D6E4F5", "EEF3FA"
ORANGE, LORANGE, NEUTRAL, GREY, DGREY = "C75B12", "F8D3B0", "F2F2F2", "D9D9D9", "595959"
OPT_FILL = {"O1": (NAVY, "FFFFFF"), "O2": (LBLUE, "000000"), "O3": (NEUTRAL, "000000"),
            "O4": (LORANGE, "000000"), "O5": (ORANGE, "FFFFFF"), "O6": ("FFFFFF", "A6A6A6")}
GR_FILL = {1: (ORANGE, "FFFFFF"), 2: (LORANGE, "000000"), 3: (NEUTRAL, "000000"), 4: (LBLUE, "000000"), 5: (NAVY, "FFFFFF")}
PR_FILL = {1: (ORANGE, "FFFFFF"), 2: (NEUTRAL, "000000"), 3: (LBLUE, "000000"), 4: (NAVY, "FFFFFF")}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ═════════════════════════════════════════════════════════════════════════════
# Data extraction
# ═════════════════════════════════════════════════════════════════════════════
def scoreMcs(mcDict):
    """MC dict ({'MC1': {'key':'O2'}} or {'MC1':'O2'}) -> (score, max, pct, {mc: optKey})."""
    opts = {}
    for k, v in (mcDict or {}).items():
        if not k.startswith("MC"):
            continue
        opts[k] = v.get("key") if isinstance(v, dict) else v
    applicable = [POINTS[o] for o in opts.values() if o in POINTS]
    score = sum(applicable)
    mx = MAX_POINTS * len(applicable)
    return score, mx, (round(score / mx * 100, 1) if mx else None), opts


def sectionPct(opts, sectionMap, section):
    """% score over the MCs belonging to one checklist section."""
    vals = [SCORE_MAP[o] for mc, o in opts.items() if sectionMap.get(mc) == section and o in SCORE_MAP]
    return round(sum(vals) / len(vals) * 100, 1) if vals else None


def checklistMeta(form, itemCode):
    """Item name, MC descriptions and MC->section map from the form's assessor_config."""
    cfg = (((form.get("assessor_config") or {}).get("checklists") or {}).get("selected") or {}).get(itemCode) or {}
    fields = cfg.get("fields") or {}
    headers = (cfg.get("extra_config") or {}).get("headers") or {}
    secMap, current = {}, "General"
    for mc in sorted(fields, key=lambda m: int(re.sub(r"\D", "", m) or 0)):
        if mc in headers and headers[mc]:
            current = headers[mc][0].get("title", current)
        secMap[mc] = current
    return cfg.get("name", itemCode), fields, secMap


def checkFiveLevels(jsonPath, itemCode, cohort="BOH1"):
    """Confirm every form's checklist options for itemCode are exactly the five levels + N/A (assessor and student),
    and no MC carries an option key outside O1-O6. Raises if not, since 0/2/3/4/5 scoring would be wrong."""
    with open(jsonPath, encoding="utf-8") as f:
        records = json.load(f)
    bad, nForms = [], 0
    for r in records:
        if r.get("cohort") != cohort:
            continue
        for form in r.get("forms") or []:
            cfg = (((form.get("assessor_config") or {}).get("checklists") or {}).get("selected") or {}).get(itemCode)
            if not cfg:
                continue
            nForms += 1
            for side, opts in ((cfg.get("extra_config") or {}).get("options") or {}).items():
                if opts != FIVE_LEVELS:
                    bad.append((form.get("id"), side, opts))
            for mc, v in (((form.get("assessor_data") or {}).get("checklists") or {}).get(itemCode) or {}).items():
                key = v.get("key") if isinstance(v, dict) else v
                if key not in FIVE_LEVELS:
                    bad.append((form.get("id"), mc, key))
    if bad:
        raise ValueError(f"Item {itemCode}: {len(bad)} non-five-level option sets/values, e.g. {bad[:3]}")
    print(f"✓ Item {itemCode}: all MCs five-level + N/A across {nForms} forms")


def tidyName(n):
    """Title-case names typed all-lower or ALL-UPPER; leave mixed-case (e.g. McKinlay) alone."""
    n = (n or "").strip()
    return n.title() if n and (n.islower() or n.isupper()) else n


def scaleKey(ad, name):
    v = ((ad.get("scales") or {}).get(name) or {}).get("key")
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def loadItemRows(jsonPath, itemCode, cohort="BOH1"):
    """One row per (record, form) that contains itemCode. Returns (rows, meta, lastUpdated)."""
    with open(jsonPath, encoding="utf-8") as f:
        records = json.load(f)
    rows, meta, lastUpdated = [], None, None
    for r in records:
        if r.get("cohort") != cohort:
            continue
        stu = r.get("student") or {}
        for form in r.get("forms") or []:
            ctxCodes = [c.get("code") for c in (form.get("form_context") or {}).get("checklists", [])]
            ad = form.get("assessor_data") or {}
            sd = form.get("student_data") or {}
            if itemCode not in ctxCodes and itemCode not in (ad.get("checklists") or {}):
                continue
            name, fields, secMap = checklistMeta(form, itemCode)
            if fields and (meta is None or len(fields) >= len(meta[1])):
                meta = (name, fields, secMap)
            score, mx, pct, opts = scoreMcs((ad.get("checklists") or {}).get(itemCode))
            sScore, sMax, sPct, sOpts = scoreMcs((sd.get("checklists") or {}).get(itemCode))
            radio = ad.get("radio") or {}
            texts = ad.get("texts") or {}
            incident = str(radio.get("clinical-incident-occurred", "")).lower() == "yes" or bool((ad.get("multi-select") or {}).get("clinical-incident"))
            upd = form.get("updated_at")
            if upd and (lastUpdated is None or upd > lastUpdated):
                lastUpdated = upd
            rows.append({
                "Record ID": r.get("id"), "Form ID": form.get("id"),
                "Date": r["datetime"][:10],   # API datetime is already Melbourne-local (+10/+11)
                "Type": r.get("type"),
                "Student ID": str(stu.get("student_number", "")),
                "First Name": tidyName(stu.get("first_name")), "Last Name": tidyName(stu.get("last_name")),
                "Student Name": f"{tidyName(stu.get('first_name'))} {tidyName(stu.get('last_name'))}".strip(),
                "Assessor": form.get("assessor_name") or "—",
                "Clinic": (form.get("form_context") or {}).get("clinic_type", ""),
                "Tooth/Area": (form.get("form_context") or {}).get("teeth_quadrant", ""),
                "Assessor Submitted": bool(form.get("submitted_by_assessor")),
                "Student Submitted": bool(form.get("submitted_by_student")),
                "Score": score, "Max Score": mx, "% Score": pct, "opts": opts,
                "Self Score": sScore if sMax else None, "Self Max": sMax or None, "Self %": sPct, "selfOpts": sOpts,
                "Global Rating": scaleKey(ad, "scale-global-rating"),
                "Practice Readiness": scaleKey(ad, "scale-practice-readiness"),
                "Professionalism": scaleKey(ad, "scale-professionalism"),
                "Position & Ergonomics": scaleKey(ad, "scale-position-ergonomics"),
                "Preparedness": scaleKey(ad, "scale-preparedness"),
                "Clinical Incident": "Yes" if incident else "",
                "Incident Details": texts.get("clinical-incident-additional-details", ""),
                "Assessor Comments": texts.get("reflection", ""),
                "Student Reflection": ((sd.get("texts") or {}).get("reflection", "")),
                "Updated": upd,
            })
    return rows, meta, lastUpdated


# ═════════════════════════════════════════════════════════════════════════════
# Styling helpers
# ═════════════════════════════════════════════════════════════════════════════
def clean(v):
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def put(ws, row, col, value, bold=False, fill=None, color="000000", align="left", wrap=False, border=True, size=10, numFmt=None, italic=False):
    c = ws.cell(row=row, column=col, value=clean(value))
    if isinstance(c.value, str) and c.value.startswith("="):
        c.data_type = "s"   # never let a comment be parsed as a formula
    c.font = Font(name="Calibri", bold=bold, color=color, size=size, italic=italic)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
    if border:
        c.border = BORDER
    if numFmt:
        c.number_format = numFmt
    return c


def banner(ws, title, subtitle, ncols):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 2))
    put(ws, 1, 1, title, bold=True, fill=NAVY, color="FFFFFF", size=14, border=False).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ncols, 2))
    put(ws, 2, 1, subtitle, italic=True, color=DGREY, fill=PALE, border=False, size=9)
    return 4


def header(ws, row, cols, fill=NAVY, height=32):
    for j, name in enumerate(cols, 1):
        put(ws, row, j, name, bold=True, fill=fill, color="FFFFFF", align="center", wrap=True)
    ws.row_dimensions[row].height = height
    return row + 1


def widths(ws, spec):
    for j, w in enumerate(spec, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def paint(cell, pair):
    if pair:
        cell.fill = PatternFill("solid", fgColor=pair[0])
        cell.font = Font(name="Calibri", size=10, color=pair[1], bold=pair[1] == "FFFFFF")


def pctBar(ws, rng, color=BLUE):
    ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=color, showValue=True))


def pctScale(ws, rng):
    # orange (low) -> white -> blue (high); colour-blind-safe
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=40, start_color="F4A261",
                                                      mid_type="num", mid_value=70, mid_color="FFFFFF",
                                                      end_type="num", end_value=100, end_color="6A9BD1"))


def finish(ws, freeze, lastCol, headerRow, landscape=True):
    for rr in (range(headerRow + 1, ws.max_row + 1) if headerRow else []):
        ws.row_dimensions[rr].height = 15   # thin rows; long text stays on one line (unwrapped)
    ws.freeze_panes = freeze
    if headerRow and ws.max_row > headerRow:
        ws.auto_filter.ref = f"A{headerRow}:{get_column_letter(lastCol)}{ws.max_row}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{headerRow}:{headerRow}" if headerRow else None
    ws.sheet_view.zoomScale = 90


def flagsFor(r):
    f = []
    if r["Global Rating"] is not None and r["Global Rating"] <= 2:
        f.append(f"GR {GR_LABEL[r['Global Rating']]}")
    if r["Practice Readiness"] == 1:
        f.append("PR L1")
    if r["% Score"] is not None and r["% Score"] < 60:
        f.append("% < 60")
    nd = [mc for mc, o in r["opts"].items() if o == "O5"]
    if nd:
        f.append("Not done: " + ", ".join(nd))
    if r["Clinical Incident"]:
        f.append("Clinical incident")
    if r["Professionalism"] == 1:
        f.append("Professionalism L1")
    if r["Preparedness"] == 1:
        f.append("Not prepared")
    return "; ".join(f)


def mcSort(mcs):
    return sorted(mcs, key=lambda m: int(re.sub(r"\D", "", m) or 0))


# ═════════════════════════════════════════════════════════════════════════════
# Session workbook
# ═════════════════════════════════════════════════════════════════════════════
def buildSessionWorkbook(rows, meta, sessionDate, itemCode, outPath, lastUpdated, cohortStudents):
    itemName, mcDesc, secMap = meta
    sections = list(dict.fromkeys(secMap[m] for m in mcSort(mcDesc)))
    mcs = mcSort(mcDesc)
    dayRows = [r for r in rows if r["Date"] == sessionDate]
    done = sorted([r for r in dayRows if r["Assessor Submitted"]], key=lambda r: (r["Last Name"].lower(), r["First Name"].lower()))
    pending = [r for r in dayRows if not r["Assessor Submitted"]]
    for r in done:
        r["Flags"] = flagsFor(r)
        for s in sections:
            r[f"{s} %"] = sectionPct(r["opts"], secMap, s)
    dateTxt = pd.to_datetime(sessionDate).strftime("%a %d %b %Y")
    sub = (f"Item {itemCode} — {itemName}  |  {dateTxt}  |  Source: DASH CAF API dump (temp 2026 caf.json), "
           f"last form update {str(lastUpdated)[:16].replace('T', ' ')}  |  Score = 0/2/3/4/5 points per MC, % = points / (5 × applicable MCs); N/A excluded; assessor-submitted forms only")
    wb = Workbook()

    # ── 1. Overview ───────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Overview"
    row = banner(ws, f"BOH1 Item {itemCode} Session Report — {dateTxt}", sub, 10)
    pcts = [r["% Score"] for r in done if r["% Score"] is not None]
    grs = [r["Global Rating"] for r in done if r["Global Rating"] is not None]
    kpis = [("Forms on session", len(dayRows)), ("Assessor-submitted", len(done)), ("Pending assessor", len(pending)),
            ("Mean % score", round(statistics.mean(pcts), 1) if pcts else None),
            ("Median % score", round(statistics.median(pcts), 1) if pcts else None),
            ("Range %", f"{min(pcts):.1f} – {max(pcts):.1f}" if pcts else None),
            ("Mean GR (1–5)", round(statistics.mean(grs), 2) if grs else None),
            ("GR Borderline / Unsat.", sum(1 for g in grs if g <= 2)),
            ("PR Level 1", sum(1 for r in done if r["Practice Readiness"] == 1)),
            ("Clinical incidents", sum(1 for r in done if r["Clinical Incident"]))]
    for j, (lab, val) in enumerate(kpis, 1):
        put(ws, row, j, lab, bold=True, fill=LBLUE, color=NAVY, align="center", wrap=True, size=9)
        c = put(ws, row + 1, j, val, bold=True, align="center", size=16)
        if lab in ("GR Borderline / Unsat.", "PR Level 1", "Clinical incidents", "Pending assessor") and val:
            c.font = Font(name="Calibri", bold=True, size=16, color=ORANGE)
    ws.row_dimensions[row].height = 30
    ws.row_dimensions[row + 1].height = 30
    row += 3

    # GR / PR distributions side by side
    put(ws, row, 1, "Global Rating distribution", bold=True, color=NAVY, border=False, size=11)
    put(ws, row, 6, "Practice Readiness distribution", bold=True, color=NAVY, border=False, size=11)
    row += 1
    for j, h in enumerate(["GR", "Label", "n", "%"], 1):
        put(ws, row, j, h, bold=True, fill=NAVY, color="FFFFFF", align="center")
    for j, h in enumerate(["PR", "Label", "n", "%"], 6):
        put(ws, row, j, h, bold=True, fill=NAVY, color="FFFFFF", align="center")
    grC = Counter(grs)
    prC = Counter(r["Practice Readiness"] for r in done if r["Practice Readiness"] is not None)
    nPr = sum(prC.values())
    start = row + 1
    for i, g in enumerate(range(1, 6)):
        rr = start + i
        paint(put(ws, rr, 1, g, align="center"), GR_FILL[g])
        put(ws, rr, 2, GR_LABEL[g]); put(ws, rr, 3, grC.get(g, 0), align="center")
        put(ws, rr, 4, round(grC.get(g, 0) / len(grs) * 100, 1) if grs else None, align="center", numFmt="0.0")
    for i, p in enumerate(range(1, 5)):
        rr = start + i
        paint(put(ws, rr, 6, p, align="center"), PR_FILL[p])
        put(ws, rr, 7, PR_LABEL[p]); put(ws, rr, 8, prC.get(p, 0), align="center")
        put(ws, rr, 9, round(prC.get(p, 0) / nPr * 100, 1) if nPr else None, align="center", numFmt="0.0")
    pctBar(ws, f"D{start}:D{start + 4}"); pctBar(ws, f"I{start}:I{start + 3}")
    row = start + 6

    # Section means
    put(ws, row, 1, "Checklist section performance (mean %)", bold=True, color=NAVY, border=False, size=11)
    row += 1
    for j, h in enumerate(["Section", "MCs", "Mean %", "Min %", "Students < 60%"], 1):
        put(ws, row, j, h, bold=True, fill=NAVY, color="FFFFFF", align="center")
    s0 = row + 1
    for i, s in enumerate(sections):
        vals = [r[f"{s} %"] for r in done if r.get(f"{s} %") is not None]
        rr = s0 + i
        put(ws, rr, 1, s, bold=True)
        put(ws, rr, 2, ", ".join(m for m in mcs if secMap[m] == s), size=9)
        put(ws, rr, 3, round(statistics.mean(vals), 1) if vals else None, align="center", numFmt="0.0")
        put(ws, rr, 4, min(vals) if vals else None, align="center", numFmt="0.0")
        put(ws, rr, 5, sum(1 for v in vals if v < 60), align="center")
    pctBar(ws, f"C{s0}:C{s0 + len(sections) - 1}")
    row = s0 + len(sections) + 1

    # Flagged students
    flagged = [r for r in done if r["Flags"]]
    flagged.sort(key=lambda r: (r["% Score"] if r["% Score"] is not None else 999))
    put(ws, row, 1, f"Students needing attention ({len(flagged)})  —  GR ≤ 2, PR L1, % < 60, any 'Not done', incident, Prof./Prep. L1",
        bold=True, color=ORANGE, border=False, size=11)
    row += 1
    fcols = ["Student Name", "Student ID", "Assessor", "% Score", "GR", "PR", "Flags"]
    for j, h in enumerate(fcols, 1):
        put(ws, row, j, h, bold=True, fill=ORANGE, color="FFFFFF", align="center")
    ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=10)
    f0 = row + 1
    for i, r in enumerate(flagged):
        rr = f0 + i
        put(ws, rr, 1, r["Student Name"]); put(ws, rr, 2, r["Student ID"], align="center")
        put(ws, rr, 3, r["Assessor"]); put(ws, rr, 4, r["% Score"], align="center", numFmt="0.0")
        paint(put(ws, rr, 5, r["Global Rating"], align="center"), GR_FILL.get(r["Global Rating"]))
        paint(put(ws, rr, 6, r["Practice Readiness"], align="center"), PR_FILL.get(r["Practice Readiness"]))
        ws.merge_cells(start_row=rr, start_column=7, end_row=rr, end_column=10)
        put(ws, rr, 7, r["Flags"], size=9)
    if flagged:
        pctBar(ws, f"D{f0}:D{f0 + len(flagged) - 1}", color=ORANGE)
    for rr in range(start, ws.max_row + 1):
        ws.row_dimensions[rr].height = 15   # thin rows below the KPI tiles
    widths(ws, [24, 26, 20, 13, 13, 15, 28, 13, 13, 16])
    finish(ws, None, 10, None, landscape=False)

    # ── 2. Results ────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Session Results")
    cols = (["Student Name", "Student ID", "Assessor", "Clinic", "Tooth/Area", "Score", "Max", "% Score"]
            + [f"{s} %" for s in sections]
            + ["GR", "GR Label", "PR", "Prof.", "Pos. & Ergo.", "Prepared", "Incident", "Self %", "Self − Assessor",
               "Flags", "Assessor Comments", "Incident Details", "Student Reflection"])
    row = banner(ws, f"BOH1 Item {itemCode} — Session Results ({dateTxt})", sub, len(cols))
    hr = row
    row = header(ws, row, cols)
    for i, r in enumerate(done):
        gap = round(r["Self %"] - r["% Score"], 1) if r["Self %"] is not None and r["% Score"] is not None else None
        vals = ([r["Student Name"], r["Student ID"], r["Assessor"], r["Clinic"], r["Tooth/Area"], r["Score"], r["Max Score"], r["% Score"]]
                + [r.get(f"{s} %") for s in sections]
                + [r["Global Rating"], GR_LABEL.get(r["Global Rating"], ""), r["Practice Readiness"], r["Professionalism"],
                   r["Position & Ergonomics"], r["Preparedness"], r["Clinical Incident"], r["Self %"], gap,
                   r["Flags"], r["Assessor Comments"], r["Incident Details"], r["Student Reflection"]])
        band = PALE if i % 2 else None
        for j, v in enumerate(vals, 1):
            name = cols[j - 1]
            longTxt = name in ("Assessor Comments", "Incident Details", "Student Reflection", "Flags")
            numeric = isinstance(v, (int, float)) and not isinstance(v, bool)
            c = put(ws, row, j, v, fill=band, wrap=False, size=9 if longTxt else 10,
                    align="center" if (numeric or name in ("Student ID", "Clinic", "Incident")) else "left",
                    numFmt="0.0" if name.endswith("%") or name == "Self − Assessor" else None)
            if name == "GR":
                paint(c, GR_FILL.get(v))
            elif name == "PR":
                paint(c, PR_FILL.get(v))
            elif name in ("Prof.", "Pos. & Ergo.", "Prepared") and v == 1:
                paint(c, (ORANGE, "FFFFFF"))
            elif name == "Incident" and v:
                paint(c, (ORANGE, "FFFFFF"))
            elif name == "Self − Assessor" and v is not None and abs(v) >= 15:
                c.font = Font(name="Calibri", bold=True, color=ORANGE if v > 0 else BLUE, size=10)
            elif name == "Flags" and v:
                c.font = Font(name="Calibri", color=ORANGE, size=9, bold=True)
        row += 1
    last = row - 1
    pc = cols.index("% Score") + 1
    pctBar(ws, f"{get_column_letter(pc)}{hr + 1}:{get_column_letter(pc)}{last}")
    for s in sections:
        cl = get_column_letter(cols.index(f"{s} %") + 1)
        pctScale(ws, f"{cl}{hr + 1}:{cl}{last}")
    widthMap = {"Student Name": 24, "Student ID": 11, "Assessor": 20, "Clinic": 7, "Tooth/Area": 12, "Score": 7, "Max": 6,
                "% Score": 12, "GR": 5, "GR Label": 13, "PR": 5, "Prof.": 6, "Pos. & Ergo.": 7, "Prepared": 8, "Incident": 8,
                "Self %": 8, "Self − Assessor": 9, "Flags": 26, "Assessor Comments": 60, "Incident Details": 30, "Student Reflection": 60}
    widths(ws, [widthMap.get(c, 11) for c in cols])
    finish(ws, f"B{hr + 1}", len(cols), hr)

    # ── 3. MC Matrix ──────────────────────────────────────────────────────────
    ws = wb.create_sheet("MC Matrix")
    cols = ["Student Name", "Assessor"] + mcs + ["% Score", "GR"]
    row = banner(ws, f"Item {itemCode} — Assessor rating per checklist item (MC)",
                 "Points: 5 Done well · 4 Done · 3 Mostly done · 2 Sometimes done · 0 Not done · blank = N/A  |  hover an MC header for its description",
                 len(cols))
    # section group row
    put(ws, row, 1, "", fill=NAVY); put(ws, row, 2, "", fill=NAVY)
    j = 3
    for s in sections:
        span = [m for m in mcs if secMap[m] == s]
        ws.merge_cells(start_row=row, start_column=j, end_row=row, end_column=j + len(span) - 1)
        put(ws, row, j, s, bold=True, fill=BLUE, color="FFFFFF", align="center")
        for k in range(j + 1, j + len(span)):
            ws.cell(row=row, column=k).border = BORDER
        j += len(span)
    row += 1
    hr = row
    header(ws, row, cols, height=20)
    for k, mc in enumerate(mcs, 3):
        ws.cell(row=row, column=k).comment = Comment(mcDesc.get(mc, ""), "Checklist")
    row += 1
    for r in done:
        put(ws, row, 1, r["Student Name"]); put(ws, row, 2, r["Assessor"], size=9)
        for k, mc in enumerate(mcs, 3):
            o = r["opts"].get(mc)
            paint(put(ws, row, k, OPT_NUM.get(o), align="center", size=9), OPT_FILL.get(o))
        put(ws, row, len(cols) - 1, r["% Score"], align="center", numFmt="0.0")
        paint(put(ws, row, len(cols), r["Global Rating"], align="center"), GR_FILL.get(r["Global Rating"]))
        row += 1
    last = row - 1
    pctBar(ws, f"{get_column_letter(len(cols) - 1)}{hr + 1}:{get_column_letter(len(cols) - 1)}{last}")
    # footer: mean % per MC and count of SD/ND
    put(ws, row, 1, "Cohort mean % (applicable)", bold=True, fill=LBLUE); put(ws, row, 2, "", fill=LBLUE)
    put(ws, row + 1, 1, "n Sometimes / Not done", bold=True, fill=LBLUE); put(ws, row + 1, 2, "", fill=LBLUE)
    for k, mc in enumerate(mcs, 3):
        v = [SCORE_MAP[r["opts"][mc]] for r in done if r["opts"].get(mc) in SCORE_MAP]
        put(ws, row, k, round(sum(v) / len(v) * 100) if v else None, bold=True, align="center", numFmt="0")
        n = sum(1 for r in done if r["opts"].get(mc) in ("O4", "O5"))
        c = put(ws, row + 1, k, n, bold=True, align="center")
        if n:
            c.font = Font(name="Calibri", bold=True, color=ORANGE)
    pctScale(ws, f"C{row}:{get_column_letter(len(mcs) + 2)}{row}")
    widths(ws, [24, 18] + [5.5] * len(mcs) + [11, 5])
    finish(ws, f"C{hr + 1}", len(cols), hr)

    # ── 4. MC Item Analysis ───────────────────────────────────────────────────
    ws = wb.create_sheet("MC Item Analysis")
    cols = ["MC", "Section", "Description", "n Applicable", "Done well", "Done", "Mostly done", "Sometimes done", "Not done",
            "N/A", "Mean %", "% ≤ Mostly done", "Student self mean %", "Self-assessor exact agreement %", "Rank (weakest = 1)"]
    row = banner(ws, f"Item {itemCode} — Checklist item analysis ({dateTxt})",
                 "Mean % = mean points / 5 (Done well 5, Done 4, Mostly 3, Sometimes 2, Not done 0) over applicable ratings. Agreement compares the student's self-rating with the assessor's for the same form.",
                 len(cols))
    hr = row
    row = header(ws, row, cols, height=36)
    stats = []
    for mc in mcs:
        c = Counter(r["opts"].get(mc) for r in done if r["opts"].get(mc))
        app = [SCORE_MAP[r["opts"][mc]] for r in done if r["opts"].get(mc) in SCORE_MAP]
        selfV = [SCORE_MAP[r["selfOpts"][mc]] for r in done if r["selfOpts"].get(mc) in SCORE_MAP]
        pairs = [(r["opts"].get(mc), r["selfOpts"].get(mc)) for r in done if r["opts"].get(mc) and r["selfOpts"].get(mc)]
        stats.append({"mc": mc, "c": c, "n": len(app), "mean": round(sum(app) / len(app) * 100, 1) if app else None,
                      "low": round(sum(1 for r in done if r["opts"].get(mc) in ("O3", "O4", "O5")) / len(app) * 100, 1) if app else None,
                      "self": round(sum(selfV) / len(selfV) * 100, 1) if selfV else None,
                      "agree": round(sum(a == b for a, b in pairs) / len(pairs) * 100, 1) if pairs else None})
    ranked = sorted([s for s in stats if s["mean"] is not None], key=lambda s: s["mean"])
    rankOf = {s["mc"]: i + 1 for i, s in enumerate(ranked)}
    for i, s in enumerate(stats):
        band = PALE if i % 2 else None
        vals = [s["mc"], secMap[s["mc"]], mcDesc.get(s["mc"], ""), s["n"]] + [s["c"].get(o, 0) for o in ["O1", "O2", "O3", "O4", "O5", "O6"]] \
               + [s["mean"], s["low"], s["self"], s["agree"], rankOf.get(s["mc"])]
        for j, v in enumerate(vals, 1):
            c = put(ws, row, j, v, fill=band, size=9 if j == 3 else 10, align="left" if j <= 3 else "center",
                    numFmt="0.0" if j in (11, 12, 13, 14) else None)
            if 5 <= j <= 10 and v:
                paint(c, OPT_FILL[["O1", "O2", "O3", "O4", "O5", "O6"][j - 5]])
            if j == 15 and v is not None and v <= 3:
                paint(c, (ORANGE, "FFFFFF"))
        row += 1
    last = row - 1
    pctBar(ws, f"K{hr + 1}:K{last}"); pctBar(ws, f"L{hr + 1}:L{last}", color=ORANGE); pctBar(ws, f"M{hr + 1}:M{last}", color="8FAADC")
    pctBar(ws, f"N{hr + 1}:N{last}", color="A6A6A6")
    widths(ws, [6, 12, 60, 10, 8, 8, 8, 10, 8, 6, 12, 12, 12, 14, 9])
    finish(ws, f"D{hr + 1}", len(cols), hr)

    # ── 5. Self vs Assessor ───────────────────────────────────────────────────
    ws = wb.create_sheet("Self vs Assessor")
    cols = ["Student Name", "Student ID", "Assessor", "Assessor %", "Self %", "Self − Assessor", "MCs compared",
            "Exact agreement %", "Student rated higher", "Student rated lower", "Calibration"]
    row = banner(ws, f"Item {itemCode} — Student self-assessment vs assessor ({dateTxt})",
                 "Positive gap = student rated themselves higher than the assessor. Calibration: Over-estimates (gap ≥ +15), Under-estimates (≤ −15), else Aligned.",
                 len(cols))
    hr = row
    row = header(ws, row, cols, height=36)
    rank = {f"O{i}": i for i in range(1, 6)}   # O1 best
    srows = []
    for r in done:
        pairs = [(r["opts"][m], r["selfOpts"][m]) for m in mcs if r["opts"].get(m) in rank and r["selfOpts"].get(m) in rank]
        gap = round(r["Self %"] - r["% Score"], 1) if r["Self %"] is not None and r["% Score"] is not None else None
        srows.append((r, pairs, gap))
    srows.sort(key=lambda t: (t[2] is None, -(t[2] or 0)))
    for i, (r, pairs, gap) in enumerate(srows):
        band = PALE if i % 2 else None
        hi = sum(1 for a, s in pairs if rank[s] < rank[a]); lo = sum(1 for a, s in pairs if rank[s] > rank[a])
        cal = "No self-assessment" if gap is None else ("Over-estimates" if gap >= 15 else ("Under-estimates" if gap <= -15 else "Aligned"))
        vals = [r["Student Name"], r["Student ID"], r["Assessor"], r["% Score"], r["Self %"], gap, len(pairs),
                round(sum(a == s for a, s in pairs) / len(pairs) * 100, 1) if pairs else None, hi, lo, cal]
        for j, v in enumerate(vals, 1):
            c = put(ws, row, j, v, fill=band, align="left" if j in (1, 3, 11) else "center", numFmt="0.0" if j in (4, 5, 6, 8) else None)
            if j == 11:
                paint(c, {"Over-estimates": (LORANGE, "000000"), "Under-estimates": (LBLUE, "000000"), "No self-assessment": (GREY, "000000")}.get(v))
        row += 1
    last = row - 1
    pctBar(ws, f"D{hr + 1}:D{last}"); pctBar(ws, f"E{hr + 1}:E{last}", color="8FAADC"); pctBar(ws, f"H{hr + 1}:H{last}", color="A6A6A6")
    ws.conditional_formatting.add(f"F{hr + 1}:F{last}", ColorScaleRule(start_type="num", start_value=-30, start_color="6A9BD1",
                                  mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=30, end_color="F4A261"))
    widths(ws, [24, 11, 20, 12, 12, 11, 10, 12, 11, 11, 18])
    finish(ws, f"B{hr + 1}", len(cols), hr)

    # ── 6. By Assessor ────────────────────────────────────────────────────────
    ws = wb.create_sheet("By Assessor")
    cols = ["Assessor", "Forms", "Mean %", "SD %", "Min %", "Max %", "Mean GR", "Mean PR", "GR ≤ 2 (n)", "Pending forms"]
    row = banner(ws, f"Item {itemCode} — By assessor ({dateTxt})", "Differences may reflect student allocation as well as assessor stringency.", len(cols))
    hr = row
    row = header(ws, row, cols)
    byA = defaultdict(list)
    for r in done:
        byA[r["Assessor"]].append(r)
    pendA = Counter(r["Assessor"] for r in pending)
    arows = []
    for a, rs in byA.items():
        p = [r["% Score"] for r in rs if r["% Score"] is not None]
        g = [r["Global Rating"] for r in rs if r["Global Rating"] is not None]
        pr = [r["Practice Readiness"] for r in rs if r["Practice Readiness"] is not None]
        arows.append([a, len(rs), round(statistics.mean(p), 1) if p else None, round(statistics.stdev(p), 1) if len(p) > 1 else None,
                      min(p) if p else None, max(p) if p else None, round(statistics.mean(g), 2) if g else None,
                      round(statistics.mean(pr), 2) if pr else None, sum(1 for x in g if x <= 2), pendA.get(a, 0)])
    for a in pendA:
        if a not in byA:
            arows.append([a, 0, None, None, None, None, None, None, 0, pendA[a]])
    arows.sort(key=lambda x: (x[2] is None, -(x[2] or 0)))
    for i, vals in enumerate(arows):
        for j, v in enumerate(vals, 1):
            c = put(ws, row, j, v, fill=PALE if i % 2 else None, align="left" if j == 1 else "center",
                    numFmt="0.0" if j in (3, 4, 5, 6) else ("0.00" if j in (7, 8) else None))
            if j in (9, 10) and v:
                c.font = Font(name="Calibri", bold=True, color=ORANGE)
        row += 1
    pctBar(ws, f"C{hr + 1}:C{row - 1}")
    widths(ws, [24, 8, 12, 8, 8, 8, 9, 9, 10, 10])
    finish(ws, f"B{hr + 1}", len(cols), hr, landscape=False)

    # ── 7. Pending & Absent ───────────────────────────────────────────────────
    ws = wb.create_sheet("Pending & Absent")
    cols = ["Status", "Student Name", "Student ID", "Assessor", "Student submitted", "Last updated / last seen"]
    row = banner(ws, f"Item {itemCode} — Forms not yet scored and students without a form ({dateTxt})",
                 "Pending = form started but not submitted by the assessor at extract time. Absent = BOH1 student seen on an earlier item-"
                 f"{itemCode} session with no form on {sessionDate}.", len(cols))
    hr = row
    row = header(ws, row, cols)
    onDay = {r["Student ID"] for r in dayRows}
    absent = sorted(((sid, info) for sid, info in cohortStudents.items() if sid not in onDay), key=lambda t: t[1]["name"])
    items = [("Pending assessor", r["Student Name"], r["Student ID"], r["Assessor"], "Yes" if r["Student Submitted"] else "No",
              str(r["Updated"])[:16].replace("T", " ")) for r in sorted(pending, key=lambda r: r["Last Name"])]
    items += [("No form on session", info["name"], sid, info["lastAssessor"], "", info["lastDate"]) for sid, info in absent]
    for i, vals in enumerate(items):
        for j, v in enumerate(vals, 1):
            c = put(ws, row, j, v, fill=PALE if i % 2 else None, align="left" if j in (2, 4) else "center")
            if j == 1:
                paint(c, (LORANGE, "000000") if v.startswith("Pending") else (GREY, "000000"))
        row += 1
    if not items:
        put(ws, row, 1, "None — every form is submitted and every student has a form.", border=False, italic=True)
    widths(ws, [20, 26, 12, 22, 12, 22])
    finish(ws, f"A{hr + 1}", len(cols), hr, landscape=False)

    wb.save(outPath)
    return done, pending


# ═════════════════════════════════════════════════════════════════════════════
# History workbook (item across all dates)
# ═════════════════════════════════════════════════════════════════════════════
def buildHistoryWorkbook(rows, meta, sessionDate, itemCode, outPath, lastUpdated):
    itemName, mcDesc, secMap = meta
    mcs = mcSort(mcDesc)
    scored = [r for r in rows if r["Assessor Submitted"] and r["Date"] <= sessionDate]
    dates = sorted({r["Date"] for r in scored})
    dateHdr = [pd.to_datetime(d).strftime("%d %b") for d in dates]
    todayIdx = dates.index(sessionDate) if sessionDate in dates else None
    todayStudents = {r["Student ID"] for r in scored if r["Date"] == sessionDate}
    sub = (f"Item {itemCode} — {itemName}  |  {len(dates)} sessions {pd.to_datetime(dates[0]).strftime('%d %b')} – "
           f"{pd.to_datetime(dates[-1]).strftime('%d %b %Y')}  |  Assessor-submitted forms, mean of repeat attempts per day, blank = no form  |  "
           f"last form update {str(lastUpdated)[:16].replace('T', ' ')}")
    wb = Workbook()

    # ── 1. Cohort by date ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Cohort by Date"
    cols = ["Date", "Weekday", "Forms", "Students", "Mean %", "Median %", "Min %", "Mean GR", "GR ≤ 2 (n)", "PR L1 (n)", "Incidents"]
    row = banner(ws, f"BOH1 Item {itemCode} — Cohort trend by session", sub, len(cols))
    hr = row
    row = header(ws, row, cols)
    for i, d in enumerate(dates):
        rs = [r for r in scored if r["Date"] == d]
        p = [r["% Score"] for r in rs if r["% Score"] is not None]
        g = [r["Global Rating"] for r in rs if r["Global Rating"] is not None]
        vals = [pd.to_datetime(d).strftime("%d %b %Y"), pd.to_datetime(d).strftime("%a"), len(rs), len({r["Student ID"] for r in rs}),
                round(statistics.mean(p), 1) if p else None, round(statistics.median(p), 1) if p else None, min(p) if p else None,
                round(statistics.mean(g), 2) if g else None, sum(1 for x in g if x <= 2),
                sum(1 for r in rs if r["Practice Readiness"] == 1), sum(1 for r in rs if r["Clinical Incident"])]
        for j, v in enumerate(vals, 1):
            c = put(ws, row, j, v, fill=(LBLUE if d == sessionDate else (PALE if i % 2 else None)), bold=(d == sessionDate),
                    align="left" if j == 1 else "center", numFmt="0.0" if j in (5, 6, 7) else ("0.00" if j == 8 else None))
            if j in (9, 10, 11) and v:
                c.font = Font(name="Calibri", bold=True, color=ORANGE)
        row += 1
    pctBar(ws, f"E{hr + 1}:E{row - 1}")
    widths(ws, [14, 9, 8, 9, 12, 10, 8, 9, 10, 10, 10])
    finish(ws, f"A{hr + 1}", len(cols), hr, landscape=False)

    # ── 2-4. Student × date pivots ────────────────────────────────────────────
    students = {}
    for r in scored:
        students.setdefault(r["Student ID"], (r["Last Name"], r["First Name"], r["Student Name"]))
    order = sorted(students, key=lambda s: (students[s][0].lower(), students[s][1].lower()))
    for sheetName, key, fmt, fills in [("% Score by Date", "% Score", "0.0", None), ("GR by Date", "Global Rating", "0.0", GR_FILL),
                                       ("PR by Date", "Practice Readiness", "0.0", PR_FILL)]:
        ws = wb.create_sheet(sheetName)
        cols = ["Student Name", "Student ID", "Scored " + pd.to_datetime(sessionDate).strftime("%d %b")] + dateHdr + ["Mean (prior)", "Latest", "Latest − prior mean", "Sessions"]
        row = banner(ws, f"Item {itemCode} — {key} by session (student × date)", sub, len(cols))
        hr = row
        row = header(ws, row, cols, height=24)
        if todayIdx is not None:
            ws.cell(row=hr, column=4 + todayIdx).fill = PatternFill("solid", fgColor=ORANGE)
        for i, sid in enumerate(order):
            byDate = defaultdict(list)
            for r in scored:
                if r["Student ID"] == sid and r[key] is not None:
                    byDate[r["Date"]].append(r[key])
            series = [round(statistics.mean(byDate[d]), 1) if byDate.get(d) else None for d in dates]
            present = [(d, v) for d, v in zip(dates, series) if v is not None]
            latest = present[-1][1] if present else None
            prior = [v for d, v in present[:-1]]
            priorMean = round(statistics.mean(prior), 1) if prior else None
            delta = round(latest - priorMean, 1) if latest is not None and priorMean is not None else None
            band = PALE if i % 2 else None
            put(ws, row, 1, students[sid][2], fill=band); put(ws, row, 2, sid, fill=band, align="center")
            put(ws, row, 3, "Yes" if sid in todayStudents else "No", fill=band if sid in todayStudents else GREY, align="center")
            for k, v in enumerate(series):
                c = put(ws, row, 4 + k, v, align="center", numFmt=fmt, fill=band)
                if fills and v is not None:
                    paint(c, fills.get(int(round(v))))
            base = 4 + len(dates)
            put(ws, row, base, priorMean, align="center", numFmt=fmt, fill=band)
            put(ws, row, base + 1, latest, align="center", numFmt=fmt, fill=band, bold=True)
            c = put(ws, row, base + 2, delta, align="center", numFmt="+0.0;-0.0;0.0", fill=band)
            if delta is not None and delta != 0:
                c.font = Font(name="Calibri", bold=True, color=BLUE if delta > 0 else ORANGE)
            put(ws, row, base + 3, len(present), align="center", fill=band)
            row += 1
        last = row - 1
        if fills is None:
            pctScale(ws, f"D{hr + 1}:{get_column_letter(3 + len(dates))}{last}")
        # cohort mean footer
        put(ws, row, 1, "Cohort mean", bold=True, fill=LBLUE); put(ws, row, 2, "", fill=LBLUE); put(ws, row, 3, "", fill=LBLUE)
        for k, d in enumerate(dates):
            v = [r[key] for r in scored if r["Date"] == d and r[key] is not None]
            put(ws, row, 4 + k, round(statistics.mean(v), 1) if v else None, bold=True, fill=LBLUE, align="center", numFmt=fmt)
        widths(ws, [24, 11, 10] + [8] * len(dates) + [10, 8, 10, 8])
        finish(ws, f"D{hr + 1}", len(cols), hr)

    # ── 5. MC by date ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("MC by Date")
    cols = ["MC", "Section", "Description"] + dateHdr + ["Change (first → latest)"]
    row = banner(ws, f"Item {itemCode} — Cohort mean % per checklist item by session", sub, len(cols))
    hr = row
    row = header(ws, row, cols, height=24)
    if todayIdx is not None:
        ws.cell(row=hr, column=4 + todayIdx).fill = PatternFill("solid", fgColor=ORANGE)
    for i, mc in enumerate(mcs):
        series = []
        for d in dates:
            v = [SCORE_MAP[r["opts"][mc]] for r in scored if r["Date"] == d and r["opts"].get(mc) in SCORE_MAP]
            series.append(round(sum(v) / len(v) * 100, 1) if v else None)
        present = [v for v in series if v is not None]
        band = PALE if i % 2 else None
        put(ws, row, 1, mc, fill=band, bold=True); put(ws, row, 2, secMap[mc], fill=band)
        put(ws, row, 3, mcDesc.get(mc, ""), fill=band, size=9)
        for k, v in enumerate(series):
            put(ws, row, 4 + k, v, align="center", numFmt="0.0")
        ch = round(present[-1] - present[0], 1) if len(present) > 1 else None
        c = put(ws, row, 4 + len(dates), ch, align="center", numFmt="+0.0;-0.0;0.0", fill=band)
        if ch:
            c.font = Font(name="Calibri", bold=True, color=BLUE if ch > 0 else ORANGE)
        row += 1
    pctScale(ws, f"D{hr + 1}:{get_column_letter(3 + len(dates))}{row - 1}")
    widths(ws, [6, 12, 55] + [8] * len(dates) + [12])
    finish(ws, f"D{hr + 1}", len(cols), hr)

    # ── 6. All sessions (long) ────────────────────────────────────────────────
    ws = wb.create_sheet("All Sessions")
    cols = ["Date", "Student Name", "Student ID", "Assessor", "Score", "Max", "% Score", "GR", "PR", "Prof.", "Pos. & Ergo.",
            "Prepared", "Incident", "Self %", "Assessor Comments", "Student Reflection"]
    row = banner(ws, f"Item {itemCode} — Every assessor-submitted form", sub, len(cols))
    hr = row
    row = header(ws, row, cols)
    for i, r in enumerate(sorted(scored, key=lambda r: (r["Date"], r["Last Name"].lower(), r["First Name"].lower()))):
        vals = [r["Date"], r["Student Name"], r["Student ID"], r["Assessor"], r["Score"], r["Max Score"], r["% Score"],
                r["Global Rating"], r["Practice Readiness"], r["Professionalism"], r["Position & Ergonomics"], r["Preparedness"],
                r["Clinical Incident"], r["Self %"], r["Assessor Comments"], r["Student Reflection"]]
        band = LBLUE if r["Date"] == sessionDate and i % 2 else (PALE if i % 2 else None)
        for j, v in enumerate(vals, 1):
            c = put(ws, row, j, v, fill=band, size=9 if j >= 15 else 10,
                    align="left" if j in (2, 4, 15, 16) else "center", numFmt="0.0" if j in (7, 14) else None)
            if j == 8:
                paint(c, GR_FILL.get(v))
            elif j == 9:
                paint(c, PR_FILL.get(v))
            elif j == 13 and v:
                paint(c, (ORANGE, "FFFFFF"))
        row += 1
    pctBar(ws, f"G{hr + 1}:G{row - 1}")
    widths(ws, [11, 24, 11, 20, 7, 6, 11, 5, 5, 6, 7, 8, 8, 8, 55, 55])
    finish(ws, f"C{hr + 1}", len(cols), hr)

    wb.save(outPath)
    return scored


# ═════════════════════════════════════════════════════════════════════════════
# Entry point
# ═════════════════════════════════════════════════════════════════════════════
def buildSessionReports(sessionDate="2026-09-29", itemCode="532", jsonPath="temp 2026 caf.json", outDir="BOH1", cohort="BOH1"):
    """Build '<cohort> Item <code> Session Report (DD-MM-YYYY).xlsx' and '<cohort> Item <code> History (to DD-MM-YYYY).xlsx'."""
    checkFiveLevels(jsonPath, itemCode, cohort)
    rows, meta, lastUpdated = loadItemRows(jsonPath, itemCode, cohort)
    # students seen on any earlier session of this item -> used for the "absent" list
    cohortStudents = {}
    for r in sorted(rows, key=lambda r: r["Date"]):
        if r["Date"] < sessionDate:
            cohortStudents[r["Student ID"]] = {"name": r["Student Name"], "lastDate": r["Date"], "lastAssessor": r["Assessor"]}
    dmy = pd.to_datetime(sessionDate).strftime("%d-%m-%Y")
    sessPath = f"{outDir}/{cohort} Item {itemCode} Session Report ({dmy}).xlsx"
    histPath = f"{outDir}/{cohort} Item {itemCode} History (to {dmy}).xlsx"
    done, pending = buildSessionWorkbook(rows, meta, sessionDate, itemCode, sessPath, lastUpdated, cohortStudents)
    buildHistoryWorkbook(rows, meta, sessionDate, itemCode, histPath, lastUpdated)
    print(f"✓ {sessPath}  ({len(done)} scored, {len(pending)} pending)")
    print(f"✓ {histPath}")
    return sessPath, histPath


if __name__ == "__main__":
    buildSessionReports()
