import sys,pickle;sys.path.insert(0,'.')
from analysis import *
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.formatting.rule import ColorScaleRule,CellIsRule
from openpyxl.utils import get_column_letter
import warnings;warnings.filterwarnings('ignore')
NAVY='094183';ORANGE='E87722';BLUE='9DC3E6';ORLIGHT='F8CBAD'
dAll=pd.read_pickle('forms.pkl'); a=pd.read_pickle('assessors.pkl'); r2=pd.read_pickle('rated_adj.pkl')
trajR=pickle.load(open('traj.pkl','rb')); trajG=pickle.load(open('traj_grouped.pkl','rb'))
sheets={}
# 1 raw
raw=pd.concat([rawTable(r,['cohort','clinicM']).assign(period='All'),rawTable(r,['cohort','period','clinicM'])])
raw=raw[['cohort','period','clinicM']+[c for c in raw.columns if c not in('cohort','period','clinicM')]].rename(columns={'clinicM':'placement'})
sheets['Placement Raw']=raw.sort_values(['cohort','period','entMean'])
# 2 adjusted
adj=[]
for coh in ['DDS4','BOH3']:
  for per in ['All','FHY','SHY']:
    s=r[(r.cohort==coh)&((r.period==per)|(per=='All'))]
    e,_=feModel(s,'ent'); l,_=feModel(s,'low'); w,_=feModel(s,'anyWeak')
    x=e.rename(columns={'adjEffect':'entEffect','se':'entSE','ciLow':'entCiLow','ciHigh':'entCiHigh','p':'entP'})
    x=x.merge(l[['clinic','adjEffect','p']].rename(columns={'adjEffect':'lowShareEffect','p':'lowP'}),on='clinic')
    x=x.merge(w[['clinic','adjEffect','p']].rename(columns={'adjEffect':'anyWeakEffect','p':'weakP'}),on='clinic')
    x.insert(0,'period',per);x.insert(0,'cohort',coh); adj.append(x.sort_values('entEffect'))
adj=pd.concat(adj).rename(columns={'clinic':'placement'})
adj['verdict']=np.select([(adj.entCiHigh<0),(adj.entCiLow>0)],['Harder (lower entrustment)','Easier (higher entrustment)'],'Not different')
sheets['Placement Adjusted']=adj
# 3 IPC within-student
p=[]
for per in ['All','FHY','SHY']:
    s=r if per=='All' else r[r.period==per]
    p.append(pairedIPC(s).assign(period=per,assessorsExcluded='None'))
    p.append(pairedIPC(s[~s.assessorId.isin(['710','835'])]).assign(period=per,assessorsExcluded='Assessors 710 & 835'))
p=pd.concat(p); sheets['IPC Within-Student']=p[['cohort','period','assessorsExcluded']+[c for c in p.columns if c not in('cohort','period','assessorsExcluded')]]
# 4 IPC by assessor
ipc=r2[r2.clinicM==IPC]
t=ipc.groupby(['cohort','assessorId','period']).agg(assessorName=('assessorName','first'),forms=('ent','size'),students=('studentNumber','nunique'),
    entMean=('ent','mean'),lowShare=('low','mean'),anyWeak=('anyWeak','mean'),stageAdjDev=('entDev','mean')).reset_index()
t['shareOfIPCForms']=t.forms/t.groupby(['cohort','period']).forms.transform('sum')
t=t.merge(a[['cohort','assessorId','adjEffect','ciLow','ciHigh','status']].rename(columns={'adjEffect':'assessorEffectAllSites'}),on=['cohort','assessorId'],how='left')
sheets['IPC by Assessor']=t.sort_values(['cohort','period','forms'],ascending=[True,True,False])
# 5 assessor effects
sheets['Assessor Effects']=a[['cohort','assessorId','assessorName','n','students','mainClinic','mainClinicShare','clinics','rawEnt','lowShare','anyWeak','concernRate','resMean','shrink','adjEffect','ciLow','ciHigh','status']].rename(columns={'n':'forms','resMean':'rawResidual'}).sort_values(['cohort','adjEffect'])
# 6 variance
sheets['Variance Decomposition']=varianceDecomp(r)
# 7 flag impact
fi=[]
for lab,tr in [('Per rotation',trajR),('Flagging windows R1-3/R4-6/R7+',trajG)]:
    for col,name in [('ent','Raw entrustment (as flagged now)'),('entPlAdj','Placement-adjusted'),('entAsAdj','Assessor-adjusted')]:
        ws=tr[col][0].dropna(subset=['z']).copy()
        ws['ipcBand']=pd.cut(ws.ipcShare,[-0.01,0,0.34,1.0],labels=['No IPC','1-34% IPC','>34% IPC']).astype(str)
        g=ws.groupby(['cohort','ipcBand']).agg(windows=('z','size'),lowWindowRate=('lowWindow','mean'),meanZ=('z','mean')).reset_index()
        g.insert(0,'entrustmentUsed',name); g.insert(0,'windows_',lab); fi.append(g)
sheets['Flag Impact Windows']=pd.concat(fi).rename(columns={'windows_':'windowDefinition'})
st={k:v[1].set_index(['cohort','studentNumber']) for k,v in trajG.items()}
ch=pd.DataFrame({'stateRaw':st['ent'].state,'statePlacementAdj':st['entPlAdj'].state,'stateAssessorAdj':st['entAsAdj'].state,
                 'lowWindowsRaw':st['ent'].lowWindows,'lowWindowsPlacementAdj':st['entPlAdj'].lowWindows})
names=dAll.groupby('studentNumber').studentName.first()
ipcF=r2[r2.clinicM==IPC].groupby(['cohort','studentNumber']).size().rename('ratedFormsAtIPC')
ch=ch.join(ipcF).fillna({'ratedFormsAtIPC':0}).reset_index(); ch.insert(2,'studentName',ch.studentNumber.map(names))
ch['changed']=(ch.stateRaw!=ch.statePlacementAdj)|(ch.stateRaw!=ch.stateAssessorAdj)
sheets['Flag Impact Students']=ch[ch.changed|ch.stateRaw.isin(ENT_FLAG_STATES)|(ch.lowWindowsRaw!=ch.lowWindowsPlacementAdj)].sort_values(['cohort','ratedFormsAtIPC'],ascending=[True,False])
# 8 volume
dv=dAll[dAll.clinic.ne('Unknown/Other')].copy()
sr=dv.groupby(['cohort','period','studentNumber','rotNum']).agg(forms=('formId','size'),patients=('nAttended','sum'),items=('nItems','sum'),paed=('nPaed','sum'),
    mainClinic=('clinic',lambda s:s.value_counts().index[0])).reset_index()
vol=sr.groupby(['cohort','mainClinic']).agg(studentRotations=('forms','size'),formsPerRotation=('forms','mean'),patientsPerRotation=('patients','mean'),
    itemsPerRotation=('items','mean'),paedPerRotation=('paed','mean')).reset_index()
vol=vol[vol.studentRotations>=5]
for c in ['formsPerRotation','patientsPerRotation','itemsPerRotation']:
    vol[c+'VsCohort']=vol[c]/vol.groupby('cohort')[c].transform(lambda s: np.average(s,weights=vol.loc[s.index,'studentRotations']))
sheets['Volume by Placement']=vol.rename(columns={'mainClinic':'placement (main clinic in rotation)'}).sort_values(['cohort','formsPerRotationVsCohort'])
pickle.dump(sheets,open('sheets.pkl','wb'))
for k,v in sheets.items(): print(k,v.shape)
