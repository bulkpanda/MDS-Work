import os,re
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image as XlImage
from PIL import Image
BASE=os.path.expanduser('~/mnt/2026/Claude outputs/Placement analysis')
DOCX=BASE+'/DDS4_BOH3_Assessor_Placement_Report 06-10-2026.docx'; XLSX=BASE+'/DDS4_BOH3_Placement_Assessor_Analysis 06-10-2026.xlsx'
NAVY='094183'; NCOL=14; LASTCOL='N'; CHARS_PER_LINE=175
doc=Document(DOCX)
# map image rIds -> files saved to temp
imgDir=os.path.expanduser('~/work/docimgs'); os.makedirs(imgDir,exist_ok=True)
rels={r.rId:r for r in doc.part.rels.values() if 'image' in r.reltype}
# walk body in order, split on Heading 1
blocks=[]; cur=None
for el in doc.element.body.iterchildren():
    if el.tag==qn('w:p'):
        p=Paragraph(el,doc); sty=p.style.name if p.style is not None else ''
        blips=el.findall('.//'+qn('a:blip'))
        if sty=='Heading 1':
            cur={'title':p.text.strip(),'items':[]}; blocks.append(cur); continue
        if cur is None: cur={'title':'Title','items':[]}; blocks.append(cur)
        if blips:
            for b in blips:
                rid=b.get(qn('r:embed')); part=rels[rid].target_part; fn=os.path.join(imgDir,os.path.basename(part.partname)); open(fn,'wb').write(part.blob)
                ext=el.find('.//'+qn('wp:extent')); wcm=int(ext.get('cx'))/360000 if ext is not None else 16
                cur['items'].append(('img',fn,wcm))
        elif p.text.strip():
            kind='h2' if sty=='Heading 2' else ('bullet' if 'List' in sty else 'p')
            italic=any(r.italic for r in p.runs) and kind=='p'
            cur['items'].append((kind,p.text.strip(),italic))
    elif el.tag==qn('w:tbl'):
        t=Table(el,doc); rows=[]
        for r in t.rows:
            row=[]
            for c in r.cells:
                sh=c._tc.find('.//'+qn('w:shd')); row.append((c.text.strip(),sh.get(qn('w:fill')) if sh is not None else None))
            rows.append(row)
        if cur is None: cur={'title':'Title','items':[]}; blocks.append(cur)
        cur['items'].append(('table',rows,None))
wb=load_workbook(XLSX)
for n in [s for s in wb.sheetnames if re.match(r'R\d\d ',s)]+(['Report Contents'] if 'Report Contents' in wb.sheetnames else []): del wb[n]
thin=Side(style='thin',color='BBBBBB'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
SHORT=['Title & Stats','Summary','Background','Data & Methods','Rotation Stage','Placement Differences','IPC Deep-Dive','Assessor Differences','Assessor Exposure','Flag Impact','Activity by Placement','Data Quality','Recommendations','Limitations','App A Placements','App B Assessors']
def sheetName(i,title):
    if i<len(SHORT): return f'R{i:02d} '+SHORT[i]
    t=re.sub(r'^\d+\.\s*','',title); t=re.sub(r'[\[\]:*?/\\]','',t)
    return f'R{i:02d} '+t[:26]
created=[]
for i,b in enumerate(blocks):
    title='Report title & summary' if b['title']=='Title' else b['title']
    ws=wb.create_sheet(sheetName(i,title)); created.append((ws.title,title))
    ws.sheet_view.showGridLines=False; ws.column_dimensions['A'].width=2
    for c in range(2,NCOL+1): ws.column_dimensions[chr(64+c)].width=12.5
    r=2
    ws.cell(r,2,title).font=Font(name='Arial',size=15,bold=True,color=NAVY); ws.cell(r,NCOL,'← Contents').hyperlink="#'Report Contents'!A1"
    ws.cell(r,NCOL).font=Font(name='Arial',size=9,color='2A6EBB',underline='single'); r+=2
    for kind,a1,a2 in b['items']:
        if kind in('p','bullet','h2'):
            text=('•  '+a1) if kind=='bullet' else a1
            ws.merge_cells(f'B{r}:{LASTCOL}{r}'); c=ws.cell(r,2,text); c.alignment=Alignment(wrap_text=True,vertical='top')
            if kind=='h2': c.font=Font(name='Arial',size=12,bold=True,color=NAVY); ws.row_dimensions[r].height=20
            else:
                c.font=Font(name='Arial',size=10,italic=a2,color='555555' if a2 else '222222')
                lines=max(1,-(-len(text)//CHARS_PER_LINE)); ws.row_dimensions[r].height=14*lines+4
            r+=1 if kind!='h2' else 1
        elif kind=='img':
            im=Image.open(a1); wpx=min(1100,int(a2*62)); hpx=int(im.size[1]*wpx/im.size[0])
            xi=XlImage(a1); xi.width=wpx; xi.height=hpx; ws.add_image(xi,f'B{r}')
            r+=int(hpx/20)+2
        elif kind=='table':
            rows=a1
            if len(rows)==1 and len(rows[0])==1:   # callout box
                txt,fill=rows[0][0]; ws.merge_cells(f'B{r}:{LASTCOL}{r}'); c=ws.cell(r,2,txt)
                c.alignment=Alignment(wrap_text=True,vertical='top'); c.fill=PatternFill('solid',fgColor=fill or 'EEF2FF'); c.font=Font(name='Arial',size=10,bold=txt.startswith(('▶','CONFIDENTIAL')) and False)
                lines=max(1,-(-len(txt)//(CHARS_PER_LINE-10))); ws.row_dimensions[r].height=14*lines+8; r+=2; continue
            ncols=len(rows[0]); span=max(1,(NCOL-1)//ncols) if ncols<=6 else 1
            for ri,row in enumerate(rows):
                maxLen=0
                for ci,(txt,fill) in enumerate(row):
                    col=2+ci*span
                    if span>1: ws.merge_cells(start_row=r,start_column=col,end_row=r,end_column=col+span-1)
                    c=ws.cell(r,col,txt); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=border
                    if ri==0: c.fill=PatternFill('solid',fgColor=NAVY); c.font=Font(name='Arial',size=9,bold=True,color='FFFFFF')
                    else:
                        c.font=Font(name='Arial',size=9)
                        if fill and fill not in('auto','FFFFFF'): c.fill=PatternFill('solid',fgColor=fill)
                    maxLen=max(maxLen,len(txt)/(span*1.0))
                ws.row_dimensions[r].height=max(15,13*(-(-int(maxLen)//14)))
                r+=1
            r+=1
# contents sheet
cs=wb.create_sheet('Report Contents',0); cs.sheet_view.showGridLines=False
cs.column_dimensions['A'].width=2; cs.column_dimensions['B'].width=60; cs.column_dimensions['C'].width=80
cs['B2']='DDS4 & BOH3 2026 — Assessor & Clinic Placement Differences'; cs['B2'].font=Font(name='Arial',size=15,bold=True,color=NAVY)
cs['B3']='Report sheets (R00–R15) contain the written report with figures; the data sheets after them hold the full tables.'; cs['B3'].font=Font(name='Arial',size=9,color='555555')
r=5
for name,title in created:
    c=cs.cell(r,2,title); c.hyperlink=f"#'{name}'!A1"; c.font=Font(name='Arial',size=10,color='2A6EBB',underline='single'); r+=1
r+=1; cs.cell(r,2,'Data sheets').font=Font(name='Arial',size=11,bold=True,color=NAVY); r+=1
desc={'Summary':'One-page summary of findings','Notes':'Methods and definitions','Placement Raw':'Unadjusted metrics by cohort × period × placement',
 'Placement Adjusted':'Same-student, stage-adjusted placement effects','IPC Within-Student':'Paired IPC vs elsewhere tests','IPC by Assessor':'IPC forms by assessor and period',
 'Assessor Effects':'All assessors, adjusted effects','Variance Decomposition':'Incremental R² by component','Flag Impact Windows':'Low-window rates raw vs adjusted',
 'Flag Impact Students':'Students whose flag state changes','Volume by Placement':'Forms / patients / items per rotation'}
for n,dd in desc.items():
    if n in wb.sheetnames:
        c=cs.cell(r,2,n); c.hyperlink=f"#'{n}'!A1"; c.font=Font(name='Arial',size=10,color='2A6EBB',underline='single'); cs.cell(r,3,dd).font=Font(name='Arial',size=9,color='555555'); r+=1
# order: contents, report sheets, then data sheets
order=['Report Contents']+[n for n,_ in created]+[s for s in wb.sheetnames if s!='Report Contents' and s not in dict(created)]
wb._sheets=[wb[n] for n in order]; wb.active=0
wb.save(XLSX); print(order)
