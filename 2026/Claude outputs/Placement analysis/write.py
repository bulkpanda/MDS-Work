import pickle,pandas as pd,numpy as np,os
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
NAVY='094183';ORANGE='F4B183';BLUE='9DC3E6'
sheets=pickle.load(open('sheets.pkl','rb'))
outDir=os.path.expanduser('~/mnt/2026/Claude outputs/Placement analysis'); os.makedirs(outDir,exist_ok=True)
fn=os.path.join(outDir,'DDS4_BOH3_Placement_Assessor_Analysis 06-10-2026.xlsx')
summary=[
("PLACEMENT & ASSESSOR DIFFICULTY ANALYSIS — DDS4 / BOH3 CAF 2026",""),
("Data","DASH CAF forms 15 Jan – 5 Oct 2026 (16,571 forms; 15,923 assessor-rated with entrustment). FHY = before 1 Jul, SHY = from 1 Jul."),
("Question","Are students placed at IPC (Wyndham Vale) being flagged because the placement / its assessors rate harder?"),
("",""),
("KEY FINDINGS",""),
("1. DDS4: IPC is the hardest placement","Same student, same stage of the year: entrustment 0.46 levels lower at IPC than at their other placements (26 of 27 students lower; p<0.001). Share of L1–2 ratings +24 pts. FHY gap −0.58 (18/18 students lower); SHY −0.40 (15/15)."),
("2. Mostly two assessors, partly the site","Assessors 710 and 835 rated ~60% of DDS4 IPC forms (raw mean 2.42 and 2.66; 64% / 35% L1–2). Without them the IPC gap halves to −0.23 but stays significant, so a smaller site-wide stricter standard remains."),
("3. Not a workload issue; weakness tagging is split","IPC rotations are not low-volume (DDS4: 1.21× cohort forms and patients per rotation). At site level IPC forms do not carry significantly more weakness tags, but assessors 710 and 835 tag a weakness on 72% / 97% of their forms while other IPC assessors tag fewer than average. The two stringent raters therefore rate lower AND document more weaknesses. DDS4 students also self-rate slightly lower at IPC (−0.11), about a quarter of the assessor gap. See the full report for detail."),
("4. Direct impact on flagging","Using the flagging tool's windows (R1-3 / R4-6 / R7+, z ≤ −1 = low): DDS4 windows with >34% of forms at IPC are judged LOW 69% of the time vs 9% with no IPC. After placement adjustment: 19% vs 12%. See 'Flag Impact' sheets for the students whose state changes."),
("5. BOH3: small, FHY only","BOH3 students are ~0.17 lower at IPC in FHY (19/26 lower) and no different in SHY. The effect is spread across several IPC assessors (700, 1031, 1033 slightly harsher; 699 typical)."),
("6. Assessor > placement > student","Share of entrustment variation explained (after rotation stage): DDS4 — student 6%, placement 9%, assessor 24%; BOH3 — student 5%, placement 3%, assessor 18%. Once the assessor is known, placement adds ~0%. Assessors are mostly fixed to one site, so 'placement difficulty' is largely 'who assesses there'."),
("7. Other placements","DDS4: Link Health (Clayton) also harder (−0.38). Congress (Alice Springs), Health Ability (Box Hill), VAHS (Fitzroy) and Northwest Health are more lenient. BOH3: North Richmond, Preston, GV Health more lenient; none significantly harder overall."),
("8. Data issue found (reports)","BOH3 and DDS4 forms use DIFFERENT clinic code lists (e.g. BOH3 EC05 = IPC, EC08 = MDC; DDS4 EC08 = IPC). processDds4BohForms applies the DDS4 CLINIC_DICT to all rows, so BOH3 rows stored as codes (1,998 forms) get the wrong clinic name in DB-based clinic reports (e.g. BOH3 MDC forms show as IPC). This analysis decodes each form from its own context_schema_snapshot. Entrustment-based flags are not affected (they don't use clinic)."),
("",""),
("SUGGESTED ACTIONS",""),
("a","Review flags for students whose low window was mainly at IPC (sheet 'Flag Impact Students') before acting on them."),
("b","Calibration / moderation with IPC assessors 710 and 835 (DDS4), and benchmarking of IPC rating against other sites."),
("c","Optionally add a placement- or assessor-adjusted entrustment option to the flagging tool (method in 'Notes')."),
("d","Fix BOH3 clinic code decoding (decode per form from context_schema_snapshot)."),
("",""),
("CAVEATS","Observational data; students are not randomly allocated, but each student acts as their own control (student fixed effects) and rotation stage is adjusted for. Entrustment (L1–L4) treated as numeric. Site and assessor cannot be fully separated because most assessors work at one site. Assessor IDs are DASH IDs; names are in the assessor sheets."),
]
notes=[
("METHODS",""),
("Outcome","Assessor entrustment (S1–S4 → 1–4) on assessor-submitted forms. Also low share (L1–2), any weakness tagged, concerns, incidents."),
("Rotation stage","Entrustment rises from 2.81 (R1) to 3.38 (R8), so every adjusted figure controls for rotation number (FHY 'Rotation N' and SHY 'RN' labels merged by number)."),
("Placement Adjusted","Per cohort & period: OLS  entrustment ~ student FE + rotation FE + placement; effects centred on the form-weighted average placement; 95% CI clustered by assessor. Placement with <30 rated forms in a cohort pooled into 'Other (small)'. Same model for low share and any-weakness (linear probability)."),
("IPC Within-Student","Students with ≥3 rated forms at IPC and ≥3 elsewhere. For each: mean (entrustment − cohort×rotation mean) at IPC minus elsewhere. One-sample t-test and Wilcoxon signed-rank."),
("Assessor Effects","Residual after student FE + rotation FE (no placement term), averaged per assessor and shrunk toward 0 (empirical Bayes, so low-volume assessors aren't over-read). 'Harsher'/'More lenient' = ≥15 forms and 95% interval excludes 0. Effects include whatever is specific to the assessor's site."),
("Variance Decomposition","Incremental R² from adding student, placement and assessor terms in turn; last row shows placement adds ~0 beyond assessor."),
("Flag Impact","Re-implements evaluateEntrustmentTrajectory (boh3_dds4_flagging.py): window z vs cohort among students with ≥5 rated forms; low = z ≤ −1; Falling behind = last window low and z drop ≥1; Consistently low = all windows low. Run on raw, placement-adjusted (form entrustment − its placement effect) and assessor-adjusted entrustment. 'ipcShare' = share of the student's forms in that window that were at IPC. Approximate mirror only: the live tool uses the DB table and its own eligibility rules."),
("Volume","Per student × rotation: forms, patients attended, items, patients <18; grouped by the student's main clinic in that rotation; ratio vs cohort average."),
("Clinic names","Coded values (ECxx) decoded per form from context_schema_snapshot; free text normalised (e.g. pc / Primary Care / RDHM PC → RDHM PC; moe → La Trobe (Moe); Gove variants → Gove / NT). 38 unusable entries (numbers, subject code) excluded."),
("Reproduce","Scripts in this folder: flatten.py (JSON → forms.csv), prep.py, analysis.py, build.py, write.py. Source: 'temp 2026 caf dds4_boh3.json'."),
]
with pd.ExcelWriter(fn,engine='openpyxl') as xw:
    pd.DataFrame(summary,columns=['Item','Detail']).to_excel(xw,'Summary',index=False)
    for k,v in sheets.items(): v.to_excel(xw,k[:31],index=False)
    pd.DataFrame(notes,columns=['Item','Detail']).to_excel(xw,'Notes',index=False)
wb=load_workbook(fn)
hdrFill=PatternFill('solid',fgColor=NAVY)
pctCols=('lowShare','anyWeak','techWeak','timeWeak','concernRate','incidentRate','mainClinicShare','shareOfIPCForms','lowWindowRate','meanLowShareDiff','lowShareEffect','anyWeakEffect','R2','adjR2','addedR2','shrink')
for ws in wb:
    for c in ws[1]: c.fill=hdrFill; c.font=Font(color='FFFFFF',bold=True); c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'
    if ws.title in('Summary','Notes'):
        ws.column_dimensions['A'].width=42; ws.column_dimensions['B'].width=130
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
            if row[0].value and row[0].value.isupper(): row[0].font=Font(bold=True,color=NAVY,size=12)
            elif row[0].value: row[0].font=Font(bold=True)
        continue
    hdr=[c.value for c in ws[1]]
    for j,h in enumerate(hdr,1):
        L=get_column_letter(j); ws.column_dimensions[L].width=min(max(12,len(str(h))+2),45)
        fmt='0.0%' if any(str(h).startswith(p) for p in pctCols) else ('0.000' if h in('p','entP','lowP','weakP','pT','pWilcoxon') else '0.00')
        for c in ws[L][1:]:
            if isinstance(c.value,float): c.number_format=fmt
        rng=f'{L}2:{L}{ws.max_row}'
        if h in('entEffect','adjEffect','meanEntDiff','stageAdjDev','assessorEffectAllSites','meanZ'):
            ws.conditional_formatting.add(rng,CellIsRule(operator='lessThan',formula=['-0.15'],fill=PatternFill('solid',fgColor=ORANGE)))
            ws.conditional_formatting.add(rng,CellIsRule(operator='greaterThan',formula=['0.15'],fill=PatternFill('solid',fgColor=BLUE)))
        if h in('verdict','status','stateRaw','statePlacementAdj','stateAssessorAdj'):
            for c in ws[L][1:]:
                v=str(c.value)
                if v.startswith('Harder') or v=='Harsher' or v in('Falling behind','Consistently low'): c.fill=PatternFill('solid',fgColor=ORANGE)
                elif v.startswith('Easier') or v=='More lenient': c.fill=PatternFill('solid',fgColor=BLUE)
    ws.auto_filter.ref=ws.dimensions
wb.move_sheet('Notes',offset=-(len(wb.sheetnames)-2))
wb.save(fn); print(fn); print(wb.sheetnames)
