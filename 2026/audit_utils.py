"""
audit_utils.py — Clinical record Audit (DASH `audit` forms) → Excel report, per cohort.

Built for DDS3 2026 (DENT90150, checklists DDS3-1 … DDS3-5, single assessor per audit).
Called from main.ipynb section "## Audit" via `from audit_utils import *`.

Public functions (everything else is module-internal):
    fetchAuditJson(bearerToken, outPath, year, cohorts)          pull from DASH (ordering=id) + dedupe
    resolveAudits(cohort, audits)                                selected audits -> {label: (start, end)}
    loadAudit(jsonPath, cohort, rosterPath, audits)              parse + per-form/per-student/item/assessor stats (dict)
    buildAuditReport(jsonPath, outPath, cohort, rosterPath, audits, ...) full styled workbook

Form model (from the JSON):
    form.checklists[<cohort>-<n>]  → section name, fields {MCx: item text}, options, rubric
    form.data.assessor[<cohort>-<n>][MCx] → answer LABEL ("Yes" / "No" / "N/A" / "Not Applicable")
    form.data.assessor["<cohort>-<n>-comment"], "additional_comments", "coordinator_concerns(_flag)"
Scoring:
    Compliance % = Yes / (Yes + No) over the checklist items (N/A and blank excluded); the Summary
    section (Yes / Yes but action required / No) is the audit OUTCOME and is not part of compliance.
Audits: each audit is a date range, {label: (first date, last date)} (AUDIT_ROUNDS[cohort] holds the defaults).
`audits=` selects what goes in the report: None = all configured audits · "Audit 2" or ["Audit 1", "Audit 2"]
= those configured audits · a dict {label: (start, end)} = custom ranges. Forms outside the selected ranges are
dropped. Several forms for one student in one audit = re-audit; attempt 1 = "First", the latest = "Final".
"""
import csv
import json
import re
import datetime as _dt
import statistics as _st
import collections as _collections

from openpyxl import Workbook as _XlWorkbook
from openpyxl.styles import Font as _XlFont, PatternFill as _XlPatternFill, Alignment as _XlAlignment, \
    Border as _XlBorder, Side as _XlSide
from openpyxl.utils import get_column_letter as _xlColLetter
from openpyxl.chart import BarChart as _XlBarChart, Reference as _XlReference
from openpyxl.chart.label import DataLabelList as _XlDataLabelList
from openpyxl.chart.series import SeriesLabel as _XlSeriesLabel
from openpyxl.formatting.rule import ColorScaleRule as _XlColorScaleRule

__all__ = ["AUDIT_ROUNDS", "AUDIT_SUMMARY_KEY", "AUDIT_LOW_COMPLIANCE_PCT",
           "fetchAuditJson", "resolveAudits", "loadAudit", "buildAuditReport"]

# ---------------------------------------------------------------- config (edit per cohort / year)
# Audits per cohort: {label: (first date, last date)} inclusive, ISO strings, in report order.
AUDIT_ROUNDS = {
    "DDS3": {"Audit 1": ("2026-01-01", "2026-07-31"),
             "Audit 2": ("2026-08-01", "2026-12-31")},
}
# Checklist key holding the overall audit outcome (auto-detected when a cohort is missing here).
AUDIT_SUMMARY_KEY = {"DDS3": "DDS3-5"}
AUDIT_LOW_COMPLIANCE_PCT = 80.0      # flag a final form below this compliance %
_DECLINE_PTS = 10.0                  # flag a compliance drop of ≥ this many points between audits
_BASE_URL = "https://api.unimelb-dash.com"

_ANSWER_NORMALISE = {"yes": "Yes", "no": "No", "n/a": "N/A", "na": "N/A", "not applicable": "N/A"}
_OUTCOMES = ["Yes", "Yes but action required", "No"]

# palette — UniMelb navy/blue + orange (colour-blind safe pairing), matches the other workbooks
_NAVY, _BLUE, _ORANGE, _INK, _SUBHDR = "1F3864", "2E5496", "ED7D31", "1E293B", "D6E0F0"
_ANS_FILL = {"Yes": ("DDEBF7", "1F4E79"), "No": ("F8CBAD", "843C0C"), "N/A": ("EDEDED", "595959"),
             "": ("FFFFFF", "A6A6A6")}
_OUT_FILL = {"Yes": ("DDEBF7", "1F4E79"), "Yes but action required": ("FFE699", "7F6000"),
             "No": ("F8CBAD", "843C0C"), "": ("FFFFFF", "A6A6A6")}
_OUT_CHART = {"Yes": "2E5496", "Yes but action required": "FFC000", "No": "ED7D31"}
_THIN = _XlSide(style="thin", color="BFBFBF")
_BORDER = _XlBorder(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ================================================================ fetch
def fetchAuditJson(bearerToken, outPath, year=2026, cohorts="DDS3", baseUrl=_BASE_URL):
    """Pull every audit form from DASH (ordering=id so paging is stable), drop repeated ids, save JSON."""
    import requests
    url = f"{baseUrl}/assessment/audit/get?page_size=max&page=1&cohort={cohorts}&year={year}&ordering=id"
    headers = {"Authorization": f"Token {bearerToken}", "Accept": "application/json"}
    records = []
    with requests.Session() as session:
        while url:                                           # follow `next` links
            resp = session.get(url, headers=headers, timeout=60)
            resp.raise_for_status()
            payload = resp.json()
            page = payload.get("results")
            if not isinstance(page, list):
                raise ValueError("Unexpected response: missing 'results' list")
            records.extend(page)
            url = payload.get("next")
    byId = {}
    for r in records: byId[r["id"]] = r                      # dedupe by id
    if len(byId) != len(records):
        print(f"WARNING: DASH returned {len(records) - len(byId)} repeated row(s) — removed.")
    records = sorted(byId.values(), key=lambda r: r["id"])
    with open(outPath, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    print(f"Saved {len(records)} unique audit records -> {outPath}")
    return records


# ================================================================ helpers
def _isTestName(name):
    n = (name or "").lower()
    return "dummy" in n or "test" in n


def _surnameKey(n):
    return (n.split()[-1].lower(), n.lower()) if n else ("", "")


def _cleanName(n):
    n = (n or "").strip()
    return n.title() if n.islower() else n


def _stripHtml(s):
    s = re.sub(r"</p>\s*<p>", "\n", s or "")
    s = re.sub(r"<[^>]+>", "", s)
    return s.replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&").replace("&nbsp;", " ").strip()


def _normAnswer(v):
    if v is None or v == "": return ""
    return _ANSWER_NORMALISE.get(str(v).strip().lower(), str(v).strip())


def _sectionNo(key):
    m = re.search(r"-(\d+)$", key)
    return int(m.group(1)) if m else 999


def _excludedNumbers(cohort):
    """Cohort remove-list + global exclusion list from variableUtils (empty if unavailable)."""
    try:
        from variableUtils import REMOVE_STUDENTS_DICT, EXCLUDED_STUDENT_NUMBERS
        return {int(x) for x in list(REMOVE_STUDENTS_DICT.get(cohort, [])) + list(EXCLUDED_STUDENT_NUMBERS)}
    except Exception:
        return set()


def resolveAudits(cohort="DDS3", audits=None):
    """Selected audits as an ordered {label: (start, end)} dict.

    audits: None -> every audit in AUDIT_ROUNDS[cohort] (or one 'All' range if the cohort has none)
            "Audit 2" / ["Audit 1", "Audit 2"] -> those audits from AUDIT_ROUNDS[cohort]
            {label: (start, end)} -> used as given (custom ranges, ISO date strings or date objects)
    """
    configured = AUDIT_ROUNDS.get(cohort) or {}
    if audits is None:
        sel = dict(configured) or {"All": ("1900-01-01", "2999-12-31")}
    elif isinstance(audits, dict):
        sel = dict(audits)
    else:
        names = [audits] if isinstance(audits, str) else list(audits)
        missing = [n for n in names if n not in configured]
        if missing: raise KeyError(f"{missing} not in AUDIT_ROUNDS['{cohort}'] ({list(configured)}) — pass a dict instead")
        sel = {n: configured[n] for n in names}
    out = {}
    for label, (a, b) in sel.items():
        a = a if isinstance(a, _dt.date) else _dt.date.fromisoformat(str(a))
        b = b if isinstance(b, _dt.date) else _dt.date.fromisoformat(str(b))
        if a > b: raise ValueError(f"{label}: start {a} is after end {b}")
        out[label] = (a, b)
    return out


def _roundOf(date, rounds):
    for label, (a, b) in rounds.items():
        if a <= date <= b: return label
    return None


def _pct(a, b):
    return 100.0 * a / b if b else None


# ================================================================ load + stats
def loadAudit(jsonPath, cohort="DDS3", rosterPath="studentEmailList.csv", audits=None, summaryKey=None):
    """Parse the audit JSON for one cohort, keep forms inside the selected audits, compute every report table."""
    raw = json.load(open(jsonPath, encoding="utf-8"))
    rounds = resolveAudits(cohort, audits)                   # {label: (startDate, endDate)}
    roundLabels = list(rounds)

    # roster (name → row), cohort-restricted first, any cohort as fallback
    rosterAll = list(csv.DictReader(open(rosterPath, encoding="utf-8")))
    rosterCohort = {r["student_name"].strip().lower(): r for r in rosterAll if r.get("cohort") == cohort}
    rosterAny = {r["student_name"].strip().lower(): r for r in rosterAll}
    excludedNos = _excludedNumbers(cohort)

    cohortRecs = [r for r in raw if r.get("cohort") == cohort]
    if not cohortRecs: raise ValueError(f"No {cohort} records in {jsonPath}")

    # ---- checklist structure: latest submitted form's config (labels were reworded between versions)
    ref = max(cohortRecs, key=lambda r: (bool(r.get("submitted")), r["id"]))
    checklists = ref["form"]["checklists"]
    sections = sorted(checklists, key=_sectionNo)
    if summaryKey is None:
        summaryKey = AUDIT_SUMMARY_KEY.get(cohort) or next(
            (k for k in sections if any("action" in str(o).lower()
                                        for o in checklists[k]["extra_config"].get("options", {}).get("assessor", {}).values())),
            None)
    itemSections = [k for k in sections if k != summaryKey]
    items = [dict(section=k, sectionName=checklists[k]["name"], mc=mc, text=txt,
                  rubric=_stripHtml(checklists[k]["extra_config"].get("rubric", {}).get(mc, {}).get("label", "")))
             for k in itemSections for mc in sorted(checklists[k]["fields"], key=lambda x: int(re.sub(r"\D", "", x) or 0))
             for txt in [checklists[k]["fields"][mc]]]
    itemKeys = [(it["section"], it["mc"]) for it in items]

    # ---- forms
    excluded, seen, forms, outside = [], set(), [], 0
    for r in sorted(cohortRecs, key=lambda r: r["id"]):
        date = _dt.date.fromisoformat(r["datetime"][:10])
        if _roundOf(date, rounds) is None:                   # not in any selected audit
            outside += 1; seen.add(r["id"]); continue
        ans = ((r.get("form") or {}).get("data") or {}).get("assessor") or {}
        nameKey = (r.get("student") or "").strip().lower()
        ros = rosterCohort.get(nameKey) or rosterAny.get(nameKey) or {}
        stNo = int(ros["student_number"]) if ros.get("student_number") else None
        reason = None
        if r["id"] in seen: reason = "Repeated form id in export"
        elif not r.get("submitted"): reason = "Not submitted (draft)"
        elif _isTestName(r.get("student")) or _isTestName(r.get("assessor")): reason = "Dummy/test student or assessor"
        elif stNo is not None and stNo in excludedNos: reason = "Student on exclusion list"
        elif not any(isinstance(ans.get(k), dict) and ans.get(k) for k in sections): reason = "No checklist answers"
        seen.add(r["id"])
        if reason:
            excluded.append(dict(id=r["id"], student=r.get("student"), assessor=r.get("assessor"),
                                 date=r["datetime"][:10], reason=reason))
            continue
        answers = {(k, mc): _normAnswer((ans.get(k) or {}).get(mc)) for k, mc in itemKeys}
        nYes = sum(v == "Yes" for v in answers.values())
        nNo = sum(v == "No" for v in answers.values())
        nNa = sum(v == "N/A" for v in answers.values())
        outcome = _normAnswer((ans.get(summaryKey) or {}).get("MC1")) if summaryKey else ""
        outcome = {"Yes But Action Required": "Yes but action required"}.get(outcome, outcome)
        forms.append(dict(
            id=r["id"], date=date, round=_roundOf(date, rounds), subject=r.get("subject"),
            student=_cleanName(ros.get("student_name") or r.get("student")), studentNo=stNo, inRoster=bool(ros),
            assessor=(r.get("assessor") or "").strip(), answers=answers, nYes=nYes, nNo=nNo, nNa=nNa,
            nBlank=len(itemKeys) - nYes - nNo - nNa, compliance=_pct(nYes, nYes + nNo), outcome=outcome,
            noItems=[f"P{_sectionNo(k)} {next(i['text'] for i in items if (i['section'], i['mc']) == (k, mc))}"
                     for (k, mc), v in answers.items() if v == "No"],
            sectionComments={k: (ans.get(f"{k}-comment") or "").strip() for k in sections},
            comments=(ans.get("additional_comments") or "").strip(),
            concernFlag=(ans.get("coordinator_concerns_flag") or "").strip(),
            concerns=(ans.get("coordinator_concerns") or "").strip()))

    # ---- attempts within an audit (re-audits)
    byStuRound = _collections.defaultdict(list)
    for f in forms: byStuRound[(f["student"], f["round"])].append(f)
    for fs in byStuRound.values():
        fs.sort(key=lambda f: (f["date"], f["id"]))
        for i, f in enumerate(fs):
            f["attempt"], f["nAttempts"] = i + 1, len(fs)
            f["isFirst"], f["isFinal"] = i == 0, i == len(fs) - 1

    allRounds = roundLabels

    # ---- students (roster ∪ audited)
    names = {f["student"]: f for f in forms}
    studentRows = {v["student_name"].strip(): dict(name=_cleanName(v["student_name"]), no=int(v["student_number"]),
                                                   inRoster=True)
                   for v in rosterCohort.values() if int(v["student_number"]) not in excludedNos}
    for n, f in names.items():
        if n not in studentRows: studentRows[n] = dict(name=n, no=f["studentNo"], inRoster=f["inRoster"])
    students = []
    for s in studentRows.values():
        s["rounds"] = {}
        for rl in allRounds:
            fs = byStuRound.get((s["name"], rl), [])
            if fs:
                first, final = fs[0], fs[-1]
                s["rounds"][rl] = dict(forms=fs, first=first, final=final, n=len(fs))
        flags = []
        for rl in roundLabels:
            rd = s["rounds"].get(rl)
            if not rd:
                flags.append(f"Not audited ({rl})"); continue
            fi, fl = rd["first"], rd["final"]
            if fi["outcome"] == "No":
                flags.append(f"Outcome No → re-audited ({rl})" if rd["n"] > 1 else f"Outcome No, no re-audit ({rl})")
            if fl["outcome"] == "Yes but action required": flags.append(f"Action required ({rl})")
            if fl["outcome"] == "No" and rd["n"] > 1: flags.append(f"Still No after re-audit ({rl})")
            if fl["compliance"] is not None and fl["compliance"] < AUDIT_LOW_COMPLIANCE_PCT:
                flags.append(f"Compliance < {AUDIT_LOW_COMPLIANCE_PCT:.0f}% ({rl})")
            if any(f["concernFlag"].lower() == "yes" or f["concerns"] for f in rd["forms"]):
                flags.append(f"Coordinator concern ({rl})")
        done = [s["rounds"][rl]["final"]["compliance"] for rl in roundLabels if rl in s["rounds"]]
        s["change"] = (done[-1] - done[0]) if len(done) >= 2 and None not in (done[0], done[-1]) else None
        if s["change"] is not None and s["change"] <= -_DECLINE_PTS: flags.append(f"Compliance fell ≥ {_DECLINE_PTS:.0f} pts")
        if not s["inRoster"]: flags.append("Not on roster")
        s["flags"] = flags
        students.append(s)
    students.sort(key=lambda s: _surnameKey(s["name"]))

    # ---- item analysis on FIRST attempts (one form per student per audit, the un-coached state)
    itemStats = []
    for it in items:
        row = dict(it, byRound={})
        for rl in allRounds:
            vals = [f["answers"][(it["section"], it["mc"])] for f in forms if f["round"] == rl and f["isFirst"]]
            c = _collections.Counter(vals)
            row["byRound"][rl] = dict(yes=c["Yes"], no=c["No"], na=c["N/A"], blank=c[""], n=len(vals),
                                      pctYes=_pct(c["Yes"], c["Yes"] + c["No"]), pctNo=_pct(c["No"], c["Yes"] + c["No"]),
                                      pctNa=_pct(c["N/A"], len(vals)))
        itemStats.append(row)

    # ---- section (part) compliance per audit, first attempts
    sectionStats = []
    for k in itemSections:
        row = dict(section=k, name=checklists[k]["name"], byRound={})
        for rl in allRounds:
            vals = [v for f in forms if f["round"] == rl and f["isFirst"] for (sk, _), v in f["answers"].items() if sk == k]
            y, n = vals.count("Yes"), vals.count("No")
            row["byRound"][rl] = _pct(y, y + n)
        sectionStats.append(row)

    # ---- round KPIs
    rosterCount = sum(1 for s in students if s["inRoster"])
    roundStats = {}
    for rl in allRounds:
        fr = [f for f in forms if f["round"] == rl]
        firsts = [f for f in fr if f["isFirst"]]
        finals = [f for f in fr if f["isFinal"]]
        comp = [f["compliance"] for f in firsts if f["compliance"] is not None]
        roundStats[rl] = dict(
            forms=len(fr), students=len(firsts), reaudits=len(fr) - len(firsts),
            notAudited=sum(1 for s in students if s["inRoster"] and rl not in s["rounds"]),
            firstOutcome={o: sum(f["outcome"] == o for f in firsts) for o in _OUTCOMES + [""]},
            finalOutcome={o: sum(f["outcome"] == o for f in finals) for o in _OUTCOMES + [""]},
            meanCompliance=_st.mean(comp) if comp else None, medianCompliance=_st.median(comp) if comp else None,
            fullCompliance=sum(c == 100 for c in comp), anyNo=sum(f["nNo"] > 0 for f in firsts),
            concerns=sum(f["concernFlag"].lower() == "yes" or bool(f["concerns"]) for f in fr),
            dates=(min(f["date"] for f in fr), max(f["date"] for f in fr)) if fr else None)

    # ---- assessors per audit (all forms)
    assessorStats = []
    for (ex, rl), fs in sorted(_groupBy(forms, lambda f: (f["assessor"], f["round"])).items(),
                               key=lambda kv: (allRounds.index(kv[0][1]), -len(kv[1]), kv[0][0])):
        comp = [f["compliance"] for f in fs if f["compliance"] is not None]
        nAns = sum(f["nYes"] + f["nNo"] + f["nNa"] for f in fs)
        assessorStats.append(dict(
            assessor=ex, round=rl, forms=len(fs), students=len({f["student"] for f in fs}),
            meanCompliance=_st.mean(comp) if comp else None,
            anyNo=_pct(sum(f["nNo"] > 0 for f in fs), len(fs)),
            outcome={o: _pct(sum(f["outcome"] == o for f in fs), len(fs)) for o in _OUTCOMES},
            naRate=_pct(sum(f["nNa"] for f in fs), nAns),
            commentRate=_pct(sum(bool(f["comments"]) or any(f["sectionComments"].values()) for f in fs), len(fs))))

    multi = [(s["name"], rl, rd["n"]) for s in students for rl, rd in s["rounds"].items() if rd["n"] > 1]
    return dict(cohort=cohort, forms=forms, excluded=excluded, students=students, items=items, itemKeys=itemKeys,
                itemStats=itemStats, sectionStats=sectionStats, roundStats=roundStats, assessorStats=assessorStats,
                rounds=allRounds, roundDefs=rounds, outsideCount=outside, sections=sections, itemSections=itemSections,
                summaryKey=summaryKey, checklists=checklists, rosterCount=rosterCount, multiForms=multi,
                notInRoster=sorted({f["student"] for f in forms if not f["inRoster"]}),
                subjects=sorted({f["subject"] for f in forms if f["subject"]}))


def _groupBy(seq, key):
    out = _collections.defaultdict(list)
    for x in seq: out[key(x)].append(x)
    return out


# ================================================================ styling primitives
def _fill(h):
    return _XlPatternFill("solid", start_color=h, end_color=h)


def _hdr(c, text, bg=_NAVY, fg="FFFFFF", sz=10):
    c.value = text
    c.font = _XlFont(bold=True, color=fg, size=sz)
    c.fill = _fill(bg)
    c.alignment = _XlAlignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = _BORDER


def _cell(ws, r, col, v, bold=False, numFmt=None, bg=None, fg=_INK, align="center", wrap=False, sz=9):
    c = ws.cell(r, col, v)
    c.font = _XlFont(bold=bold, color=fg, size=sz)
    c.alignment = _XlAlignment(horizontal=align, vertical="center" if not wrap else "top", wrap_text=wrap)
    c.border = _BORDER
    if bg: c.fill = _fill(bg)
    if numFmt: c.number_format = numFmt
    return c


def _ansCell(ws, r, col, v):
    bg, fg = _ANS_FILL.get(v, ("FFFFFF", _INK))
    return _cell(ws, r, col, v or "—", bg=bg, fg=fg, bold=v == "No")


def _outCell(ws, r, col, v):
    bg, fg = _OUT_FILL.get(v, ("FFFFFF", _INK))
    return _cell(ws, r, col, v or "—", bg=bg, fg=fg, bold=v == "No")


def _compCell(ws, r, col, v):
    c = _cell(ws, r, col, None if v is None else v / 100, numFmt="0%")
    if v is not None and v < AUDIT_LOW_COMPLIANCE_PCT:
        c.fill, c.font = _fill(_OUT_FILL["No"][0]), _XlFont(bold=True, color=_OUT_FILL["No"][1], size=9)
    return c


def _title(ws, text, note=None, width=12):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(1, 1, text)
    c.font, c.fill = _XlFont(bold=True, size=14, color="FFFFFF"), _fill(_NAVY)
    c.alignment = _XlAlignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 26
    if note:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        n = ws.cell(2, 1, note)
        n.font, n.alignment = _XlFont(italic=True, size=9, color="595959"), _XlAlignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 30


def _widths(ws, widths):
    for i, w in enumerate(widths, 1): ws.column_dimensions[_xlColLetter(i)].width = w


def _fmtDate(d):
    return d.strftime("%d/%m/%Y") if d else None


# ================================================================ sheets
def _writeSummary(ws, m):
    rs, rounds = m["roundStats"], m["rounds"]
    _title(ws, f"{m['cohort']} Clinical Record Audit — Summary",
           f"Subject {', '.join(m['subjects'])} · roster {m['rosterCount']} students · generated "
           f"{_dt.date.today():%d/%m/%Y}. Compliance and outcome counts use each student's FIRST audit in an audit; "
           f"'Final outcome' uses the latest (after any re-audit).", width=2 + len(rounds))
    r = 4
    _hdr(ws.cell(r, 1), "Measure"); [_hdr(ws.cell(r, 2 + i), rl) for i, rl in enumerate(rounds)]
    rows = [
        ("Audit dates", lambda s: f"{_fmtDate(s['dates'][0])} – {_fmtDate(s['dates'][1])}" if s["dates"] else "—", None),
        ("Students audited", lambda s: s["students"], "0"),
        ("Roster students not audited", lambda s: s["notAudited"], "0"),
        ("Forms submitted", lambda s: s["forms"], "0"),
        ("Re-audit forms", lambda s: s["reaudits"], "0"),
        ("Mean compliance (first audit)", lambda s: None if s["meanCompliance"] is None else s["meanCompliance"] / 100, "0.0%"),
        ("Median compliance (first audit)", lambda s: None if s["medianCompliance"] is None else s["medianCompliance"] / 100, "0.0%"),
        ("Students at 100% compliance", lambda s: s["fullCompliance"], "0"),
        ("Students with ≥1 'No' item", lambda s: s["anyNo"], "0"),
        ("First outcome — Yes", lambda s: s["firstOutcome"]["Yes"], "0"),
        ("First outcome — Yes but action required", lambda s: s["firstOutcome"]["Yes but action required"], "0"),
        ("First outcome — No", lambda s: s["firstOutcome"]["No"], "0"),
        ("Final outcome — Yes", lambda s: s["finalOutcome"]["Yes"], "0"),
        ("Final outcome — Yes but action required", lambda s: s["finalOutcome"]["Yes but action required"], "0"),
        ("Final outcome — No", lambda s: s["finalOutcome"]["No"], "0"),
        ("Forms with coordinator concern", lambda s: s["concerns"], "0"),
    ]
    for label, fn, fmt in rows:
        r += 1
        _cell(ws, r, 1, label, align="left", bold=True)
        for i, rl in enumerate(rounds):
            _cell(ws, r, 2 + i, fn(rs[rl]), numFmt=fmt)

    # part-level compliance
    r += 2
    _hdr(ws.cell(r, 1), "Checklist part — compliance % (first audit)")
    [_hdr(ws.cell(r, 2 + i), rl) for i, rl in enumerate(rounds)]
    for s in m["sectionStats"]:
        r += 1
        _cell(ws, r, 1, s["name"], align="left")
        for i, rl in enumerate(rounds): _compCell(ws, r, 2 + i, s["byRound"][rl])

    # outcome chart data (first attempts) + stacked bar
    r += 2
    chartTop = r
    _hdr(ws.cell(r, 1), "First outcome (students)"); [_hdr(ws.cell(r, 2 + i), o) for i, o in enumerate(_OUTCOMES)]
    for rl in rounds:
        r += 1
        _cell(ws, r, 1, rl, align="left")
        for i, o in enumerate(_OUTCOMES): _cell(ws, r, 2 + i, rs[rl]["firstOutcome"][o])
    ch = _XlBarChart(); ch.type, ch.grouping, ch.overlap = "bar", "percentStacked", 100
    ch.title, ch.height, ch.width = "Audit outcome (first audit)", 6 + 0.6 * len(rounds), 16
    ch.add_data(_XlReference(ws, min_col=2, max_col=1 + len(_OUTCOMES), min_row=chartTop, max_row=r), titles_from_data=True)
    ch.set_categories(_XlReference(ws, min_col=1, min_row=chartTop + 1, max_row=r))
    for s, o in zip(ch.series, _OUTCOMES):
        s.graphicalProperties.solidFill = _OUT_CHART[o]; s.graphicalProperties.line.solidFill = _OUT_CHART[o]
    ch.dataLabels = _XlDataLabelList(); ch.dataLabels.showVal = True
    ch.y_axis.majorGridlines = None; ch.legend.position = "b"
    ch.y_axis.delete = False; ch.x_axis.delete = False
    ws.add_chart(ch, f"{_xlColLetter(4 + len(rounds))}4")
    _widths(ws, [44] + [18] * max(len(rounds), len(_OUTCOMES)))
    ws.freeze_panes = "B5"


def _writeStudentResults(ws, m):
    rounds = m["rounds"]
    per = ["Date", "Assessor", "Forms", "First outcome", "Final outcome", "Compliance %", "No items", "Items marked No"]
    nCols = 3 + len(per) * len(rounds) + 2
    _title(ws, f"{m['cohort']} Audit — Student Results",
           "One row per roster student (surname order). Compliance % and 'Items marked No' are from the FINAL form "
           "in each audit; 'First outcome' is the first audit before any re-audit. Change = last audit − first audit "
           "compliance (points).", width=nCols)
    r1, r2 = 4, 5
    for c, t in enumerate(["#", "Student No", "Student"], 1):
        ws.merge_cells(start_row=r1, start_column=c, end_row=r2, end_column=c); _hdr(ws.cell(r1, c), t)
    col = 4
    for i, rl in enumerate(rounds):
        ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + len(per) - 1)
        _hdr(ws.cell(r1, col), rl, bg=_BLUE if i % 2 == 0 else _NAVY)
        for j, p in enumerate(per): _hdr(ws.cell(r2, col + j), p, bg=_SUBHDR, fg=_INK, sz=9)
        col += len(per)
    for t in ["Change (pts)", "Flags"]:
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col); _hdr(ws.cell(r1, col), t); col += 1

    r = r2
    for n, s in enumerate(m["students"], 1):
        r += 1
        _cell(ws, r, 1, n); _cell(ws, r, 2, s["no"]); _cell(ws, r, 3, s["name"], align="left")
        col = 4
        for rl in rounds:
            rd = s["rounds"].get(rl)
            if not rd:
                for j in range(len(per)): _cell(ws, r, col + j, "Not audited" if j == 0 else None, fg="A6A6A6", bg="F2F2F2")
            else:
                fl = rd["final"]
                _cell(ws, r, col, _fmtDate(fl["date"]))
                _cell(ws, r, col + 1, " / ".join(dict.fromkeys(f["assessor"] for f in rd["forms"])), align="left")
                _cell(ws, r, col + 2, rd["n"], bg=_OUT_FILL["Yes but action required"][0] if rd["n"] > 1 else None)
                _outCell(ws, r, col + 3, rd["first"]["outcome"]); _outCell(ws, r, col + 4, fl["outcome"])
                _compCell(ws, r, col + 5, fl["compliance"])
                _cell(ws, r, col + 6, fl["nNo"], bold=fl["nNo"] > 0, fg=_OUT_FILL["No"][1] if fl["nNo"] else _INK)
                _cell(ws, r, col + 7, "; ".join(fl["noItems"]), align="left", wrap=False)
            col += len(per)
        ch = s["change"]
        c = _cell(ws, r, col, None if ch is None else round(ch, 1), numFmt="+0.0;-0.0;0.0")
        if ch is not None and ch <= -_DECLINE_PTS: c.fill = _fill(_OUT_FILL["No"][0])
        _cell(ws, r, col + 1, "; ".join(s["flags"]), align="left",
              bg=_OUT_FILL["Yes but action required"][0] if s["flags"] else None)
    _widths(ws, [5, 11, 26] + [11, 18, 7, 13, 13, 12, 9, 40] * len(rounds) + [11, 60])
    ws.freeze_panes = ws.cell(r2 + 1, 4)
    ws.auto_filter.ref = f"A{r2}:{_xlColLetter(nCols)}{r}"


def _writeItemMatrix(ws, m):
    items = m["items"]
    lead = ["Form id", "Audit", "Date", "Student No", "Student", "Assessor"]
    tail = ["Compliance %", "Yes", "No", "N/A"]
    nCols = len(lead) + len(items) + len(tail)
    _title(ws, f"{m['cohort']} Audit — Item Matrix (every submitted form)",
           "Yes = blue · No = orange · N/A = grey · — = unanswered. Item headers carry the checklist part number.",
           width=nCols)
    r1, r2 = 4, 5
    for c, t in enumerate(lead, 1):
        ws.merge_cells(start_row=r1, start_column=c, end_row=r2, end_column=c); _hdr(ws.cell(r1, c), t)
    col = len(lead) + 1
    for i, (k, grp) in enumerate(_groupConsecutive(items, lambda it: it["section"])):
        ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + len(grp) - 1)
        _hdr(ws.cell(r1, col), grp[0]["sectionName"], bg=_BLUE if i % 2 == 0 else _NAVY, sz=9)
        for j, it in enumerate(grp): _hdr(ws.cell(r2, col + j), it["text"], bg=_SUBHDR, fg=_INK, sz=8)
        col += len(grp)
    for t in tail:
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col); _hdr(ws.cell(r1, col), t); col += 1
    ws.row_dimensions[r2].height = 60

    r = r2
    order = {s["name"]: i for i, s in enumerate(m["students"])}
    for f in sorted(m["forms"], key=lambda f: (m["rounds"].index(f["round"]), order.get(f["student"], 9e9), f["attempt"])):
        r += 1
        vals = [f["id"], f["round"], _fmtDate(f["date"]), f["studentNo"], f["student"], f["assessor"]]
        for c, v in enumerate(vals, 1): _cell(ws, r, c, v, align="left" if c in (5, 6) else "center")
        col = len(lead) + 1
        for key in m["itemKeys"]: _ansCell(ws, r, col, f["answers"][key]); col += 1
        _compCell(ws, r, col, f["compliance"])
        _cell(ws, r, col + 1, f["nYes"]); _cell(ws, r, col + 2, f["nNo"]); _cell(ws, r, col + 3, f["nNa"])
    _widths(ws, [8, 9, 11, 11, 24, 18] + [11] * len(items) + [12, 6, 6, 6])
    ws.freeze_panes = ws.cell(r2 + 1, 6)
    ws.auto_filter.ref = f"A{r2}:{_xlColLetter(nCols)}{r}"


def _groupConsecutive(seq, key):
    out = []
    for x in seq:
        if out and key(out[-1][1][0]) == key(x): out[-1][1].append(x)
        else: out.append((key(x), [x]))
    return out


def _writeItemAnalysis(ws, m):
    rounds = m["rounds"]
    per = ["Yes", "No", "N/A", "% Yes", "% No", "% N/A"]
    nCols = 3 + len(per) * len(rounds)
    _title(ws, f"{m['cohort']} Audit — Item Analysis",
           "First audit per student per audit. % Yes / % No are of applicable answers (Yes + No); % N/A is of all "
           "answers. Lowest-compliance items are the teaching targets.", width=nCols)
    r1, r2 = 4, 5
    for c, t in enumerate(["Part", "Code", "Item"], 1):
        ws.merge_cells(start_row=r1, start_column=c, end_row=r2, end_column=c); _hdr(ws.cell(r1, c), t)
    col = 4
    for i, rl in enumerate(rounds):
        ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + len(per) - 1)
        _hdr(ws.cell(r1, col), rl, bg=_BLUE if i % 2 == 0 else _NAVY)
        for j, p in enumerate(per): _hdr(ws.cell(r2, col + j), p, bg=_SUBHDR, fg=_INK, sz=9)
        col += len(per)
    r = r2
    for it in m["itemStats"]:
        r += 1
        _cell(ws, r, 1, it["sectionName"], align="left"); _cell(ws, r, 2, f"{it['section']} {it['mc']}")
        _cell(ws, r, 3, it["text"], align="left")
        col = 4
        for rl in rounds:
            b = it["byRound"][rl]
            for j, v in enumerate([b["yes"], b["no"], b["na"]]): _cell(ws, r, col + j, v)
            for j, v in enumerate([b["pctYes"], b["pctNo"], b["pctNa"]]):
                _cell(ws, r, col + 3 + j, None if v is None else v / 100, numFmt="0%")
            col += len(per)
    # colour scales: % Yes (orange→blue), % No (blue→orange)
    for i in range(len(rounds)):
        base = 4 + i * len(per)
        ws.conditional_formatting.add(f"{_xlColLetter(base + 3)}{r2 + 1}:{_xlColLetter(base + 3)}{r}",
                                      _XlColorScaleRule(start_type="num", start_value=0.7, start_color="F8CBAD",
                                                        end_type="num", end_value=1, end_color="DDEBF7"))
        ws.conditional_formatting.add(f"{_xlColLetter(base + 4)}{r2 + 1}:{_xlColLetter(base + 4)}{r}",
                                      _XlColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                        end_type="num", end_value=0.2, end_color="F8CBAD"))
    # % No chart, one series per audit
    ch = _XlBarChart(); ch.type = "bar"
    ch.title, ch.height, ch.width = "% 'No' by item (first audit)", max(8, 0.55 * len(m["itemStats"])), 18
    for i, rl in enumerate(rounds):
        c = 4 + i * len(per) + 4
        ch.add_data(_XlReference(ws, min_col=c, min_row=r2 + 1, max_row=r), titles_from_data=False)
        ch.series[-1].tx = _XlSeriesLabel(v=rl)
        colr = [_BLUE, _ORANGE, "7F7F7F"][i % 3]
        ch.series[-1].graphicalProperties.solidFill = colr; ch.series[-1].graphicalProperties.line.solidFill = colr
    ch.set_categories(_XlReference(ws, min_col=3, min_row=r2 + 1, max_row=r))
    ch.x_axis.scaling.orientation = "maxMin"; ch.y_axis.number_format = "0%"; ch.y_axis.majorGridlines = None
    ch.y_axis.delete = False; ch.x_axis.delete = False; ch.legend.position = "b"
    ws.add_chart(ch, f"{_xlColLetter(nCols + 2)}4")
    _widths(ws, [30, 11, 32] + [6, 6, 6, 8, 8, 8] * len(rounds))
    ws.freeze_panes = ws.cell(r2 + 1, 4)


def _writeAssessors(ws, m):
    heads = ["Audit", "Assessor", "Forms", "Students", "Mean compliance", "% forms with a No",
             "% Outcome Yes", "% Yes but action req.", "% Outcome No", "N/A rate", "% forms with comments"]
    _title(ws, f"{m['cohort']} Audit — Assessor Analysis",
           "All submitted forms (re-audits included). Large differences in compliance / outcome mix or N/A use "
           "between assessors in the same audit suggest calibration is needed.", width=len(heads))
    for c, t in enumerate(heads, 1): _hdr(ws.cell(4, c), t)
    r = 4
    for a in m["assessorStats"]:
        r += 1
        _cell(ws, r, 1, a["round"]); _cell(ws, r, 2, a["assessor"], align="left")
        _cell(ws, r, 3, a["forms"]); _cell(ws, r, 4, a["students"])
        _compCell(ws, r, 5, a["meanCompliance"])
        vals = [a["anyNo"], a["outcome"]["Yes"], a["outcome"]["Yes but action required"], a["outcome"]["No"],
                a["naRate"], a["commentRate"]]
        for j, v in enumerate(vals): _cell(ws, r, 6 + j, None if v is None else v / 100, numFmt="0%")
    _widths(ws, [10, 22, 8, 9, 12, 12, 11, 13, 11, 9, 13])
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{_xlColLetter(len(heads))}{r}"


def _writeComments(ws, m):
    secs = m["sections"]
    heads = ["Audit", "Date", "Student No", "Student", "Assessor", "Attempt", "Outcome"] + \
            [m["checklists"][k]["name"].split(":")[0] + " comment" for k in secs] + \
            ["Additional comments", "Coordinator concern", "Concern details"]
    _title(ws, f"{m['cohort']} Audit — Comments", "Forms with at least one comment or concern.", width=len(heads))
    for c, t in enumerate(heads, 1): _hdr(ws.cell(4, c), t)
    r = 4
    order = {s["name"]: i for i, s in enumerate(m["students"])}
    for f in sorted(m["forms"], key=lambda f: (m["rounds"].index(f["round"]), order.get(f["student"], 9e9), f["attempt"])):
        if not (f["comments"] or f["concerns"] or f["concernFlag"].lower() == "yes" or any(f["sectionComments"].values())):
            continue
        r += 1
        vals = [f["round"], _fmtDate(f["date"]), f["studentNo"], f["student"], f["assessor"],
                f"{f['attempt']} of {f['nAttempts']}"]
        for c, v in enumerate(vals, 1): _cell(ws, r, c, v, align="left" if c in (4, 5) else "center")
        _outCell(ws, r, 7, f["outcome"])
        for j, k in enumerate(secs): _cell(ws, r, 8 + j, f["sectionComments"][k] or None, align="left", wrap=True)
        c0 = 8 + len(secs)
        _cell(ws, r, c0, f["comments"] or None, align="left", wrap=True)
        _cell(ws, r, c0 + 1, f["concernFlag"] or None,
              bg=_OUT_FILL["No"][0] if f["concernFlag"].lower() == "yes" else None)
        _cell(ws, r, c0 + 2, f["concerns"] or None, align="left", wrap=True)
    _widths(ws, [9, 11, 11, 22, 18, 9, 14] + [30] * len(secs) + [50, 11, 40])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{_xlColLetter(len(heads))}{r}"


def _writeNotAudited(ws, m):
    _title(ws, f"{m['cohort']} Audit — Not Audited", "Roster students with no submitted audit form in the audit.", width=4)
    for c, t in enumerate(["Audit", "#", "Student No", "Student"], 1): _hdr(ws.cell(4, c), t)
    r = 4
    for rl in m["rounds"]:
        miss = [s for s in m["students"] if s["inRoster"] and rl not in s["rounds"]]
        for i, s in enumerate(miss, 1):
            r += 1
            _cell(ws, r, 1, rl); _cell(ws, r, 2, i); _cell(ws, r, 3, s["no"]); _cell(ws, r, 4, s["name"], align="left")
    _widths(ws, [10, 6, 12, 30])
    ws.freeze_panes = "A5"


def _writeDataQuality(ws, m):
    _title(ws, f"{m['cohort']} Audit — Data Quality", "Excluded forms, re-audits and roster mismatches.", width=6)
    r = 4
    _hdr(ws.cell(r, 1), "Excluded forms"); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    for c, t in enumerate(["Form id", "Date", "Student", "Assessor", "Reason"], 1): _hdr(ws.cell(r, c), t, bg=_SUBHDR, fg=_INK)
    for e in m["excluded"]:
        r += 1
        for c, k in enumerate(["id", "date", "student", "assessor", "reason"], 1): _cell(ws, r, c, e[k], align="left")
    r += 2
    _hdr(ws.cell(r, 1), "Students with more than one form in an audit (re-audits)")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    for c, t in enumerate(["Student", "Audit", "Forms"], 1): _hdr(ws.cell(r, c), t, bg=_SUBHDR, fg=_INK)
    for name, rl, n in m["multiForms"]:
        r += 1
        _cell(ws, r, 1, name, align="left"); _cell(ws, r, 2, rl); _cell(ws, r, 3, n)
    r += 2
    _hdr(ws.cell(r, 1), "Audited students not matched to the cohort roster")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    for n in m["notInRoster"] or ["(none)"]:
        r += 1; _cell(ws, r, 1, n, align="left")
    r += 2
    _cell(ws, r, 1, f"Forms outside the selected audit date ranges (not in this report): {m['outsideCount']}",
          align="left", bold=True)
    _widths(ws, [28, 12, 26, 22, 34])


def _writeDescription(ws, m):
    _title(ws, f"{m['cohort']} Audit — Method & Checklist", width=5)
    notes = [
        ("Source", "DASH `audit` forms (JSON export). Only submitted forms for the cohort; drafts, repeated ids, "
                   "test/dummy names and excluded student numbers are removed (see Data Quality)."),
        ("Audits", "; ".join(f"{l}: {_fmtDate(a)} – {_fmtDate(b)}" for l, (a, b) in m["roundDefs"].items()) +
                   ". Selected with buildAuditReport(audits=...); defaults in AUDIT_ROUNDS in audit_utils.py."),
        ("Answers", "'Not Applicable' and 'N/A' are treated as the same answer (N/A)."),
        ("Compliance %", "Yes ÷ (Yes + No) across checklist items in Parts other than the Summary. N/A and blank "
                         f"are excluded. Values below {AUDIT_LOW_COMPLIANCE_PCT:.0f}% are highlighted."),
        ("Outcome", f"The {m['summaryKey']} Summary answer: Yes / Yes but action required / No. A 'No' triggers a "
                    "re-audit within the same audit."),
        ("First vs Final", "First = earliest form for a student in an audit (used for cohort statistics, item and "
                           "part analysis). Final = latest form (used on Student Results)."),
        ("Flags", "Not audited · Outcome No (re-audited or not) · Action required · Still No after re-audit · "
                  f"Compliance < {AUDIT_LOW_COMPLIANCE_PCT:.0f}% · Coordinator concern · Compliance fell ≥ "
                  f"{_DECLINE_PTS:.0f} pts between audits · Not on roster."),
    ]
    r = 3
    for k, v in notes:
        _cell(ws, r, 1, k, bold=True, align="left", bg=_SUBHDR)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        _cell(ws, r, 2, v, align="left", wrap=True); ws.row_dimensions[r].height = 32; r += 1
    r += 1
    for c, t in enumerate(["Part", "Code", "Item", "Options", "Rubric / what the assessor checks"], 1): _hdr(ws.cell(r, c), t)
    for k in m["sections"]:
        cl = m["checklists"][k]
        opts = " / ".join(cl["extra_config"].get("options", {}).get("assessor", {}).values())
        for mc, txt in cl["fields"].items():
            r += 1
            rub = _stripHtml(cl["extra_config"].get("rubric", {}).get(mc, {}).get("label", ""))
            for c, v in enumerate([cl["name"], f"{k} {mc}", txt, opts, rub], 1):
                _cell(ws, r, c, v, align="left", wrap=True)
            ws.row_dimensions[r].height = min(15 * max(1, rub.count("\n") + 1), 300)
    _widths(ws, [30, 11, 28, 24, 90])


# ================================================================ build
def buildAuditReport(jsonPath, outPath=None, cohort="DDS3", rosterPath="studentEmailList.csv", audits=None,
                     summaryKey=None):
    """Build the audit workbook for one cohort and the selected audits (see resolveAudits). Returns the model.

    outPath None -> <cohort>/Audit/<cohort>_Audit_Report_<year>[ (<audit labels>)].xlsx — the labels are added
    when only some of the configured audits (or custom ranges) are selected.
    """
    import os
    m = loadAudit(jsonPath, cohort=cohort, rosterPath=rosterPath, audits=audits, summaryKey=summaryKey)
    if not m["forms"]: raise ValueError(f"No submitted {cohort} forms inside {list(m['rounds'])}")
    year = m["forms"][0]["date"].year
    if outPath is None:
        allCfg = list(AUDIT_ROUNDS.get(cohort) or {})
        suffix = "" if (audits is None or list(m["rounds"]) == allCfg) else f" ({', '.join(m['rounds'])})"
        outPath = f"{cohort}/Audit/{cohort}_Audit_Report_{year}{suffix}.xlsx"
    if os.path.dirname(outPath): os.makedirs(os.path.dirname(outPath), exist_ok=True)

    wb = _XlWorkbook()
    sheets = [("Summary", _writeSummary), ("Student Results", _writeStudentResults),
              ("Item Analysis", _writeItemAnalysis), ("Item Matrix", _writeItemMatrix),
              ("Assessors", _writeAssessors), ("Comments", _writeComments), ("Not Audited", _writeNotAudited),
              ("Data Quality", _writeDataQuality), ("Method & Checklist", _writeDescription)]
    for i, (name, fn) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.sheet_view.showGridLines = False
        fn(ws, m)
    wb.save(outPath)

    rs = m["roundStats"]
    print(f"{cohort} audit [{', '.join(m['rounds'])}]: {len(m['forms'])} forms · {len(m['excluded'])} excluded · "
          f"{m['outsideCount']} outside selected dates · roster {m['rosterCount']}")
    for rl in m["rounds"]:
        s = rs[rl]
        mc = "—" if s["meanCompliance"] is None else f"{s['meanCompliance']:.1f}%"
        print(f"  {rl}: {s['students']} students audited, {s['notAudited']} not audited, {s['reaudits']} re-audit(s), "
              f"mean compliance {mc}, first outcome {s['firstOutcome']['Yes']}/"
              f"{s['firstOutcome']['Yes but action required']}/{s['firstOutcome']['No']} (Yes/Action/No)")
    print(f"Saved -> {outPath}")
    return m
