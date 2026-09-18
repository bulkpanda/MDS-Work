"""Seed the BOH2 reference workbooks: SHY session schedule + student roster.

    python make_boh2_reference.py BOH2      # writes into the BOH2 folder
    python make_boh2_reference.py .         # writes into the current folder

Creates the two files that `loadSessionSchedule('BOH2')` and `loadRoster('BOH2')` read:

    BOH2/BOH2 Sim Sessions 2026.xlsx   week_no, date, weekday, task, is_choice_week,
                                       choice_options, counts_toward_total, notes
    BOH2/BOH2 Roster 2026.xlsx         sort_order, student_number, first_name, last_name,
                                       status, supplied_student_number, notes

WHEN YOU NEED THIS
------------------
Normally you don't — **edit the .xlsx files directly in Excel**, which is the whole point of
keeping them as spreadsheets. Reach for this script only to:
  * regenerate from scratch after the files are lost or badly mangled,
  * seed the same pair for another cohort (copy SESSIONS / ROSTER below and change the
    cohort string — the loaders are cohort-generic, only the data here is BOH2's),
  * start a new academic year.

RUNNING IT OVERWRITES the files, discarding any edits made in Excel. Nothing else reads
this script; general_utils only reads the spreadsheets.

DATA NOTES (from the schedule and roster supplied 2026-08-17, verified against the DB)
--------------------------------------------------------------------------------------
* 15 sessions carry counts_toward_total=1, matching the stated "15 checklists by year end".
  Week 37 is the break and Week 43 is a spare/extension task, so neither counts.
* Choice weeks are 27 and 38 — the covering note said "30 and 38", but the schedule table
  and the data both show the two-checklist week at 27 (30 June: 36MO x28, 41MIBL x12).
* Jackie Tran: the supplied roster said 1745117, DASH has 1475117 (digits transposed).
  student_number holds the DASH value so the join finds her; supplied_student_number keeps
  the original. If DASH is corrected, change both here.
* 1234567 "Test1 Student1" appears in the BOH2 sim data and is deliberately NOT on the
  roster — reports append and flag unknown students rather than dropping them.
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY, WHITE, GREY, AMBER = "010D44", "FFFFFF", "F2F2F2", "FFEB9C"


def _style(ws, nCols, freeze="A2", widths=None):
    thin = Side(style="thin", color="D9D9D9")
    for c in range(1, nCols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", start_color=NAVY, end_color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = freeze
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w


# ── Sessions ─────────────────────────────────────────────────────────────────
# week_no, date (None = no session), task, is_choice, options, counts_toward_total
SESSIONS = [
    (27, "2026-06-30", "41MIBL (524 578) or 36MO (532)", True,
     "41MIBL (524 578) | 36MO (532)", 1),
    (28, "2026-07-07", "14MODB (534 577) (preparation)", False, "", 1),
    (29, "2026-07-14", "14MODB (534 577) (restoration)", False, "", 1),
    (30, "2026-07-21", "16MODB (534 577)", False, "", 1),
    (31, "2026-07-28", "26O (531) & 64DO 65MO (532)", False, "", 1),
    (32, "2026-08-04", "51B DP (525)", False, "", 1),
    (33, "2026-08-11", "85 (587) & 75O (414 531)", False, "", 1),
    (34, "2026-08-18", "55MO (414 627 586)", False, "", 1),
    (35, "2026-08-25", "75 (414 627 586) & 55 (655)", False, "", 1),
    (36, "2026-09-01", "75 (414 627) and Q3 (222)", False, "", 1),
    (37, None, "BREAK", False, "", 0),
    (38, "2026-09-15", "Q4 (114 222) & 85MO (414 627 586) or Q4 (114 222) & 54DO 55MO (532)",
     True, "Q4 (114 222) & 85MO (414 627 586) | Q4 (114 222) & 54DO 55MO (532)", 1),
    (39, "2026-09-22", "11MIBP (414 579) & 71 53 55 (311)", False, "", 1),
    (40, "2026-09-29", "21 (386) & 71 53 55 (311)", False, "", 1),
    (41, "2026-10-06", "61 53 75 (311)", False, "", 1),
    (42, "2026-10-13", "61 53 75 (311)", False, "", 1),
    (43, "2026-10-20", "", False, "", 0),   # spare / extension task, not in the 15
]

sessions = pd.DataFrame(SESSIONS, columns=[
    "week_no", "date", "task", "is_choice_week", "choice_options", "counts_toward_total"])
sessions.insert(0, "cohort", "BOH2")
sessions.insert(1, "period", "SHY")
sessions["weekday"] = pd.to_datetime(sessions["date"]).dt.strftime("%a")
sessions["notes"] = ""
sessions.loc[sessions.week_no == 37, "notes"] = "Mid-semester break, no session"
sessions.loc[sessions.week_no == 43, "notes"] = "Spare / extension task — not part of the 15"
sessions.loc[sessions.is_choice_week, "notes"] = (
    "Students choose ONE of two tasks; expect two checklist codes on this date")
sessions = sessions[["cohort", "period", "week_no", "date", "weekday", "task",
                     "is_choice_week", "choice_options", "counts_toward_total", "notes"]]

# ── Roster (as supplied 2026-08-17) ──────────────────────────────────────────
ROSTER = [
    (1746763, "Tala", "Aktam Anjrini"), (1746782, "Riana", "Amit"),
    (1775573, "Amal", "Barakat"), (1756996, "Thy", "Bui"),
    (1608030, "Zipei", "Cheng"), (1756811, "Cadence", "Chun"),
    (1615593, "Thanh", "Do"), (1756115, "Angelyna", "Eng"),
    (1606071, "Halima", "Fiqi"), (1734402, "Tara", "Fitzgerald"),
    (1756193, "Jecka Angela", "Gloria"), (1757718, "Bianca", "Goldwyn"),
    (1779552, "Khatheeja", "Hidhayathulla"), (1757345, "Lexy Mhaye San Pablo", "Huang"),
    (1605779, "Melissa", "Huynh"), (1613317, "Yejin", "Kim"),
    (1757036, "Regan", "Kwok"), (1353219, "Joane", "Lam"),
    (1640181, "Jamie", "Lao"), (1746298, "Amy", "Le"),
    (1746181, "Henry", "Le"), (1746153, "Anna", "Ma"),
    (1472831, "Usman", "Majoo"), (1616141, "Kate", "Mclennan Arnott"),
    (1746975, "Sundus", "Mohamed"), (1638283, "Stephane", "Ndiyamba"),
    (1615131, "Katie", "Nguyen"), (1758681, "Ley", "Nguyen"),
    (1758012, "Teresa", "Nguyen"), (1634259, "Hilda", "Petros"),
    (1755988, "Chelsea", "Pham"), (1088656, "Linda", "Pham"),
    (1748772, "Avaya", "Pola"), (1760557, "Elizabeth", "Puthumana"),
    (1321463, "Yuvarnika", "Ramesh"), (1757029, "Rem", "Said"),
    (1634242, "James", "Segalla"), (1756610, "Chloe", "Shi"),
    (1758179, "Anda", "Sudampanthorn"), (1606301, "Lea", "Tabbara"),
    (1749315, "Tasnim", "Tanisha"), (1616618, "Bubby", "Tereva"),
    (1475117, "Jackie", "Tran"), (1746168, "My Linh (Jocelyn)", "Tran"),
    (1452737, "Nina", "Tran"), (1535538, "Qiwen", "Xue"),
    (1764857, "Kai Xin", "Ye"), (1746202, "Selina", "Yu"),
    (1746304, "Heifa", "Zeki"),
]

roster = pd.DataFrame(ROSTER, columns=["student_number", "first_name", "last_name"])
roster.insert(0, "cohort", "BOH2")
roster["status"] = "active"
roster["notes"] = ""
# The supplied list had 1745117 for Jackie Tran; DASH has 1475117 (digits transposed).
# The DASH value is used so the join works; the supplied value is kept for traceability.
roster.loc[roster.student_number == 1475117, "notes"] = (
    "Roster supplied 1745117; DASH has 1475117 (transposed). Using the DASH value so the "
    "join finds her — correct in DASH or here, whichever is wrong.")
roster["supplied_student_number"] = roster["student_number"]
roster.loc[roster.student_number == 1475117, "supplied_student_number"] = 1745117
roster = roster.sort_values(["last_name", "first_name"], kind="stable").reset_index(drop=True)
roster.insert(1, "sort_order", range(1, len(roster) + 1))
roster = roster[["cohort", "sort_order", "student_number", "first_name", "last_name",
                 "status", "supplied_student_number", "notes"]]

if __name__ == "__main__":
    import sys
    outDir = sys.argv[1] if len(sys.argv) > 1 else "."

    p1 = f"{outDir}/BOH2 Sim Sessions 2026.xlsx"
    with pd.ExcelWriter(p1, engine="openpyxl") as w:
        sessions.to_excel(w, sheet_name="sessions", index=False)
        _style(w.sheets["sessions"], len(sessions.columns), "D2",
               {"A": 8, "B": 8, "C": 9, "D": 12, "E": 9, "F": 62, "G": 14, "H": 62,
                "I": 12, "J": 52})
        for r in range(2, len(sessions) + 2):
            if w.sheets["sessions"].cell(row=r, column=7).value:      # is_choice_week
                for c in range(1, len(sessions.columns) + 1):
                    w.sheets["sessions"].cell(row=r, column=c).fill = PatternFill(
                        "solid", start_color=AMBER, end_color=AMBER)
    print(f"wrote {p1}  ({len(sessions)} rows, "
          f"{int(sessions.counts_toward_total.sum())} counting sessions)")

    p2 = f"{outDir}/BOH2 Roster 2026.xlsx"
    with pd.ExcelWriter(p2, engine="openpyxl") as w:
        roster.to_excel(w, sheet_name="roster", index=False)
        _style(w.sheets["roster"], len(roster.columns), "C2",
               {"A": 8, "B": 11, "C": 16, "D": 24, "E": 20, "F": 10, "G": 22, "H": 70})
    print(f"wrote {p2}  ({len(roster)} students)")
