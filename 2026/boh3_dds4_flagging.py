"""
boh3_dds4_flagging.py
Cohort-level "lagging student" flagging for the final-year cohorts (DDS4 / BOH3).

STAGE 1: DATA LAYER + METRICS TABLE.
STAGE 1b: SECTION-LEVEL entrustment / weakness, cohort-relative status + colour pivots
          (buildSectionPerformance, pivotSectionPerformance, writeSectionPivotSheet).
STAGE 2: count flags (3 threshold modes), entrustment / weakness trajectory states, Summary +
         styled workbook -> buildLaggingFlagWorkbook (the deliverable).
    getCohortFormFrame      one row per form   (SQL, 1 query)
    getCohortItemFrame      one row per form x item code (SQL, 1 query)
    buildMetricsLong        per-student metrics, Overall + every rotation group (pandas)
    pivotMetrics            long -> wide (one row per student) for eyeballing
    getExpectedSections     which sections / sub-sections the cohort actually does
    cohortWindowStats       cohort mean / SD per metric+window and what each of the
                            three threshold modes resolves to (for tuning defaults)
    resolveThreshold        fixed | meanStd (mean - k*SD) | meanFrac (f*mean)
    buildFlaggingMetrics    convenience: runs all of the above, returns a dict

Reuses boh3_dds4_utils (weakness SQL fragments, entrustment CASE, rotationNumber,
PBN / leave readers) and Utils (readDf, section mapping) so numbers agree with the
existing PDFs / Summary_Details workbooks.

Usage in notebook:
    from boh3_dds4_flagging import *

Naming: everything a notebook cell calls is PUBLIC (no leading underscore) because
main.ipynb uses `from x import *`.
"""

import numpy as np
import pandas as pd

from Utils import readDf, _loadSectionMapping, _mergeSection
from boh3_dds4_utils import (
    WEAKNESS_KEY_LABELS, _weaknessCountSelect, _ENT_CASE_SQL, rotationNumber,
    getLeaveDaysByStudent, getPbnCountByStudent,
)

# ═══════════════════════════════════════════════════════════════════════════
# 0. Configuration defaults (all overridable per call)
# ═══════════════════════════════════════════════════════════════════════════

OVERALL_WINDOW = "Overall"

# Count-type metrics (forms, patients, age bands, sections): overall + these groups.
# Any number of groups; each is (label, [rotation numbers]).
DEFAULT_ROTATION_GROUPS = [
    ("R1-4", [1, 2, 3, 4]),
    ("R5+", list(range(5, 21))),
]

# Entrustment / weakness trajectory windows (3 non-overlapping points judge
# "consistent" better than 2). Used in stage 2; computed here so you can inspect.
DEFAULT_TRAJECTORY_GROUPS = [
    ("R1-3", [1, 2, 3]),
    ("R4-6", [4, 5, 6]),
    ("R7+", list(range(7, 21))),
]

# Students under this many forms (Overall) are left OUT of the cohort mean / SD so
# one near-empty record (long leave, withdrawn) cannot inflate every SD.
DEFAULT_MIN_FORMS_FLOOR = 20

# A section is "expected" for the cohort when at least this share of students
# (above the forms floor) have any activity in it.
DEFAULT_EXPECTED_MIN_SHARE = 0.80

# Never treated as expected / flaggable coverage buckets.
NON_FLAG_SECTIONS = ("Unmapped", "Miscellaneous")

WEAKNESS_LABELS = list(WEAKNESS_KEY_LABELS.values())

# metric -> domain (drives sheet grouping + which threshold family applies later)
DOMAIN_ACTIVITY = "Activity"
DOMAIN_PATIENTS = "Patients"
DOMAIN_AGE = "Age mix"
DOMAIN_ENTRUSTMENT = "Entrustment"
DOMAIN_WEAKNESS = "Weakness"
DOMAIN_COVERAGE = "Coverage"
DOMAIN_CONTEXT = "Context"

SECTION_PREFIX = "Section: "
SUBSECTION_PREFIX = "Sub-section: "


# ═══════════════════════════════════════════════════════════════════════════
# 1. Data layer (two cohort-wide queries)
# ═══════════════════════════════════════════════════════════════════════════

_PATIENT_ARRAY_SQL = ("jsonb_array_elements(CASE WHEN jsonb_typeof(f.patient_data)='array' "
                      "THEN f.patient_data ELSE '[]'::jsonb END)")

# safe int age: non-numeric ages become NULL instead of raising
_AGE_SQL = "CASE WHEN pd->>'patient_age' ~ '^\\d+$' THEN (pd->>'patient_age')::int END"


def getCohortFormFrame(engine, cohort, formsTable="dds4_boh3_forms_v3"):
    """ONE ROW PER FORM for the whole cohort — the single source for every non-item metric.

    Columns: assessmentId, studentNumber, studentName, date, rotation, rotationNum, clinic,
    assessor, submittedByStudent, submittedByAssessor, entrustment (1-4), readiness (1-4, self),
    one tag-count column per weakness category (WEAKNESS_LABELS), incidentTags, concernFlag,
    commendationTags, patientsAttended, patientsFta, age0to6, age7to17, age18plus.
    Patient / age counts follow getAgeCountsBatch / getPatientPerStudent (attended only for age).
    """
    sql = f"""
    SELECT
      f.assessmentid                                        AS "assessmentId",
      f.student_number                                      AS "studentNumber",
      f.student_name                                        AS "studentName",
      f.datetimeutc::date                                   AS "date",
      f.rotation                                            AS "rotation",
      COALESCE(NULLIF(f.external_clinic,''), f.clinic)      AS "clinic",
      f.assessor_name                                       AS "assessor",
      COALESCE(f.submitted_by_student, false)               AS "submittedByStudent",
      COALESCE(f.submitted_by_assessor, false)              AS "submittedByAssessor",
      {_ENT_CASE_SQL}                                       AS "entrustment",
      CASE NULLIF(f.student_data->'scales'->'scale-practice-readiness'->>'key','')
        WHEN 'S1' THEN 1 WHEN 'S2' THEN 2 WHEN 'S3' THEN 3 WHEN 'S4' THEN 4
      END::smallint                                         AS "readiness",
      {_weaknessCountSelect()},
      jsonb_array_length(COALESCE(f.assessor_data->'multi-select'->'clinical-incident','[]'::jsonb))
                                                            AS "incidentTags",
      (NULLIF(TRIM(f.additional_concerns),'') IS NOT NULL)::int AS "concernFlag",
      jsonb_array_length(COALESCE(f.assessor_data->'multi-select'->'strengths','[]'::jsonb))
                                                            AS "commendationTags",
      COALESCE(pt."patientsAttended", 0)                    AS "patientsAttended",
      COALESCE(pt."patientsFta", 0)                         AS "patientsFta",
      COALESCE(pt."age0to6", 0)                             AS "age0to6",
      COALESCE(pt."age7to17", 0)                            AS "age7to17",
      COALESCE(pt."age18plus", 0)                           AS "age18plus"
    FROM {formsTable} f
    LEFT JOIN LATERAL (
      SELECT
        COUNT(*) FILTER (WHERE (pd->>'patient_attended')::boolean)::int       AS "patientsAttended",
        COUNT(*) FILTER (WHERE NOT (pd->>'patient_attended')::boolean)::int   AS "patientsFta",
        COUNT(*) FILTER (WHERE (pd->>'patient_attended')::boolean AND {_AGE_SQL} BETWEEN 0 AND 6)::int  AS "age0to6",
        COUNT(*) FILTER (WHERE (pd->>'patient_attended')::boolean AND {_AGE_SQL} BETWEEN 7 AND 17)::int AS "age7to17",
        COUNT(*) FILTER (WHERE (pd->>'patient_attended')::boolean AND {_AGE_SQL} >= 18)::int            AS "age18plus"
      FROM {_PATIENT_ARRAY_SQL} pd
    ) pt ON TRUE
    WHERE f.cohort = :cohort AND f.student_name IS NOT NULL
    ORDER BY f.student_name, f.datetimeutc, f.assessmentid;
    """
    df = readDf(engine, sql, {"cohort": cohort})
    if df is None or df.empty:
        return pd.DataFrame()
    df["date"] = pd.to_datetime(df["date"])
    df["rotationNum"] = df["rotation"].map(rotationNumber)
    return df


def getCohortItemFrame(engine, cohort, formsTable="dds4_boh3_forms_v3", mappingFile=None):
    """ONE ROW PER form x item code (attended patients only; assessmentId kept so a form's
    entrustment / weakness can be attributed to the sections it touched), with
    Section / Sub-section merged from item_section_mapping.xlsx (Utils._mergeSection).
    Same filters as getItemCodesPerStudentBatch, plus the rotation so it can be windowed.
    """
    sql = f"""
    SELECT
      f.assessmentid   AS "assessmentId",
      f.student_number AS "studentNumber",
      f.student_name   AS "studentName",
      f.rotation       AS "rotation",
      ic->>'code'      AS "Item Code",
      SUM(COALESCE(NULLIF(ic->>'quantity','')::int, 1))::int AS "qty"
    FROM {formsTable} f
    CROSS JOIN LATERAL {_PATIENT_ARRAY_SQL} pd
    CROSS JOIN LATERAL jsonb_array_elements(
      CASE WHEN jsonb_typeof(pd->'item_codes')='array' THEN pd->'item_codes' ELSE '[]'::jsonb END) ic
    WHERE f.cohort = :cohort AND f.student_name IS NOT NULL
      AND ic ? 'code' AND (pd->>'patient_attended')::boolean = true
    GROUP BY f.assessmentid, f.student_number, f.student_name, f.rotation, ic->>'code';
    """
    df = readDf(engine, sql, {"cohort": cohort})
    if df is None or df.empty:
        return pd.DataFrame(columns=["assessmentId", "studentNumber", "studentName", "rotation", "rotationNum",
                                     "Item Code", "qty", "Section", "Sub-section"])
    df["Item Code"] = df["Item Code"].astype(str).str.strip()
    df = _mergeSection(df, _loadSectionMapping(mappingFile))
    df["rotationNum"] = df["rotation"].map(rotationNumber)
    return df


# ═══════════════════════════════════════════════════════════════════════════
# 2. Windows
# ═══════════════════════════════════════════════════════════════════════════

def normaliseRotationGroups(groups):
    """[(label, iterable of rotation numbers)] -> [(label, set)]; validates no rotation
    sits in two groups (overlap would double-count)."""
    out, seen = [], {}
    for label, nums in (groups or []):
        numSet = {int(n) for n in nums}
        for n in numSet:
            if n in seen:
                raise ValueError(f"Rotation {n} is in both '{seen[n]}' and '{label}'")
            seen[n] = label
        out.append((str(label), numSet))
    return out


def windowMasks(df, groups, rotationCol="rotationNum"):
    """{windowLabel: boolean mask} — Overall first, then each group in order."""
    masks = {OVERALL_WINDOW: pd.Series(True, index=df.index)}
    for label, numSet in normaliseRotationGroups(groups):
        masks[label] = df[rotationCol].isin(numSet)
    return masks


# ═══════════════════════════════════════════════════════════════════════════
# 3. Metrics (long form: studentNumber, studentName, window, domain, metric, value)
# ═══════════════════════════════════════════════════════════════════════════

def _formMetricsOneWindow(sub, roster):
    """All per-form-derived metrics for one window. `sub` = formDf rows in the window.
    Returns wide DataFrame indexed by studentNumber (every roster student present)."""
    idx = roster.index
    g = sub.groupby("studentNumber")
    assessed = sub[sub["submittedByAssessor"]]
    ga = assessed.groupby("studentNumber")
    rated = sub[sub["entrustment"].notna()]
    gr = rated.groupby("studentNumber")

    out = {}
    # ── Activity ──
    out[(DOMAIN_ACTIVITY, "Forms")] = g.size()
    out[(DOMAIN_ACTIVITY, "Forms assessor-submitted")] = ga.size()
    out[(DOMAIN_ACTIVITY, "Forms not assessor-submitted")] = g["submittedByAssessor"].apply(lambda s: int((~s).sum()))
    out[(DOMAIN_ACTIVITY, "Forms not student-submitted")] = g["submittedByStudent"].apply(lambda s: int((~s).sum()))
    out[(DOMAIN_ACTIVITY, "Active days")] = g["date"].nunique()
    # ── Patients ──
    out[(DOMAIN_PATIENTS, "Patients attended")] = g["patientsAttended"].sum()
    out[(DOMAIN_PATIENTS, "Patients FTA")] = g["patientsFta"].sum()
    out[(DOMAIN_PATIENTS, "Clinics")] = g["clinic"].nunique()
    # ── Age mix ──
    for col, label in (("age0to6", "Age 0-6"), ("age7to17", "Age 7-17"), ("age18plus", "Age 18+")):
        out[(DOMAIN_AGE, label)] = g[col].sum()
    # ── Context counts ──
    out[(DOMAIN_CONTEXT, "Supervisors")] = ga["assessor"].nunique()
    out[(DOMAIN_CONTEXT, "Incident forms")] = ga["incidentTags"].apply(lambda s: int((s > 0).sum()))
    out[(DOMAIN_CONTEXT, "Concern forms")] = g["concernFlag"].sum()

    wide = pd.DataFrame(out).reindex(idx).fillna(0)

    # ── Rates (NaN when the denominator is 0 — never a fake zero) ──
    nAssessed = wide[(DOMAIN_ACTIVITY, "Forms assessor-submitted")].replace(0, np.nan)
    pts = (wide[(DOMAIN_PATIENTS, "Patients attended")] + wide[(DOMAIN_PATIENTS, "Patients FTA")]).replace(0, np.nan)
    wide[(DOMAIN_PATIENTS, "FTA rate")] = wide[(DOMAIN_PATIENTS, "Patients FTA")] / pts

    wide[(DOMAIN_ENTRUSTMENT, "Entrustment n")] = gr.size().reindex(idx).fillna(0)
    wide[(DOMAIN_ENTRUSTMENT, "Entrustment avg")] = gr["entrustment"].mean().reindex(idx)
    wide[(DOMAIN_ENTRUSTMENT, "Low share (L1-2)")] = gr["entrustment"].apply(lambda s: float((s <= 2).mean())).reindex(idx)
    both = sub[sub["entrustment"].notna() & sub["readiness"].notna()]
    wide[(DOMAIN_ENTRUSTMENT, "Self minus supervisor")] = (
        (both["readiness"] - both["entrustment"]).groupby(both["studentNumber"]).mean().reindex(idx))

    # weakness RATE = assessor-submitted forms carrying >=1 tag in the category / assessor-submitted forms
    for label in WEAKNESS_LABELS:
        tagged = (assessed[label] > 0).groupby(assessed["studentNumber"]).sum().reindex(idx).fillna(0)
        wide[(DOMAIN_WEAKNESS, f"{label} rate")] = tagged / nAssessed
        wide[(DOMAIN_WEAKNESS, f"{label} forms")] = tagged
    anyTag = (assessed[WEAKNESS_LABELS].sum(axis=1) > 0).groupby(assessed["studentNumber"]).sum().reindex(idx).fillna(0)
    wide[(DOMAIN_WEAKNESS, "Any weakness rate")] = anyTag / nAssessed
    tagSum = assessed[WEAKNESS_LABELS].sum(axis=1).groupby(assessed["studentNumber"]).sum().reindex(idx).fillna(0)
    wide[(DOMAIN_WEAKNESS, "Weakness tags per form")] = tagSum / nAssessed
    comm = ga["commendationTags"].sum().reindex(idx).fillna(0)
    wide[(DOMAIN_CONTEXT, "Commendations per form")] = comm / nAssessed
    return wide


def _itemMetricsOneWindow(sub, roster):
    """Coverage metrics for one window: totals, distinct codes, per-Section and
    per-Sub-section quantities. Indexed by studentNumber."""
    idx = roster.index
    out = {}
    g = sub.groupby("studentNumber")
    out[(DOMAIN_COVERAGE, "Items total")] = g["qty"].sum()
    out[(DOMAIN_COVERAGE, "Distinct item codes")] = g["Item Code"].nunique()
    wide = pd.DataFrame(out).reindex(idx).fillna(0)
    if sub.empty:
        return wide
    sec = sub.pivot_table(index="studentNumber", columns="Section", values="qty", aggfunc="sum", fill_value=0)
    sec.columns = pd.MultiIndex.from_tuples([(DOMAIN_COVERAGE, f"{SECTION_PREFIX}{c}") for c in sec.columns])
    subKey = sub["Section"].astype(str) + " / " + sub["Sub-section"].astype(str)
    ss = sub.assign(_ss=subKey).pivot_table(index="studentNumber", columns="_ss", values="qty",
                                            aggfunc="sum", fill_value=0)
    ss.columns = pd.MultiIndex.from_tuples([(DOMAIN_COVERAGE, f"{SUBSECTION_PREFIX}{c}") for c in ss.columns])
    return pd.concat([wide, sec.reindex(idx).fillna(0), ss.reindex(idx).fillna(0)], axis=1)


def buildMetricsLong(formDf, itemDf, rotationGroups=None, year=2026, addPbnLeave=True):
    """Every metric x window x student in LONG form.

    rotationGroups: [(label, [rotation numbers]), ...] — any number of groups. Overall is
    always added. Returns columns: studentNumber, studentName, window, domain, metric, value.
    Students with no forms in a window still get a row (counts 0, rates NaN) so a gap shows
    as a gap, not as a missing student.
    """
    groups = DEFAULT_ROTATION_GROUPS if rotationGroups is None else rotationGroups
    roster = (formDf.groupby("studentNumber")["studentName"].first().to_frame())
    fMasks, iMasks = windowMasks(formDf, groups), windowMasks(itemDf, groups)

    frames = []
    for window in fMasks:
        wide = pd.concat([_formMetricsOneWindow(formDf[fMasks[window]], roster),
                          _itemMetricsOneWindow(itemDf[iMasks[window]], roster)], axis=1)
        if window == OVERALL_WINDOW:
            lastDate = formDf.groupby("studentNumber")["date"].max()
            wide[(DOMAIN_ACTIVITY, "Days since last form")] = (formDf["date"].max() - lastDate).dt.days.reindex(roster.index)
            wide[(DOMAIN_ACTIVITY, "Forms without rotation")] = (
                formDf[formDf["rotationNum"].isna()].groupby("studentNumber").size().reindex(roster.index).fillna(0))
            if addPbnLeave:
                pbnMap, leaveMap = _pbnLeaveMaps(roster, year)
                wide[(DOMAIN_CONTEXT, f"PBNs {year}")] = roster.index.map(lambda s: pbnMap.get(int(s), 0))
                wide[(DOMAIN_CONTEXT, f"Leave days {year}")] = roster.index.map(lambda s: leaveMap.get(int(s), 0))
        # wide -> long without DataFrame.stack (its behaviour differs across pandas versions)
        wide.index.name = "studentNumber"
        longDf = (pd.concat({col: wide[col] for col in wide.columns}, names=["domain", "metric"])
                    .rename("value").reset_index())
        longDf.insert(1, "window", window)
        frames.append(longDf)

    out = pd.concat(frames, ignore_index=True)
    out.insert(1, "studentName", out["studentNumber"].map(roster["studentName"]))
    return out[["studentNumber", "studentName", "window", "domain", "metric", "value"]]


def _pbnLeaveMaps(roster, year):
    """PBN + leave lookups via the existing readers; a missing file degrades to zeros."""
    rosterDf = roster.reset_index().rename(columns={"studentNumber": "Student ID", "studentName": "Student Name"})
    try:
        pbnMap = getPbnCountByStudent(rosterDf, year)
    except Exception as ex:
        print(f"[flagging] PBN read failed ({ex}); filling 0")
        pbnMap = {}
    try:
        leaveMap = getLeaveDaysByStudent(rosterDf["Student ID"].dropna().astype(int), year)
    except Exception as ex:
        print(f"[flagging] Leave read failed ({ex}); filling 0")
        leaveMap = {}
    return pbnMap, leaveMap


def pivotMetrics(metricsLong, domains=None, windows=None, metrics=None):
    """LONG -> WIDE, one row per student, columns '<metric> | <window>' (metric-major so
    Overall / R1-4 / R5+ of the same metric sit side by side). Optional filters."""
    df = metricsLong
    if domains is not None:
        df = df[df["domain"].isin(list(domains))]
    if windows is not None:
        df = df[df["window"].isin(list(windows))]
    if metrics is not None:
        df = df[df["metric"].isin(list(metrics))]
    metricOrder = list(dict.fromkeys(df["metric"]))
    windowOrder = list(dict.fromkeys(df["window"]))
    # index on studentNumber only: dropna=False with a 2-level index would build the
    # number x name cartesian product
    nameMap = df.drop_duplicates("studentNumber").set_index("studentNumber")["studentName"]
    wide = df.pivot_table(index="studentNumber", columns=["metric", "window"],
                          values="value", aggfunc="first", dropna=False)
    cols = [(m, w) for m in metricOrder for w in windowOrder if (m, w) in wide.columns]
    wide = wide[cols].dropna(axis=1, how="all")   # Overall-only metrics have no group columns
    wide.columns = [f"{m} | {w}" for m, w in wide.columns]
    wide = wide.reset_index()
    wide.insert(1, "studentName", wide["studentNumber"].map(nameMap))
    return wide.sort_values("studentName", key=lambda s: s.astype(str).str.casefold()).reset_index(drop=True)


# ═══════════════════════════════════════════════════════════════════════════
# 4. Cohort reference: eligible students, expected sections, threshold preview
# ═══════════════════════════════════════════════════════════════════════════

def getEligibleStudents(metricsLong, minFormsFloor=DEFAULT_MIN_FORMS_FLOOR):
    """Student numbers with >= minFormsFloor forms Overall — the population the cohort
    mean / SD is computed over. Everyone else is reported as 'insufficient activity'."""
    forms = metricsLong[(metricsLong["window"] == OVERALL_WINDOW) & (metricsLong["metric"] == "Forms")]
    return set(forms.loc[forms["value"] >= minFormsFloor, "studentNumber"])


def getExpectedSections(metricsLong, minShare=DEFAULT_EXPECTED_MIN_SHARE,
                        minFormsFloor=DEFAULT_MIN_FORMS_FLOOR):
    """Derive what the cohort is 'supposed' to do from what it actually does.
    One row per Section / Sub-section x window: share of eligible students with any
    activity, cohort median / mean, and `expected` (share >= minShare, not Unmapped/Misc)."""
    eligible = getEligibleStudents(metricsLong, minFormsFloor)
    cov = metricsLong[(metricsLong["domain"] == DOMAIN_COVERAGE)
                      & metricsLong["metric"].str.startswith((SECTION_PREFIX, SUBSECTION_PREFIX))
                      & metricsLong["studentNumber"].isin(eligible)]
    if cov.empty:
        return pd.DataFrame(columns=["metric", "window", "level", "shareWithActivity", "median", "mean", "expected"])
    g = cov.groupby(["metric", "window"], sort=False)["value"]
    out = pd.DataFrame({"shareWithActivity": g.apply(lambda s: float((s.fillna(0) > 0).mean())),
                        "median": g.median(), "mean": g.mean()}).reset_index()
    out["level"] = np.where(out["metric"].str.startswith(SECTION_PREFIX), "Section", "Sub-section")
    nonFlag = out["metric"].apply(lambda m: any(n in m for n in NON_FLAG_SECTIONS))
    out["expected"] = (out["shareWithActivity"] >= minShare) & ~nonFlag
    return out[["metric", "window", "level", "shareWithActivity", "median", "mean", "expected"]].round(3)


THRESHOLD_MODES = ("fixed", "meanStd", "meanFrac")


def resolveThreshold(values, spec, window=OVERALL_WINDOW):
    """Turn a threshold spec into a number for one metric+window.
        {"mode": "fixed",    "value": 60}  or  {"mode": "fixed", "value": {"Overall": 60, "R1-4": 25}}
        {"mode": "meanStd",  "k": 1.5}     -> mean - k*SD   (floored at 0)
        {"mode": "meanFrac", "f": 0.5}     -> f * mean
    `values` = the eligible students' values. Returns NaN when it cannot be resolved
    (e.g. fixed dict with no entry for this window)."""
    mode = spec.get("mode", "meanStd")
    if mode not in THRESHOLD_MODES:
        raise ValueError(f"Unknown threshold mode '{mode}'. Use one of {THRESHOLD_MODES}")
    vals = pd.Series(values, dtype=float).dropna()
    if mode == "fixed":
        v = spec.get("value")
        v = v.get(window) if isinstance(v, dict) else v
        return float(v) if v is not None else np.nan
    if vals.empty:
        return np.nan
    above = spec.get("direction", "below") == "above"      # "above": flag HIGH values (mean + k*SD)
    if mode == "meanStd":
        if above:
            return float(vals.mean() + spec.get("k", 1.5) * vals.std(ddof=1))
        return max(0.0, float(vals.mean() - spec.get("k", 1.5) * vals.std(ddof=1)))
    return float(spec.get("f", 0.5) * vals.mean())


def cohortWindowStats(metricsLong, domains=(DOMAIN_ACTIVITY, DOMAIN_PATIENTS, DOMAIN_AGE, DOMAIN_COVERAGE),
                      k=1.5, f=0.5, minFormsFloor=DEFAULT_MIN_FORMS_FLOOR):
    """TUNING AID. Per metric x window over eligible students: n, mean, SD, median, min, and —
    for BOTH relative modes at the given k / f — the resolved threshold and how many students
    would fall below it. 'meanStd usable' is False when mean - k*SD <= 0 (the rule can never
    fire; pick meanFrac or fixed for that metric)."""
    eligible = getEligibleStudents(metricsLong, minFormsFloor)
    df = metricsLong[metricsLong["domain"].isin(list(domains)) & metricsLong["studentNumber"].isin(eligible)]
    rows = []
    for (domain, metric, window), s in df.groupby(["domain", "metric", "window"], sort=False)["value"]:
        s = s.astype(float)
        tStd = resolveThreshold(s, {"mode": "meanStd", "k": k})
        tFrac = resolveThreshold(s, {"mode": "meanFrac", "f": f})
        rows.append({"domain": domain, "metric": metric, "window": window, "n": int(s.notna().sum()),
                     "mean": s.mean(), "sd": s.std(ddof=1), "median": s.median(), "min": s.min(),
                     f"meanStd(k={k})": tStd, "below meanStd": int((s < tStd).sum()),
                     "meanStd usable": bool(tStd > 0),
                     f"meanFrac(f={f})": tFrac, "below meanFrac": int((s < tFrac).sum())})
    return pd.DataFrame(rows).round(2)


# ═══════════════════════════════════════════════════════════════════════════
# 5. Section-level entrustment & weakness (is the student strong THROUGHOUT?)
# ═══════════════════════════════════════════════════════════════════════════
# Attribution is CO-OCCURRENCE: entrustment / weakness are recorded once per form, so a
# form counts toward EVERY section its item codes belong to (same convention as the
# "weakness by item code" views). Sections on almost every form (e.g. Diagnostics)
# therefore track the student's overall figure.
# Flags are COHORT-RELATIVE PER SECTION + WINDOW, because sections differ in difficulty
# (everyone scores lower in some) and the whole cohort improves through the year.

DEFAULT_SECTION_MIN_FORMS = 5       # student needs this many rated forms in a cell to be judged
DEFAULT_SECTION_MIN_STUDENTS = 8    # cell's cohort mean/SD needs this many judged students
DEFAULT_SECTION_K = 1.0             # LOW  : z <= -k   (weakness: z >= +k)
DEFAULT_SECTION_WATCH_K = 0.5       # WATCH: z <= -watchK

STATUS_LOW, STATUS_WATCH, STATUS_OK, STATUS_NA = "low", "watch", "ok", "n/a"
SECTION_STATUS_FILLS = {STATUS_LOW: "FFC7CE", STATUS_WATCH: "FFEB9C", STATUS_OK: "C6EFCE", STATUS_NA: "EDEDED"}


def getFormSectionFrame(formDf, itemDf, level="Section"):
    """One row per form x section (or 'Section / Sub-section' when level='Sub-section'),
    carrying that form's entrustment + weakness columns. De-duplicated so a form with
    three Restorative codes counts ONCE for Restorative."""
    if level == "Sub-section":
        key = itemDf["Section"].astype(str) + " / " + itemDf["Sub-section"].astype(str)
    else:
        key = itemDf["Section"].astype(str)
    pairs = pd.DataFrame({"assessmentId": itemDf["assessmentId"], "section": key}).drop_duplicates()
    cols = ["assessmentId", "studentNumber", "studentName", "rotationNum", "submittedByAssessor",
            "entrustment"] + WEAKNESS_LABELS
    return pairs.merge(formDf[cols], on="assessmentId", how="inner")


def buildSectionPerformanceLong(formDf, itemDf, trajectoryGroups=None, level="Section"):
    """Per student x section x window: forms, entrustmentN, entrustmentAvg, lowShare (L1-2),
    weaknessRate (assessor forms with >=1 weakness tag / assessor forms), tagsPerForm,
    topWeakness (most-tagged category). Windows = Overall + trajectoryGroups."""
    groups = DEFAULT_TRAJECTORY_GROUPS if trajectoryGroups is None else trajectoryGroups
    fs = getFormSectionFrame(formDf, itemDf, level)
    frames = []
    for window, mask in windowMasks(fs, groups).items():
        sub = fs[mask]
        if sub.empty:
            continue
        keys = ["studentNumber", "section"]
        rated = sub[sub["entrustment"].notna()]
        assessed = sub[sub["submittedByAssessor"]].copy()
        assessed["_tags"] = assessed[WEAKNESS_LABELS].sum(axis=1)
        out = pd.DataFrame({"forms": sub.groupby(keys).size()})
        out["entrustmentN"] = rated.groupby(keys).size()
        out["entrustmentAvg"] = rated.groupby(keys)["entrustment"].mean()
        out["lowShare"] = rated.groupby(keys)["entrustment"].apply(lambda s: float((s <= 2).mean()))
        ga = assessed.groupby(keys)
        out["assessedForms"] = ga.size()
        out["weaknessRate"] = ga["_tags"].apply(lambda s: float((s > 0).mean()))
        out["tagsPerForm"] = ga["_tags"].mean()
        catSums = ga[WEAKNESS_LABELS].sum()
        out["topWeakness"] = catSums.idxmax(axis=1).where(catSums.sum(axis=1) > 0)
        out = out.reset_index()
        out.insert(2, "window", window)
        frames.append(out)
    if not frames:
        return pd.DataFrame()
    res = pd.concat(frames, ignore_index=True)
    nameMap = formDf.drop_duplicates("studentNumber").set_index("studentNumber")["studentName"]
    res.insert(1, "studentName", res["studentNumber"].map(nameMap))
    res[["entrustmentN", "assessedForms"]] = res[["entrustmentN", "assessedForms"]].fillna(0).astype(int)
    return res


def flagSectionPerformance(sectionLong, eligibleStudents=None, k=DEFAULT_SECTION_K,
                           watchK=DEFAULT_SECTION_WATCH_K, minForms=DEFAULT_SECTION_MIN_FORMS,
                           minStudents=DEFAULT_SECTION_MIN_STUDENTS):
    """Add cohort-relative z-scores + a status per cell.
        entZ  = (student entrustmentAvg - cohort mean) / cohort SD   within section+window
        weakZ = same on weaknessRate
        entStatus : low (entZ <= -k) | watch (<= -watchK) | ok | n/a
        weakStatus: low (weakZ >= +k) | watch (>= +watchK) | ok | n/a     ('low' = bad, both)
    n/a when the student has < minForms in the cell, the cell has < minStudents judged
    students, the cohort SD is 0, or the section is Unmapped / Miscellaneous.
    Cohort mean/SD use only eligibleStudents (forms floor) with >= minForms in the cell."""
    df = sectionLong.copy()
    nonFlag = df["section"].apply(lambda s: any(n in str(s) for n in NON_FLAG_SECTIONS))
    inPool = pd.Series(True, index=df.index) if eligibleStudents is None else df["studentNumber"].isin(eligibleStudents)

    def _score(valueCol, nCol, sign):
        judged = (df[nCol] >= minForms) & df[valueCol].notna() & ~nonFlag
        pool = df[judged & inPool].groupby(["section", "window"])[valueCol]
        stats = pd.DataFrame({"mean": pool.mean(), "sd": pool.std(ddof=1), "n": pool.size()})
        merged = df[["section", "window"]].merge(stats, left_on=["section", "window"], right_index=True, how="left")
        merged.index = df.index
        usable = judged & (merged["n"] >= minStudents) & (merged["sd"] > 0)
        z = ((df[valueCol] - merged["mean"]) / merged["sd"]).where(usable)
        bad = sign * z                                # >0 means worse than cohort
        status = pd.Series(STATUS_NA, index=df.index)
        status[usable] = STATUS_OK
        status[usable & (bad >= watchK)] = STATUS_WATCH
        status[usable & (bad >= k)] = STATUS_LOW
        return merged["mean"], z, status

    df["cohortEntrustmentAvg"], df["entZ"], df["entStatus"] = _score("entrustmentAvg", "entrustmentN", -1)
    df["cohortWeaknessRate"], df["weakZ"], df["weakStatus"] = _score("weaknessRate", "assessedForms", +1)
    return df


def summariseSectionFlags(sectionFlagged, prefix="Sections"):
    """One row per student — the columns that go into the Summary:
        '<prefix> low entrustment'      sections LOW on the Overall window
        '<prefix> persistently low ent' LOW in every judged group window (>= 2 judged)
        '<prefix> high weakness' / '<prefix> persistently high weakness'   same for weakness
        '<prefix> judged'               sections with an Overall verdict
        'Strong throughout (<prefix>)'  True when NO cell (any window) is LOW on either measure
                                        and at least one section was judged."""
    rows = []
    for sn, g in sectionFlagged.groupby("studentNumber"):
        overall, grp = g[g["window"] == OVERALL_WINDOW], g[g["window"] != OVERALL_WINDOW]
        row = {"studentNumber": sn}
        for col, label in (("entStatus", "low entrustment"), ("weakStatus", "high weakness")):
            lowOverall = sorted(overall.loc[overall[col] == STATUS_LOW, "section"])
            judged = grp[grp[col] != STATUS_NA]
            agg = judged.groupby("section")[col].agg(n="size", low=lambda s: int((s == STATUS_LOW).sum()))
            persistent = sorted(agg.index[(agg["n"] >= 2) & (agg["low"] == agg["n"])]) if len(agg) else []
            row[f"{prefix} {label}"] = "; ".join(lowOverall)
            row[f"{prefix} {label} (n)"] = len(lowOverall)
            row[f"{prefix} persistently {label}"] = "; ".join(persistent)
        nJudged = int((overall["entStatus"] != STATUS_NA).sum())
        anyLow = bool(((g["entStatus"] == STATUS_LOW) | (g["weakStatus"] == STATUS_LOW)).any())
        row[f"{prefix} judged"] = nJudged
        row[f"Strong throughout ({prefix})"] = (nJudged > 0) and not anyLow
        rows.append(row)
    return pd.DataFrame(rows)


def pivotSectionPerformance(sectionFlagged, valueCol="entrustmentAvg", statusCol="entStatus"):
    """(valuesWide, statusWide): one row per student, columns '<section> | <window>'
    (Overall first, then each group), sections ordered A-Z. statusWide has the same shape
    and drives the cell colours."""
    df = sectionFlagged
    windowOrder = list(dict.fromkeys(df["window"]))
    sections = sorted(df["section"].unique(), key=lambda s: (any(n in s for n in NON_FLAG_SECTIONS), s))
    nameMap = df.drop_duplicates("studentNumber").set_index("studentNumber")["studentName"]

    def _wide(col):
        w = df.pivot(index="studentNumber", columns=["section", "window"], values=col)
        cols = [(s, x) for s in sections for x in windowOrder if (s, x) in w.columns]
        w = w[cols]
        w.columns = [f"{s} | {x}" for s, x in cols]
        w = w.reset_index()
        w.insert(1, "studentName", w["studentNumber"].map(nameMap))
        return w.sort_values("studentName", key=lambda s: s.astype(str).str.casefold()).reset_index(drop=True)

    return _wide(valueCol), _wide(statusCol).fillna(STATUS_NA)


def writeSectionPivotSheet(writer, valuesWide, statusWide, sheetName, title, numberFormat="0.00"):
    """Colour-coded pivot: red = LOW vs cohort for that section+window, amber = WATCH,
    green = OK, grey = not enough forms to judge. Row 1 = title/legend, row 2 = headers."""
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    valuesWide.to_excel(writer, sheet_name=sheetName, index=False, startrow=1)
    ws = writer.sheets[sheetName]
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12, color="010D44")
    for i, (status, text) in enumerate(((STATUS_LOW, "Low vs cohort"), (STATUS_WATCH, "Watch"),
                                        (STATUS_OK, "OK"), (STATUS_NA, "Too few forms"))):
        c = ws.cell(row=1, column=4 + i, value=text)
        c.fill = PatternFill("solid", fgColor=SECTION_STATUS_FILLS[status])
        c.alignment = Alignment(horizontal="center")
    fills = {s: PatternFill("solid", fgColor=hexCol) for s, hexCol in SECTION_STATUS_FILLS.items()}
    headerFill, headerFont = PatternFill("solid", fgColor="010D44"), Font(bold=True, color="FFFFFF", size=9)
    for j, colName in enumerate(valuesWide.columns, start=1):
        h = ws.cell(row=2, column=j)
        h.value = str(colName).replace(" | ", "\n")
        h.fill, h.font = headerFill, headerFont
        h.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = 12 if j == 1 else (26 if j == 2 else 11)
        if j <= 2:
            continue
        statusCol = statusWide[colName]
        for i in range(len(valuesWide)):
            cell = ws.cell(row=3 + i, column=j)
            cell.fill = fills.get(statusCol.iloc[i], fills[STATUS_NA])
            cell.number_format = numberFormat
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "C3"


def buildSectionPerformance(formDf, itemDf, eligibleStudents=None, trajectoryGroups=None, level="Section",
                            k=DEFAULT_SECTION_K, watchK=DEFAULT_SECTION_WATCH_K,
                            minForms=DEFAULT_SECTION_MIN_FORMS, minStudents=DEFAULT_SECTION_MIN_STUDENTS):
    """Long + flagged + per-student summary for one level. Returns (flaggedLong, summaryDf)."""
    longDf = buildSectionPerformanceLong(formDf, itemDf, trajectoryGroups, level)
    if longDf.empty:
        return longDf, pd.DataFrame(columns=["studentNumber"])
    flagged = flagSectionPerformance(longDf, eligibleStudents, k, watchK, minForms, minStudents)
    prefix = "Sections" if level == "Section" else "Sub-sections"
    return flagged, summariseSectionFlags(flagged, prefix)


# ═══════════════════════════════════════════════════════════════════════════
# 6. Convenience driver (stage 1)
# ═══════════════════════════════════════════════════════════════════════════

def buildFlaggingMetrics(engine, cohort, formsTable="dds4_boh3_forms_v3", rotationGroups=None,
                         trajectoryGroups=None, excludeStudentNumbers=None,
                         minFormsFloor=DEFAULT_MIN_FORMS_FLOOR, expectedMinShare=DEFAULT_EXPECTED_MIN_SHARE,
                         k=1.5, f=0.5, year=2026, outPath=None,
                         sectionK=DEFAULT_SECTION_K, sectionWatchK=DEFAULT_SECTION_WATCH_K,
                         sectionMinForms=DEFAULT_SECTION_MIN_FORMS,
                         sectionMinStudents=DEFAULT_SECTION_MIN_STUDENTS):
    """Run the whole of stage 1 for one cohort. Returns a dict of DataFrames:
        formDf, itemDf            raw frames
        metricsLong               counts windows (rotationGroups)
        trajectoryLong            entrustment + weakness on trajectoryGroups (for stage 2)
        metricsWide               one row per student, non-coverage metrics
        coverageWide              one row per student, Section / Sub-section quantities
        expectedSections          derived cohort structure
        windowStats               threshold tuning table
        insufficient              students under the forms floor
    outPath: optional .xlsx — dumps the review tables (plain, unstyled) for eyeballing."""
    formDf = getCohortFormFrame(engine, cohort, formsTable)
    if formDf.empty:
        raise LookupError(f"No forms found for cohort '{cohort}' in {formsTable}")
    itemDf = getCohortItemFrame(engine, cohort, formsTable)
    if excludeStudentNumbers:
        drop = {int(s) for s in excludeStudentNumbers}
        formDf = formDf[~formDf["studentNumber"].isin(drop)]
        itemDf = itemDf[~itemDf["studentNumber"].isin(drop)]

    metricsLong = buildMetricsLong(formDf, itemDf, rotationGroups, year=year)
    trajGroups = DEFAULT_TRAJECTORY_GROUPS if trajectoryGroups is None else trajectoryGroups
    trajectoryLong = buildMetricsLong(formDf, itemDf, trajGroups, year=year, addPbnLeave=False)
    trajectoryLong = trajectoryLong[trajectoryLong["domain"].isin([DOMAIN_ENTRUSTMENT, DOMAIN_WEAKNESS])
                                    | (trajectoryLong["metric"] == "Forms assessor-submitted")]

    isCoverageSplit = metricsLong["metric"].str.startswith((SECTION_PREFIX, SUBSECTION_PREFIX))
    eligible = getEligibleStudents(metricsLong, minFormsFloor)
    overallForms = metricsLong[(metricsLong["window"] == OVERALL_WINDOW) & (metricsLong["metric"] == "Forms")]
    result = {
        "formDf": formDf, "itemDf": itemDf,
        "metricsLong": metricsLong, "trajectoryLong": trajectoryLong,
        "metricsWide": pivotMetrics(metricsLong[~isCoverageSplit]),
        "coverageWide": pivotMetrics(metricsLong[isCoverageSplit]),
        "trajectoryWide": pivotMetrics(trajectoryLong),
        "expectedSections": getExpectedSections(metricsLong, expectedMinShare, minFormsFloor),
        "windowStats": cohortWindowStats(metricsLong, k=k, f=f, minFormsFloor=minFormsFloor),
        "insufficient": overallForms.loc[~overallForms["studentNumber"].isin(eligible),
                                         ["studentNumber", "studentName", "value"]].rename(columns={"value": "Forms"}),
    }
    # section-level entrustment / weakness (cohort-relative, colour-coded pivots + summary columns)
    secArgs = dict(eligibleStudents=eligible, trajectoryGroups=trajGroups, k=sectionK, watchK=sectionWatchK,
                   minForms=sectionMinForms, minStudents=sectionMinStudents)
    result["sectionPerf"], secSummary = buildSectionPerformance(formDf, itemDf, level="Section", **secArgs)
    result["subSectionPerf"], subSummary = buildSectionPerformance(formDf, itemDf, level="Sub-section", **secArgs)
    result["sectionSummary"] = secSummary.merge(subSummary, on="studentNumber", how="outer")
    result["metricsWide"] = result["metricsWide"].merge(result["sectionSummary"], on="studentNumber", how="left")

    print(f"[{cohort}] {formDf['studentNumber'].nunique()} students, {len(formDf)} forms, "
          f"{int(formDf['rotationNum'].isna().sum())} forms without a rotation, "
          f"{len(result['insufficient'])} under the {minFormsFloor}-form floor")
    if outPath:
        with pd.ExcelWriter(outPath, engine="openpyxl") as writer:
            for sheet in ("metricsWide", "trajectoryWide", "coverageWide", "expectedSections",
                          "windowStats", "insufficient"):
                result[sheet].to_excel(writer, sheet_name=sheet, index=False)
            for key, tag in (("sectionPerf", "Sec"), ("subSectionPerf", "SubSec")):
                perf = result[key]
                if perf.empty:
                    continue
                for valueCol, statusCol, name, title, fmt in (
                        ("entrustmentAvg", "entStatus", f"{tag} Entrustment",
                         "Mean entrustment by section x window — colour = vs cohort in the SAME section+window", "0.00"),
                        ("weaknessRate", "weakStatus", f"{tag} Weakness",
                         "Share of forms with >=1 weakness tag by section x window — colour = vs cohort", "0%")):
                    vals, stat = pivotSectionPerformance(perf, valueCol, statusCol)
                    writeSectionPivotSheet(writer, vals, stat, name, title, fmt)
                perf.to_excel(writer, sheet_name=f"{tag} detail", index=False)
        print(f"[{cohort}] review workbook -> {outPath}")
    return result


# ═══════════════════════════════════════════════════════════════════════════
# 7. STAGE 2 — flags
# ═══════════════════════════════════════════════════════════════════════════
# 7a. COUNT metrics (forms / patients / age bands / items / sections): per-metric threshold
#     spec, three modes (fixed | meanStd | meanFrac), evaluated on Overall + every rotation group.
# 7b. TRAJECTORY (entrustment + the 7 weakness categories): cohort-relative z per window,
#     then a state per student (consistently low / falling behind / ... ; persistent / emerging).

SECTION_WILDCARD = SECTION_PREFIX + "*"          # one spec for every expected Section
SUBSECTION_WILDCARD = SUBSECTION_PREFIX + "*"    # one spec for every expected Sub-section

# Default per-metric specs. Override any of them (or add "Section: Endodontics" etc.) via
# thresholds={...}. A spec may carry "direction": "above" (default "below") and, for fixed,
# "value" as a number or a {window: number} dict. Set a metric to None to switch it off.
DEFAULT_COUNT_THRESHOLDS = {
    "Forms":                    {"mode": "meanStd", "k": 1.5},
    "Forms assessor-submitted": {"mode": "meanStd", "k": 1.5},
    "Patients attended":        {"mode": "meanStd", "k": 1.5},
    "Age 0-6":                  {"mode": "meanStd", "k": 1.5},
    "Age 7-17":                 {"mode": "meanStd", "k": 1.5},
    "Age 18+":                  {"mode": "meanStd", "k": 1.5},
    "Items total":              {"mode": "meanStd", "k": 1.5},
    "Distinct item codes":      {"mode": "meanStd", "k": 1.5},
    "Days since last form":     {"mode": "meanStd", "k": 1.5, "direction": "above"},
    SECTION_WILDCARD:           {"mode": "meanStd", "k": 1.5},
    SUBSECTION_WILDCARD:        {"mode": "meanStd", "k": 1.5},
}

DEFAULT_TRAJECTORY_K = 1.0            # window is LOW when z <= -k (weakness: z >= +k)
DEFAULT_TRAJECTORY_MIN_FORMS = 5      # rated forms needed in a window to judge it
DEFAULT_WEAKNESS_MIN_TAGGED = 3       # tagged forms needed before a category can be HIGH
DEFAULT_FALLING_DROP = 1.0            # z drop first->last window that counts as "falling behind"

ENT_CONSISTENTLY_LOW = "Consistently low"
ENT_FALLING_BEHIND = "Falling behind"
ENT_LOW_RECENT = "Low recently"
ENT_CATCHING_UP = "Catching up"
ENT_ON_TRACK = "On track"
ENT_INSUFFICIENT = "Insufficient data"
ENT_FLAG_STATES = (ENT_CONSISTENTLY_LOW, ENT_FALLING_BEHIND)     # red
ENT_WATCH_STATES = (ENT_LOW_RECENT,)                             # amber

WEAK_PERSISTENT, WEAK_EMERGING, WEAK_RESOLVING, WEAK_CLEAR = "Persistent", "Emerging", "Resolving", "Clear"


def specFor(metric, thresholds):
    """Exact metric spec, else the Section / Sub-section wildcard, else None (not flagged)."""
    if metric in thresholds:
        return thresholds[metric]
    if metric.startswith(SUBSECTION_PREFIX):
        return thresholds.get(SUBSECTION_WILDCARD)
    if metric.startswith(SECTION_PREFIX):
        return thresholds.get(SECTION_WILDCARD)
    return None


def evaluateCountFlags(metricsLong, expectedSections, thresholds=None, minFormsFloor=DEFAULT_MIN_FORMS_FLOOR):
    """Apply the threshold specs. Returns LONG rows for every flaggable metric x window x student:
        value, threshold, mode, direction, usable, status ('low' | 'ok' | 'n/a').
    - cohort mean/SD come from eligible students only (forms floor);
    - Section / Sub-section metrics are judged only where `expected` for that window;
    - in a rotation-group window a student with 0 forms there is 'n/a' (not reached yet);
    - students under the forms floor are 'n/a' everywhere (reported as insufficient activity);
    - usable=False when a relative 'below' threshold resolves to <= 0 (rule can never fire);
    - comparison: fixed is inclusive (<= / >=), relative modes are strict (< / >)."""
    specs = dict(DEFAULT_COUNT_THRESHOLDS)
    specs.update(thresholds or {})
    eligible = getEligibleStudents(metricsLong, minFormsFloor)
    expectedKeys = set(map(tuple, expectedSections.loc[expectedSections["expected"], ["metric", "window"]].values))
    formsByWindow = (metricsLong[metricsLong["metric"] == "Forms"]
                     .set_index(["studentNumber", "window"])["value"])

    out = []
    for (metric, window), g in metricsLong.groupby(["metric", "window"], sort=False):
        spec = specFor(metric, specs)
        if not spec:
            continue
        isSection = metric.startswith((SECTION_PREFIX, SUBSECTION_PREFIX))
        if isSection and (metric, window) not in expectedKeys:
            continue
        if g["value"].notna().sum() == 0:
            continue
        direction = spec.get("direction", "below")
        pool = g.loc[g["studentNumber"].isin(eligible), "value"]
        thr = resolveThreshold(pool, spec, window)
        usable = bool(pd.notna(thr) and (thr > 0 or direction == "above" or spec.get("mode") == "fixed"))
        vals = g["value"].astype(float)
        if direction == "above":
            hit = vals >= thr if spec.get("mode") == "fixed" else vals > thr
        else:
            hit = vals <= thr if spec.get("mode") == "fixed" else vals < thr
        reached = g.apply(lambda r: formsByWindow.get((r["studentNumber"], window), 0) > 0, axis=1)
        judged = usable & g["studentNumber"].isin(eligible) & reached & vals.notna()
        status = np.where(~judged, STATUS_NA, np.where(hit, STATUS_LOW, STATUS_OK))
        out.append(pd.DataFrame({
            "studentNumber": g["studentNumber"].values, "studentName": g["studentName"].values,
            "domain": g["domain"].values, "metric": metric, "window": window, "value": vals.values,
            "threshold": thr, "mode": spec.get("mode", "meanStd"), "direction": direction,
            "param": spec.get("k", 1.5) if spec.get("mode", "meanStd") == "meanStd"
                     else spec.get("f", 0.5) if spec.get("mode") == "meanFrac" else np.nan,
            "usable": usable, "status": status}))
    return pd.concat(out, ignore_index=True) if out else pd.DataFrame()


def _trajectoryWindows(trajectoryLong):
    return [w for w in dict.fromkeys(trajectoryLong["window"]) if w != OVERALL_WINDOW]


def _zByWindow(valueWide, nWide, eligible, minN):
    """z-score per window (columns) of each student (rows) against eligible, judged peers."""
    judged = (nWide >= minN) & valueWide.notna()
    pool = valueWide.where(judged & valueWide.index.to_series().isin(eligible).values[:, None])
    mean, sd = pool.mean(), pool.std(ddof=1)
    z = ((valueWide - mean) / sd.where(sd > 0)).where(judged)
    return z, mean


def evaluateEntrustmentTrajectory(trajectoryLong, eligible, k=DEFAULT_TRAJECTORY_K,
                                  minForms=DEFAULT_TRAJECTORY_MIN_FORMS, fallingDrop=DEFAULT_FALLING_DROP):
    """One row per student: entrustment avg + z per trajectory window, overall avg / low share,
    and a STATE —
        Consistently low : every judged window z <= -k (>= 2 judged)            -> FLAG
        Falling behind   : last window z <= -k AND z fell >= fallingDrop         -> FLAG
        Low recently     : last window z <= -k only                              -> watch
        Catching up      : first window low, last window not
        On track / Insufficient data (fewer than 2 judged windows)."""
    windows = _trajectoryWindows(trajectoryLong)
    t = trajectoryLong
    def _wide(metric):
        return (t[t["metric"] == metric].pivot(index="studentNumber", columns="window", values="value")
                .reindex(columns=windows))
    avg, n = _wide("Entrustment avg"), _wide("Entrustment n").fillna(0)
    z, cohortMean = _zByWindow(avg, n, eligible, minForms)
    overall = t[t["window"] == OVERALL_WINDOW].pivot(index="studentNumber", columns="metric", values="value")

    rows = []
    for sn in avg.index:
        zs = z.loc[sn].dropna()
        if len(zs) < 2:
            state = ENT_INSUFFICIENT
        else:
            low = zs <= -k
            if low.all():
                state = ENT_CONSISTENTLY_LOW
            elif low.iloc[-1] and (zs.iloc[0] - zs.iloc[-1]) >= fallingDrop:
                state = ENT_FALLING_BEHIND
            elif low.iloc[-1]:
                state = ENT_LOW_RECENT
            elif low.iloc[0]:
                state = ENT_CATCHING_UP
            else:
                state = ENT_ON_TRACK
        row = {"studentNumber": sn, "Entrustment state": state,
               "Entrustment avg | Overall": overall.at[sn, "Entrustment avg"] if sn in overall.index else np.nan,
               "Low share (L1-2) | Overall": overall.at[sn, "Low share (L1-2)"] if sn in overall.index else np.nan}
        for w in windows:
            row[f"Entrustment avg | {w}"] = avg.at[sn, w]
        for w in windows:
            row[f"Entrustment z | {w}"] = z.at[sn, w]
        row["Entrustment z change"] = (zs.iloc[-1] - zs.iloc[0]) if len(zs) >= 2 else np.nan
        rows.append(row)
    res = pd.DataFrame(rows)
    res.attrs["cohortMean"] = cohortMean.to_dict()
    return res


def evaluateWeaknessTrajectory(trajectoryLong, eligible, k=DEFAULT_TRAJECTORY_K,
                               minForms=DEFAULT_TRAJECTORY_MIN_FORMS, minTagged=DEFAULT_WEAKNESS_MIN_TAGGED):
    """One row per student x weakness category: rate + z per window and a STATUS —
        Persistent : HIGH (z >= +k and >= minTagged tagged forms) in EVERY judged window (>= 2)  -> FLAG
        Emerging   : HIGH in the last judged window but not the first                           -> watch
        Resolving  : HIGH in the first judged window but not the last
        Clear      : otherwise.
    Cohort-relative per category + window, so heavily-used categories (Technical Skills) and rare
    ones (Professional Behaviour) are each judged against their own norm."""
    windows = _trajectoryWindows(trajectoryLong)
    t = trajectoryLong
    nAssessed = (t[t["metric"] == "Forms assessor-submitted"]
                 .pivot(index="studentNumber", columns="window", values="value").reindex(columns=windows).fillna(0))
    rows = []
    for cat in WEAKNESS_LABELS:
        def _wide(metric):
            return (t[t["metric"] == metric].pivot(index="studentNumber", columns="window", values="value")
                    .reindex(columns=windows))
        rate, tagged = _wide(f"{cat} rate"), _wide(f"{cat} forms").fillna(0)
        z, cohortMean = _zByWindow(rate, nAssessed.reindex(rate.index).fillna(0), eligible, minForms)
        high = (z >= k) & (tagged >= minTagged)
        for sn in rate.index:
            judgedCols = [w for w in windows if pd.notna(z.at[sn, w])]
            h = high.loc[sn, judgedCols] if judgedCols else pd.Series(dtype=bool)
            if len(h) >= 2 and h.all():
                status = WEAK_PERSISTENT
            elif len(h) >= 2 and h.iloc[-1] and not h.iloc[0]:
                status = WEAK_EMERGING
            elif len(h) >= 2 and h.iloc[0] and not h.iloc[-1]:
                status = WEAK_RESOLVING
            else:
                status = WEAK_CLEAR
            row = {"studentNumber": sn, "category": cat, "status": status}
            for w in windows:
                row[f"rate | {w}"] = rate.at[sn, w]
                row[f"cohort | {w}"] = cohortMean.get(w, np.nan)
                row[f"high | {w}"] = bool(high.at[sn, w]) if pd.notna(z.at[sn, w]) else None
            rows.append(row)
    return pd.DataFrame(rows)


# ═══════════════════════════════════════════════════════════════════════════
# 8. Summary table (one row per student, every flag in one place)
# ═══════════════════════════════════════════════════════════════════════════

SUMMARY_COUNT_METRICS = ["Forms", "Patients attended", "Age 0-6", "Age 7-17", "Age 18+",
                         "Items total", "Distinct item codes"]


def buildFlagSummary(result, countFlags, entTraj, weakTraj):
    """Assemble the Summary. Returns (summaryDf, statusDf) — statusDf has the same shape and
    holds 'low' / 'watch' / 'ok' / 'n/a' / '' per cell to drive the colours.
    'Flags' counts RED items only, one per: count metric (any window low), low coverage section,
    entrustment state, each persistent weakness category, each persistently-low / persistently
    high-weakness section."""
    mw = result["metricsWide"].set_index("studentNumber")
    insufficient = set(result["insufficient"]["studentNumber"])
    ent = entTraj.set_index("studentNumber")
    windows = [w for w in dict.fromkeys(result["metricsLong"]["window"])]
    cf = countFlags
    lowCf = cf[cf["status"] == STATUS_LOW]
    isCov = cf["metric"].str.startswith((SECTION_PREFIX, SUBSECTION_PREFIX))

    rows, stats = [], []
    for sn in mw.index:
        row, st = {"ID": sn, "Student": mw.at[sn, "studentName"]}, {}
        flagTexts, watchTexts = [], []

        # ── count metrics: value per window, coloured; one flag per metric ──
        mine = cf[cf["studentNumber"] == sn]
        for metric in SUMMARY_COUNT_METRICS:
            lowWindows = []
            for w in windows:
                col = f"{metric} | {w}"
                if col not in mw.columns:
                    continue
                row[col] = mw.at[sn, col]
                hit = mine[(mine["metric"] == metric) & (mine["window"] == w)]
                st[col] = hit["status"].iloc[0] if len(hit) else ""
                if st[col] == STATUS_LOW:
                    lowWindows.append(w)
            if lowWindows:
                flagTexts.append(f"Low {metric} ({', '.join(lowWindows)})")
        dsl = mine[(mine["metric"] == "Days since last form") & (mine["window"] == OVERALL_WINDOW)]
        row["Days since last form"] = mw.at[sn, "Days since last form | Overall"] if "Days since last form | Overall" in mw.columns else np.nan
        st["Days since last form"] = dsl["status"].iloc[0] if len(dsl) else ""
        if st["Days since last form"] == STATUS_LOW:
            flagTexts.append("Inactive recently")

        # ── coverage: sections / sub-sections below threshold (Overall) ──
        for prefix, label in ((SECTION_PREFIX, "Low-volume sections"), (SUBSECTION_PREFIX, "Low-volume sub-sections")):
            sel = lowCf[(lowCf["studentNumber"] == sn) & (lowCf["window"] == OVERALL_WINDOW)
                        & lowCf["metric"].str.startswith(prefix)]
            names = sorted(m[len(prefix):] for m in sel["metric"])
            row[label] = "; ".join(names)
            st[label] = STATUS_LOW if names else STATUS_OK
            if names and prefix == SECTION_PREFIX:
                flagTexts += [f"Low volume: {n}" for n in names]

        # ── entrustment trajectory ──
        if sn in ent.index:
            for c in [c for c in ent.columns if c.startswith(("Entrustment avg", "Low share"))]:
                row[c] = ent.at[sn, c]
                w = c.split(" | ")[1]
                zc = f"Entrustment z | {w}"
                if zc in ent.columns and pd.notna(ent.at[sn, zc]):
                    st[c] = STATUS_LOW if ent.at[sn, zc] <= -DEFAULT_TRAJECTORY_K else (
                        STATUS_WATCH if ent.at[sn, zc] <= -DEFAULT_SECTION_WATCH_K else STATUS_OK)
            state = ent.at[sn, "Entrustment state"]
            row["Entrustment state"] = state
            st["Entrustment state"] = (STATUS_LOW if state in ENT_FLAG_STATES else STATUS_WATCH if state in ENT_WATCH_STATES
                                       else STATUS_NA if state == ENT_INSUFFICIENT else STATUS_OK)
            if state in ENT_FLAG_STATES:
                flagTexts.append(f"Entrustment: {state.lower()}")
            elif state in ENT_WATCH_STATES:
                watchTexts.append(f"Entrustment: {state.lower()}")

        # ── weakness trajectory ──
        w_ = weakTraj[weakTraj["studentNumber"] == sn]
        for status, label, level in ((WEAK_PERSISTENT, "Persistent weakness areas", STATUS_LOW),
                                     (WEAK_EMERGING, "Emerging weakness areas", STATUS_WATCH),
                                     (WEAK_RESOLVING, "Resolving weakness areas", STATUS_OK)):
            cats = list(w_.loc[w_["status"] == status, "category"])
            row[label] = "; ".join(cats)
            st[label] = level if cats else (STATUS_OK if status != WEAK_RESOLVING else "")
            if status == WEAK_PERSISTENT:
                flagTexts += [f"Persistent weakness: {c}" for c in cats]
            elif status == WEAK_EMERGING:
                watchTexts += [f"Emerging weakness: {c}" for c in cats]

        # ── section-level entrustment / weakness (from stage 1b) ──
        for col, tag in (("Sections persistently low entrustment", "Section entrustment low"),
                         ("Sections persistently high weakness", "Section weakness high")):
            val = mw.at[sn, col] if col in mw.columns and pd.notna(mw.at[sn, col]) else ""
            row[col] = val
            st[col] = STATUS_LOW if val else STATUS_OK
            flagTexts += [f"{tag}: {s}" for s in val.split("; ") if s]
        for col in ("Sections low entrustment", "Sections high weakness"):
            val = mw.at[sn, col] if col in mw.columns and pd.notna(mw.at[sn, col]) else ""
            row[col + " (Overall)"] = val
            st[col + " (Overall)"] = STATUS_WATCH if val else STATUS_OK
        strongCol = "Strong throughout (Sections)"
        row["Strong in every section"] = bool(mw.at[sn, strongCol]) if strongCol in mw.columns and pd.notna(mw.at[sn, strongCol]) else False
        st["Strong in every section"] = STATUS_OK if row["Strong in every section"] else ""

        # ── context (never flagged, shown for the panel) ──
        for col in [c for c in mw.columns if c.startswith(("Incident forms | Overall", "Concern forms | Overall", "PBNs", "Leave days"))]:
            row[col.replace(" | Overall", "")] = mw.at[sn, col]

        if sn in insufficient:
            flagTexts, watchTexts = ["Insufficient activity (under forms floor)"], []
        row["Flags"], row["Watch"] = len(flagTexts), len(watchTexts)
        row["What flagged"] = "\n".join(flagTexts)
        row["Watch items"] = "\n".join(watchTexts)
        rows.append(row)
        stats.append(st)

    summary = pd.DataFrame(rows)
    lead = ["ID", "Student", "Flags", "Watch", "What flagged", "Watch items"]
    summary = summary[lead + [c for c in summary.columns if c not in lead]]
    status = pd.DataFrame(stats).reindex(columns=summary.columns).fillna("")
    order = summary.sort_values(["Flags", "Watch", "Student"], ascending=[False, False, True]).index
    return summary.loc[order].reset_index(drop=True), status.loc[order].reset_index(drop=True)


# ═══════════════════════════════════════════════════════════════════════════
# 9. Styled workbook
# ═══════════════════════════════════════════════════════════════════════════

FLAG_NAVY = "010D44"
FLAG_TEXT_COLOURS = {STATUS_LOW: "9C0006", STATUS_WATCH: "9C6500", STATUS_OK: "276221", STATUS_NA: "7F7F7F"}


def styleFlagSheet(ws, headerRow=1, wideCols=(), textCols=(), freeze="C2", numberFormats=None, headerHeight=42):
    """House style for every sheet: navy header, wrapped header text, borders, freeze, filter,
    sensible widths. wideCols = name-like columns, textCols = long wrapped text columns."""
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    side = Side(style="thin", color="D9D9D9")
    border = Border(left=side, right=side, top=side, bottom=side)
    headers = [c.value for c in ws[headerRow]]
    for j, name in enumerate(headers, start=1):
        h = ws.cell(row=headerRow, column=j)
        h.value = str(name).replace(" | ", "\n") if name is not None else name
        h.fill = PatternFill("solid", fgColor=FLAG_NAVY)
        h.font = Font(bold=True, color="FFFFFF", size=10)
        h.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        h.border = border
        letter = get_column_letter(j)
        if name in textCols:
            ws.column_dimensions[letter].width = 46
        elif name in wideCols:
            ws.column_dimensions[letter].width = 26
        else:
            longest = max((len(p) for p in str(name).replace(" | ", "\n").split("\n")), default=8)
            ws.column_dimensions[letter].width = max(10, min(22, longest + 2))
        fmt = (numberFormats or {}).get(name)
        for i in range(headerRow + 1, ws.max_row + 1):
            c = ws.cell(row=i, column=j)
            c.border = border
            c.font = Font(size=10)
            if name in textCols:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(vertical="top", horizontal="left" if name in wideCols else "center")
            if fmt and isinstance(c.value, (int, float)):
                c.number_format = fmt
            elif isinstance(c.value, float):
                c.number_format = "0.00"
    ws.row_dimensions[headerRow].height = headerHeight
    ws.freeze_panes = freeze
    if ws.max_row > headerRow:
        ws.auto_filter.ref = f"A{headerRow}:{get_column_letter(ws.max_column)}{ws.max_row}"


def applyStatusFills(ws, statusDf, headerRow=1, boldLow=True):
    """Colour cells from a same-shape status DataFrame ('low' red / 'watch' amber / 'ok' green /
    'n/a' grey / '' untouched)."""
    from openpyxl.styles import PatternFill, Font
    fills = {s: PatternFill("solid", fgColor=hexCol) for s, hexCol in SECTION_STATUS_FILLS.items()}
    for j, col in enumerate(statusDf.columns, start=1):
        colStatus = statusDf[col].values
        for i, s in enumerate(colStatus):
            if s in fills:
                c = ws.cell(row=headerRow + 1 + i, column=j)
                c.fill = fills[s]
                c.font = Font(size=10, bold=(boldLow and s == STATUS_LOW), color=FLAG_TEXT_COLOURS[s])


def _writeStyled(writer, df, sheetName, statusDf=None, **styleArgs):
    df.to_excel(writer, sheet_name=sheetName, index=False)
    ws = writer.sheets[sheetName]
    styleFlagSheet(ws, **styleArgs)
    if statusDf is not None:
        applyStatusFills(ws, statusDf)
    return ws


def _countFlagPivot(countFlags, metricsFilter):
    """(valuesWide, statusWide) for a set of count-flag rows, columns '<metric> | <window>'."""
    cf = countFlags[metricsFilter(countFlags["metric"])]
    if cf.empty:
        return pd.DataFrame(), pd.DataFrame()
    metricOrder, windowOrder = list(dict.fromkeys(cf["metric"])), list(dict.fromkeys(cf["window"]))
    nameMap = cf.drop_duplicates("studentNumber").set_index("studentNumber")["studentName"]
    def _wide(col):
        w = cf.pivot(index="studentNumber", columns=["metric", "window"], values=col)
        cols = [(m, x) for m in metricOrder for x in windowOrder if (m, x) in w.columns]
        w = w[cols]
        w.columns = [f"{m.replace(SECTION_PREFIX, '').replace(SUBSECTION_PREFIX, '')} | {x}" for m, x in cols]
        w = w.reset_index()
        w.insert(1, "studentName", w["studentNumber"].map(nameMap))
        return w.sort_values("studentName", key=lambda s: s.astype(str).str.casefold()).reset_index(drop=True)
    vals, stat = _wide("value"), _wide("status")
    stat[["studentNumber", "studentName"]] = ""
    return vals, stat.fillna("")


def buildFlagLegend(cohort, countFlags, settings):
    """Legend rows: how to read the colours, every resolved threshold, and the run settings."""
    rows = [("HOW TO READ", "", ""),
            ("Red", "Flag — below the threshold / consistently low / persistent weakness", ""),
            ("Amber", "Watch — drifting low, emerging weakness, or low on Overall only", ""),
            ("Green", "Judged and fine", ""),
            ("Grey", "Not judged — too few forms, rotation group not reached, or rule unusable", ""),
            ("", "", ""),
            ("Count thresholds", "fixed = your number (inclusive) · meanStd = mean - k*SD · meanFrac = f*mean; "
             "mean/SD use students above the forms floor only", ""),
            ("Entrustment state", "z-score vs cohort in the SAME window. Consistently low = low in every judged window; "
             "Falling behind = low in the last window and dropped >= 1 SD; Low recently = low in the last window only", ""),
            ("Weakness status", "Share of assessor-submitted forms tagged in the category, z vs cohort per window. "
             "Persistent = high in every judged window; Emerging = high only recently; Resolving = high only early", ""),
            ("Section entrustment / weakness", "A form counts toward every section its item codes touch (co-occurrence). "
             "Judged against the cohort in the same section + window", ""),
            ("", "", ""), ("RUN SETTINGS", "", "")]
    rows += [(k, str(v), "") for k, v in settings.items()]
    legend = pd.DataFrame(rows, columns=["Item", "Meaning / value", "Note"])
    thr = (countFlags.groupby(["metric", "window"], sort=False)
           .agg(mode=("mode", "first"), param=("param", "first"), direction=("direction", "first"),
                threshold=("threshold", "first"), usable=("usable", "first"),
                flagged=("status", lambda s: int((s == STATUS_LOW).sum()))).reset_index())
    thr["threshold"] = thr["threshold"].round(2)
    return legend, thr


def buildLaggingFlagWorkbook(engine, cohort, outPath, thresholds=None, rotationGroups=None,
                             trajectoryGroups=None, excludeStudentNumbers=None,
                             minFormsFloor=DEFAULT_MIN_FORMS_FLOOR, expectedMinShare=DEFAULT_EXPECTED_MIN_SHARE,
                             trajectoryK=DEFAULT_TRAJECTORY_K, trajectoryMinForms=DEFAULT_TRAJECTORY_MIN_FORMS,
                             weaknessMinTagged=DEFAULT_WEAKNESS_MIN_TAGGED, fallingDrop=DEFAULT_FALLING_DROP,
                             sectionK=DEFAULT_SECTION_K, sectionWatchK=DEFAULT_SECTION_WATCH_K,
                             sectionMinForms=DEFAULT_SECTION_MIN_FORMS, sectionMinStudents=DEFAULT_SECTION_MIN_STUDENTS,
                             formsTable="dds4_boh3_forms_v3", year=2026, result=None):
    """THE deliverable: one styled workbook per cohort.
        Summary (flags first) · Counts · Coverage (Sections) · Coverage (Sub-sections) ·
        Entrustment Trajectory · Weakness Trajectory · Sec/SubSec Entrustment + Weakness ·
        Thresholds · Legend · Insufficient.
    thresholds: {metric: spec} overrides merged over DEFAULT_COUNT_THRESHOLDS (see its comment).
    result: pass a dict from buildFlaggingMetrics to skip re-querying. Returns the result dict
    extended with countFlags / entTrajectory / weakTrajectory / summary."""
    if result is None:
        result = buildFlaggingMetrics(engine, cohort, formsTable, rotationGroups, trajectoryGroups,
                                      excludeStudentNumbers, minFormsFloor, expectedMinShare, year=year,
                                      sectionK=sectionK, sectionWatchK=sectionWatchK,
                                      sectionMinForms=sectionMinForms, sectionMinStudents=sectionMinStudents)
    eligible = getEligibleStudents(result["metricsLong"], minFormsFloor)
    countFlags = evaluateCountFlags(result["metricsLong"], result["expectedSections"], thresholds, minFormsFloor)
    entTraj = evaluateEntrustmentTrajectory(result["trajectoryLong"], eligible, trajectoryK, trajectoryMinForms, fallingDrop)
    weakTraj = evaluateWeaknessTrajectory(result["trajectoryLong"], eligible, trajectoryK, trajectoryMinForms, weaknessMinTagged)
    summary, summaryStatus = buildFlagSummary(result, countFlags, entTraj, weakTraj)
    result.update(countFlags=countFlags, entTrajectory=entTraj, weakTrajectory=weakTraj,
                  summary=summary, summaryStatus=summaryStatus)

    nameMap = result["metricsWide"].set_index("studentNumber")["studentName"]
    pctCols = {c: "0%" for c in summary.columns if c.startswith("Low share")}
    textCols = ("What flagged", "Watch items", "Low-volume sections", "Low-volume sub-sections",
                "Persistent weakness areas", "Emerging weakness areas", "Resolving weakness areas",
                "Sections persistently low entrustment", "Sections persistently high weakness",
                "Sections low entrustment (Overall)", "Sections high weakness (Overall)")

    with pd.ExcelWriter(outPath, engine="openpyxl") as writer:
        # 1 ── Summary
        ws = _writeStyled(writer, summary, "Summary", summaryStatus, wideCols=("Student",), textCols=textCols,
                          freeze="E2", numberFormats=pctCols)
        from openpyxl.styles import PatternFill, Font
        for i in range(len(summary)):                     # Flags / Watch counters
            for j, col, fill, colour in ((3, "Flags", "FFC7CE", "9C0006"), (4, "Watch", "FFEB9C", "9C6500")):
                if summary.at[i, col] > 0:
                    c = ws.cell(row=2 + i, column=j)
                    c.fill, c.font = PatternFill("solid", fgColor=fill), Font(bold=True, size=11, color=colour)

        # 2 ── Counts + coverage pivots (coloured by threshold)
        isSec = lambda m: m.str.startswith(SECTION_PREFIX)
        isSub = lambda m: m.str.startswith(SUBSECTION_PREFIX)
        for name, filt in (("Counts", lambda m: ~(isSec(m) | isSub(m))),
                           ("Coverage (Sections)", isSec), ("Coverage (Sub-sections)", isSub)):
            vals, stat = _countFlagPivot(countFlags, filt)
            if not vals.empty:
                _writeStyled(writer, vals, name, stat, wideCols=("studentName",), headerHeight=60)

        # 3 ── Entrustment trajectory
        entOut = entTraj.copy()
        entOut.insert(1, "studentName", entOut["studentNumber"].map(nameMap))
        entOut = entOut.sort_values("Entrustment z change").reset_index(drop=True)
        entStat = pd.DataFrame("", index=entOut.index, columns=entOut.columns)
        for c in [c for c in entOut.columns if c.startswith("Entrustment avg | ") and not c.endswith(OVERALL_WINDOW)]:
            zc = c.replace("Entrustment avg", "Entrustment z")
            entStat[c] = np.where(entOut[zc].isna(), STATUS_NA, np.where(entOut[zc] <= -trajectoryK, STATUS_LOW,
                                  np.where(entOut[zc] <= -sectionWatchK, STATUS_WATCH, STATUS_OK)))
            entStat[zc] = entStat[c]
        entStat["Entrustment state"] = entOut["Entrustment state"].map(
            lambda s: STATUS_LOW if s in ENT_FLAG_STATES else STATUS_WATCH if s in ENT_WATCH_STATES
            else STATUS_NA if s == ENT_INSUFFICIENT else STATUS_OK)
        _writeStyled(writer, entOut, "Entrustment Trajectory", entStat, wideCols=("studentName", "Entrustment state"),
                     numberFormats={c: "0%" for c in entOut.columns if c.startswith("Low share")})

        # 4 ── Weakness trajectory: student x (category | window) rates + status per category
        windows = _trajectoryWindows(result["trajectoryLong"])
        wRows, wStats = [], []
        for sn, g in weakTraj.groupby("studentNumber", sort=False):
            row, st = {"studentNumber": sn, "studentName": nameMap.get(sn)}, {}
            for _, r in g.iterrows():
                for w in windows:
                    col = f"{r['category']} | {w}"
                    row[col] = r[f"rate | {w}"]
                    st[col] = STATUS_NA if r[f"high | {w}"] is None else (STATUS_LOW if r[f"high | {w}"] else STATUS_OK)
                col = f"{r['category']} | status"
                row[col] = r["status"]
                st[col] = {WEAK_PERSISTENT: STATUS_LOW, WEAK_EMERGING: STATUS_WATCH, WEAK_RESOLVING: STATUS_OK}.get(r["status"], "")
            wRows.append(row); wStats.append(st)
        weakWide = pd.DataFrame(wRows)
        weakStat = pd.DataFrame(wStats).reindex(columns=weakWide.columns).fillna("")
        order = weakWide.sort_values("studentName", key=lambda s: s.astype(str).str.casefold()).index
        _writeStyled(writer, weakWide.loc[order].reset_index(drop=True), "Weakness Trajectory",
                     weakStat.loc[order].reset_index(drop=True), wideCols=("studentName",), headerHeight=60,
                     numberFormats={c: "0%" for c in weakWide.columns if not c.endswith("status")})

        # 5 ── Section-level entrustment / weakness pivots (stage 1b)
        for key, tag in (("sectionPerf", "Sec"), ("subSectionPerf", "SubSec")):
            perf = result[key]
            if perf.empty:
                continue
            for valueCol, statusCol, name, title, fmt in (
                    ("entrustmentAvg", "entStatus", f"{tag} Entrustment",
                     "Mean entrustment by section x window — colour = vs cohort in the SAME section+window", "0.00"),
                    ("weaknessRate", "weakStatus", f"{tag} Weakness",
                     "Share of forms with >=1 weakness tag by section x window — colour = vs cohort", "0%")):
                vals, stat = pivotSectionPerformance(perf, valueCol, statusCol)
                writeSectionPivotSheet(writer, vals, stat, name, title, fmt)

        # 6 ── Thresholds, Legend, Insufficient
        settings = {"Cohort": cohort, "Students": len(summary), "Forms floor": minFormsFloor,
                    "Expected-section share": expectedMinShare,
                    "Rotation groups (counts)": [w for w in dict.fromkeys(result["metricsLong"]["window"]) if w != OVERALL_WINDOW],
                    "Trajectory windows": windows, "Trajectory k": trajectoryK,
                    "Trajectory min forms / window": trajectoryMinForms, "Weakness min tagged forms": weaknessMinTagged,
                    "Falling-behind z drop": fallingDrop, "Section k / watch k": f"{sectionK} / {sectionWatchK}",
                    "Section min forms / min students": f"{sectionMinForms} / {sectionMinStudents}"}
        legend, thr = buildFlagLegend(cohort, countFlags, settings)
        thrStat = pd.DataFrame("", index=thr.index, columns=thr.columns)
        thrStat["usable"] = np.where(thr["usable"], STATUS_OK, STATUS_NA)
        _writeStyled(writer, thr, "Thresholds", thrStat, wideCols=("metric",), freeze="A2")
        _writeStyled(writer, legend, "Legend", None, textCols=("Meaning / value",), wideCols=("Item",), freeze="A2")
        _writeStyled(writer, result["insufficient"], "Insufficient", None, wideCols=("studentName",), freeze="A2")

    print(f"[{cohort}] {int((summary['Flags'] > 0).sum())} of {len(summary)} students carry >= 1 flag "
          f"(>=3 flags: {int((summary['Flags'] >= 3).sum())}) -> {outPath}")
    return result
